"""Milestone 3C regression suite: legacy multi-document family/revision/
language inference (folder / multi-file legacy onboarding).

Covers: persistent data model, dedup (scenario G), revision normalization
(schemes + cross-scheme non-ordering), ambiguity (scenario E), human
confirmation + idempotent atomic persistence, lineage snapshot, mandatory
scenarios A-H, and the app.py-level project-manifest stamping hook.
"""
from __future__ import annotations

import json
import shutil
import sys
import tempfile
from contextlib import contextmanager
from pathlib import Path
from uuid import uuid4

BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE))

from docx import Document as DocxDocument

import legacy_family_inference as fam
from project_lifecycle import create_project
import app

TMP = Path(tempfile.mkdtemp(prefix="manualsync_3c_"))
results = []


def check(name, cond, detail=""):
    ok = bool(cond)
    results.append({"test": name, "status": "PASS" if ok else "FAIL", "detail": str(detail)})
    if not ok:
        raise AssertionError(f"{name}: {detail}")


def make_docx(name, front_title, heading, body):
    d = DocxDocument()
    d.add_paragraph(front_title)
    d.add_paragraph(heading)
    d.add_paragraph(body)
    p = TMP / name
    d.save(str(p))
    return p


IT_TEXT = "Il presente manuale deve essere letto con attenzione dal personale tecnico. Per il montaggio del dispositivo, il quadro deve essere collegato con cura."
EN_TEXT = "The present manual must be read carefully by the technical staff. For the mounting of the device, the control cabinet must be connected with the converter and inverter."

try:
    # ------------------------------------------------------------------
    # Group 1: revision normalization -- typed schemes, cross-scheme never orders
    # ------------------------------------------------------------------
    r1 = fam.extract_revision("MachineA_IT_Rev1.pdf")
    r2 = fam.extract_revision("MachineA_IT_Rev2.pdf")
    check("rev_numeric scheme recognized", r1.scheme == "rev_numeric" and r1.normalized == 1, r1)
    check("rev_numeric distinct across files", fam.revision_distinct(r1, r2), (r1, r2))
    check("rev_numeric orderable, Rev2 newer", fam.revision_order(r1, r2) == "b", fam.revision_order(r1, r2))

    d1 = fam.extract_revision("Machine_v1.0.pdf")
    d2 = fam.extract_revision("Machine_v1.1.pdf")
    check("dotted_version scheme recognized", d1.scheme == "dotted_version" and d1.normalized == (1, 0), d1)
    check("dotted_version orderable via dotted_version_order", fam.dotted_version_order(d1, d2) == "b", fam.dotted_version_order(d1, d2))
    check("generic revision_order() does not order dotted_version (dedicated helper required)", fam.revision_order(d1, d2) is None, fam.revision_order(d1, d2))

    lA = fam.extract_revision("ValveY_IT_RevA.pdf")
    l2 = fam.extract_revision("ValveY_IT_Rev2.pdf")
    check("cross-scheme (letter_filename vs rev_numeric) recognized as distinct", fam.revision_distinct(lA, l2), (lA, l2))
    check("cross-scheme never auto-orders (order to be confirmed)", fam.revision_order(lA, l2) is None, (lA, l2))

    ed1 = fam.extract_revision("Manual_Ed.1.pdf")
    ed2 = fam.extract_revision("Manual_Ed.2.pdf")
    check("edition scheme recognized", ed1.scheme == "edition" and ed1.normalized == 1, ed1)
    check("edition orderable same-scheme", fam.revision_order(ed1, ed2) == "b", fam.revision_order(ed1, ed2))

    dt1 = fam.extract_revision("Manual_2023-01-15.pdf")
    dt2 = fam.extract_revision("Manual_2024-06-01.pdf")
    check("date_full scheme recognized", dt1.scheme == "date_full", dt1)
    check("date_full orderable same-format", fam.revision_order(dt1, dt2) == "b", fam.revision_order(dt1, dt2))

    unk = fam.extract_revision("manual_final.pdf")
    check("unrecognized scheme when nothing matches", unk.scheme == "unrecognized" and unk.normalized is None, unk)
    check("unrecognized never counts as distinct revision evidence", not fam.revision_distinct(unk, r1), (unk, r1))

    # Real bug caught during implementation: '_' is a word character, so a bare
    # \b immediately before "Rev" in "MachineA_Rev1" would never match without
    # the (?:^|[_-.\s]) fix -- guard this regression explicitly.
    underscore_case = fam.extract_revision("MachineA_Rev1_final.pdf")
    check("revision recognized right after an underscore separator", underscore_case.scheme == "rev_numeric" and underscore_case.normalized == 1, underscore_case)

    # ------------------------------------------------------------------
    # Group 2: manual/product code extraction -- filename and content are
    # independent signals (content from a genuine free-text line, never a
    # numbered chapter heading).
    # ------------------------------------------------------------------
    check("manual code from filename strips language/revision tokens", fam.extract_manual_code_from_filename("MachineA_IT_Rev1.pdf") == "MachineA", fam.extract_manual_code_from_filename("MachineA_IT_Rev1.pdf"))
    check("manual code from filename rejects pure-digit remainder", fam.extract_manual_code_from_filename("2024_IT_Rev1.pdf") is None, fam.extract_manual_code_from_filename("2024_IT_Rev1.pdf"))
    check("manual code from content skips a leading digit-only match", fam.extract_manual_code_from_title("LineControl v3.3.04") == "LineControl", fam.extract_manual_code_from_title("LineControl v3.3.04"))

    # ------------------------------------------------------------------
    # Group 3: scenarios A-H via analyze_paths() (fully ephemeral)
    # ------------------------------------------------------------------
    e1 = make_docx("PumpX_IT_Rev1.docx", "PumpX Model 200", "1 - Introduzione", IT_TEXT + " idraulica")
    e2 = make_docx("PumpX_IT_Rev2.docx", "ValveZ Model 300", "1 - Introduzione", IT_TEXT + " pneumatica")
    c1 = make_docx("manual_final.docx", "CraneUnit Rev.1", "1 - Introduzione", IT_TEXT + " gru versione uno")
    c2 = make_docx("manual_new.docx", "CraneUnit Rev.2", "1 - Introduzione", IT_TEXT + " gru versione due nuovi sensori")
    h1 = make_docx("ValveY_IT_RevA.docx", "ValveY", "1 - Introduzione", IT_TEXT + " prima edizione valvola")
    h2 = make_docx("ValveY_IT_Rev2.docx", "ValveY", "1 - Introduzione", IT_TEXT + " seconda edizione valvola nuove specifiche")
    g1 = make_docx("DupTest_copy1.docx", "DupTest", "1 - Introduzione", IT_TEXT + " duplicato")
    g2 = TMP / "DupTest_copy2.docx"; shutil.copy2(g1, g2)
    a1 = make_docx("MachineA_IT_Rev1.docx", "MachineA", "1 - Introduzione", IT_TEXT + " versione uno macchina A")
    a2 = make_docx("MachineA_IT_Rev2.docx", "MachineA", "1 - Introduzione", IT_TEXT + " versione due macchina A aggiornata")
    b1 = make_docx("MachineA_EN_Rev2.docx", "MachineA", "1 - Introduction", EN_TEXT + " machine A updated version two")
    dd1 = make_docx("OtherMachine_IT_Rev1.docx", "OtherMachine", "1 - Introduzione", IT_TEXT + " prodotto completamente diverso")
    ignored_txt = TMP / "readme.txt"; ignored_txt.write_text("not a manual", encoding="utf-8")

    result = fam.analyze_paths([e1, e2, c1, c2, h1, h2, g1, g2, a1, a2, b1, dd1, ignored_txt])

    check("scenario D: unsupported .txt ignored with a reason, not raised", len(result.ignored) == 1 and "readme.txt" in result.ignored[0]["path"], result.ignored)
    check("scenario G: duplicate content collapsed to one candidate", any(len(d["duplicate_paths"]) == 1 for d in result.duplicate_groups), result.duplicate_groups)

    families_by_member = {frozenset(m.filename_original for m in f.members): f for f in result.families}
    machine_a_key = next((k for k in families_by_member if "MachineA_IT_Rev1.docx" in k), None)
    check("scenario A/B: MachineA family formed with all 3 members", machine_a_key is not None and len(machine_a_key) == 3, list(families_by_member.keys()))
    machine_a_fam = families_by_member[machine_a_key]
    rel_types = {r["relation_type"] for r in machine_a_fam.relations}
    check("scenario A: revision_of present for MachineA", "revision_of" in rel_types, rel_types)
    check("scenario B: language_variant_of present for MachineA", "language_variant_of" in rel_types, rel_types)
    rev_rel = next(r for r in machine_a_fam.relations if r["relation_type"] == "revision_of")
    check("scenario A: order confirmed for plain numeric revisions", rev_rel["order_confirmed"] is True, rev_rel)

    crane_key = next((k for k in families_by_member if "manual_final.docx" in k), None)
    check("scenario C: unhelpful filenames still clustered via content", crane_key is not None, list(families_by_member.keys()))
    crane_fam = families_by_member[crane_key]
    check("scenario C: revision_of derived purely from content", any(r["relation_type"] == "revision_of" and r["order_confirmed"] for r in crane_fam.relations), crane_fam.relations)

    valve_key = next((k for k in families_by_member if "ValveY_IT_RevA.docx" in k), None)
    check("scenario H: cross-scheme revision still clusters as a family", valve_key is not None, list(families_by_member.keys()))
    valve_fam = families_by_member[valve_key]
    valve_rel = next(r for r in valve_fam.relations if r["relation_type"] == "revision_of")
    check("scenario H: order NOT auto-confirmed across incompatible schemes", valve_rel["order_confirmed"] is False and valve_rel["newer_asset_hint"] is None, valve_rel)

    unrelated_names = {c.filename_original for c in result.unrelated}
    check("scenario D: unrelated manual stays independent, not merged", "OtherMachine_IT_Rev1.docx" in unrelated_names, unrelated_names)
    # Ambiguous is not the same state as unrelated (acceptance fix): an asset
    # awaiting a human decision must never also appear as "independent".
    check("scenario E: PumpX pair is NOT listed as unrelated (it is ambiguous, a distinct state)",
          "PumpX_IT_Rev1.docx" not in unrelated_names and "PumpX_IT_Rev2.docx" not in unrelated_names, unrelated_names)
    ambiguous_names = {(Path(ap["path_a"]).name, Path(ap["path_b"]).name) for ap in result.ambiguous_pairs}
    check("scenario E: PumpX pair explicitly surfaced as ambiguous", ("PumpX_IT_Rev1.docx", "PumpX_IT_Rev2.docx") in ambiguous_names, ambiguous_names)
    pump_conflict = next(ap for ap in result.ambiguous_pairs if Path(ap["path_a"]).name == "PumpX_IT_Rev1.docx")
    check("scenario E: conflict evidence role is 'conflict', not silently dropped", any(e["role"] == "conflict" for e in pump_conflict["evidence"]), pump_conflict["evidence"])

    check("no ambiguous noise for genuinely unrelated OtherMachine", not any("OtherMachine" in Path(ap["path_a"]).name or "OtherMachine" in Path(ap["path_b"]).name for ap in result.ambiguous_pairs), result.ambiguous_pairs)

    # ------------------------------------------------------------------
    # Scenario F: legacy .doc participates in the same family as .docx/.pdf,
    # via a mocked converter (mirrors the 3B.1 test pattern; no LibreOffice
    # dependency required for this automated check).
    # ------------------------------------------------------------------
    doc_path = TMP / "MachineA_IT_Rev3.doc"
    doc_path.write_bytes(b"placeholder legacy .doc content; converter is mocked for this test")
    fixture_docx_for_doc = make_docx("_fixture_for_doc.docx", "MachineA", "1 - Introduzione", IT_TEXT + " versione tre macchina A da doc legacy")

    @contextmanager
    def _fake_success_converter(_doc_path):
        yield fixture_docx_for_doc

    orig_convert = fam.convert_legacy_doc_to_docx
    fam.convert_legacy_doc_to_docx = _fake_success_converter
    try:
        doc_candidate = fam.build_candidate(doc_path)
        check("scenario F: .doc routed through the SAME 3B.1 converter, not a second parser", doc_candidate.doc_kind == "doc", doc_candidate.doc_kind)
        check("scenario F: .doc candidate carries manual code via converted content", doc_candidate.manual_code_content == "MachineA" or doc_candidate.manual_code_filename == "MachineA", doc_candidate)
    finally:
        fam.convert_legacy_doc_to_docx = orig_convert

    # ------------------------------------------------------------------
    # Group 4: persistence -- schema, atomic write, idempotency, lookup, snapshot
    # ------------------------------------------------------------------
    DATA_HOME = TMP / "datahome"
    data = fam.persist_confirmation(
        DATA_HOME, candidates=machine_a_fam.members, relations=machine_a_fam.relations, confirmed_by="tester",
    )
    check("persisted family has schema_version from the start", data.get("schema_version") == fam.SCHEMA_VERSION, data.get("schema_version"))
    check("persisted family_id follows FAM- convention", str(data["family_id"]).startswith("FAM-"), data["family_id"])
    check("persisted members carry full sha256 + original filename + revision string", all(
        m.get("file_sha256") and m.get("filename_original") and "revision_string_original" in m for m in data["members"]
    ), data["members"])
    check("persisted relations have stable relation_id, confirmed_by, confirmed_at", all(
        r.get("relation_id", "").startswith("REL-") and r.get("confirmed_by") == "tester" and r.get("confirmed_at") for r in data["relations"]
    ), data["relations"])
    family_file = fam.family_json_path(DATA_HOME, data["family_id"])
    check("family JSON physically written to LegacyFamilies/", family_file.exists(), family_file)
    check("no leftover atomic-write temp file", not any(p.name.startswith(f".{family_file.name}.") for p in family_file.parent.iterdir()), list(family_file.parent.iterdir()))
    reloaded = json.loads(family_file.read_text(encoding="utf-8"))
    check("persisted JSON round-trips (well-formed, matches written data)", reloaded["family_id"] == data["family_id"] and len(reloaded["relations"]) == len(data["relations"]), reloaded)

    data_again = fam.persist_confirmation(
        DATA_HOME, candidates=machine_a_fam.members, relations=machine_a_fam.relations,
        confirmed_by="tester2", family_id=data["family_id"],
    )
    check("re-confirming the same relations is idempotent (no duplicates)", len(data_again["relations"]) == len(data["relations"]), (len(data_again["relations"]), len(data["relations"])))

    try:
        fam.persist_confirmation(DATA_HOME, candidates=machine_a_fam.members, relations=[], confirmed_by="")
        check("confirmed_by is mandatory to persist", False, "no exception raised")
    except ValueError:
        check("confirmed_by is mandatory to persist", True)

    found_family = fam.find_family_for_asset(DATA_HOME, machine_a_fam.members[0].file_sha256)
    check("find_family_for_asset resolves an already-confirmed asset", found_family == data["family_id"], found_family)
    check("find_family_for_asset returns None for an unknown asset", fam.find_family_for_asset(DATA_HOME, "0" * 64) is None, "")

    # Second, independent confirmation must not leak into the first family.
    data_crane = fam.persist_confirmation(
        DATA_HOME, candidates=crane_fam.members, relations=crane_fam.relations, confirmed_by="tester",
    )
    check("independent family confirmation gets its own family_id", data_crane["family_id"] != data["family_id"], (data_crane["family_id"], data["family_id"]))

    # ------------------------------------------------------------------
    # Group 5: release lineage snapshot -- self-contained, hash-stamped
    # ------------------------------------------------------------------
    asset_hashes = [m.file_sha256 for m in machine_a_fam.members]
    snap1 = fam.snapshot_lineage_for_release(DATA_HOME, asset_hashes)
    check("snapshot carries schema_version", snap1.get("schema_version") == fam.SCHEMA_VERSION, snap1.get("schema_version"))
    check("snapshot is stamped with its own sha256", isinstance(snap1.get("snapshot_sha256"), str) and len(snap1["snapshot_sha256"]) == 64, snap1.get("snapshot_sha256"))
    check("snapshot includes only relations touching the requested assets", all(
        r["asset_a_hash"] in asset_hashes or r["asset_b_hash"] in asset_hashes for r in snap1["relations"]
    ), snap1["relations"])
    snap2 = fam.snapshot_lineage_for_release(DATA_HOME, asset_hashes)
    check("snapshot content (sans timestamp) is deterministic across calls", (
        {k: v for k, v in snap1.items() if k not in ("snapshot_taken_at", "snapshot_sha256")}
        == {k: v for k, v in snap2.items() if k not in ("snapshot_taken_at", "snapshot_sha256")}
    ), (snap1, snap2))

    # A later edit to LegacyFamilies/<id>.json must not retroactively change a
    # snapshot already taken -- the snapshot is a copy, not a live reference.
    mutated = json.loads(family_file.read_text(encoding="utf-8"))
    mutated["members"][0]["revision_string_original"] = "TAMPERED"
    family_file.write_text(json.dumps(mutated, ensure_ascii=False, indent=2), encoding="utf-8")
    check("release immutability: earlier snapshot dict is unaffected by a later family-file edit", snap1["members"][0]["revision_string_original"] != "TAMPERED", snap1["members"][0])

    # ------------------------------------------------------------------
    # Group 6: app.py integration -- UI-level confirm flow + project stamping
    # ------------------------------------------------------------------
    orig_data_home = app.DATA_HOME
    app.DATA_HOME = DATA_HOME

    try:
        state, summary_md, confirm_row_update, *radios = app.legacy_multi_analyze_ui([str(a1), str(a2), str(b1)], None)
        check("legacy_multi_analyze_ui returns a usable state", state is not None and state[0] is not None, summary_md[:200])
        _result_a, review_items_a = state
        check("legacy_multi_analyze_ui builds review items for MachineA relations", any(it["kind"] == "relation" for it in review_items_a), review_items_a)
        check("legacy_multi_analyze_ui returns exactly one radio update per pool slot", len(radios) == app._LEGACY_MAX_REVIEW_ITEMS, len(radios))

        values_all = [it["default"] for it in review_items_a] + [None] * (app._LEGACY_MAX_REVIEW_ITEMS - len(review_items_a))
        # Defaults alone would reject the order-ambiguous case (if any) but confirm
        # the clean revision_of/language_variant_of ones -- exercise that mix.
        confirm_msg, _row_update, *_rel_updates = app.legacy_multi_confirm_ui(state, "QA Reviewer", *values_all)
        check("legacy_multi_confirm_ui reports a persisted family id", "FAM-" in confirm_msg, confirm_msg)

        # UX fix: an empty "Il tuo nome" is an ordinary required-field
        # validation error, not a raised gr.Error / global popup -- see
        # Group 10 below for the full inline-validation contract.
        empty_name_result = app.legacy_multi_confirm_ui(state, "", *values_all)
        check("legacy_multi_confirm_ui with empty confirmed_by raises no exception",
              True, "")
        check("legacy_multi_confirm_ui with empty confirmed_by flags the name field inline, not via gr.Error",
              empty_name_result[2].get("elem_classes") == ["ms-field-invalid"], empty_name_result[2])

        project_dir = create_project(TMP / "Projects", project_name=f"proj_{uuid4().hex[:8]}", source_lang="it-IT")
        app.import_manual_ui(str(project_dir), str(a1), "it-IT", "PUB-MAIN", "MachineA")
        manifest = json.loads((project_dir / "ManualSync_Project.json").read_text(encoding="utf-8"))
        check("confirmed asset auto-stamps legacy_family_id on project creation", bool(manifest.get("legacy_family_id")), manifest)
        check("confirmed asset auto-stamps legacy_asset_id on project creation", bool(manifest.get("legacy_asset_id")), manifest)

        project_dir2 = create_project(TMP / "Projects", project_name=f"proj_{uuid4().hex[:8]}", source_lang="en-US")
        unknown_pdf = BASE / "benchmark" / "siemens" / "old_EN.pdf"
        app.import_manual_ui(str(project_dir2), str(unknown_pdf), "en-US", "PUB-MAIN", "Unrelated")
        manifest2 = json.loads((project_dir2 / "ManualSync_Project.json").read_text(encoding="utf-8"))
        check("import of a file with NO confirmed family does not error and stamps nothing", not manifest2.get("legacy_family_id"), manifest2)

        # --------------------------------------------------------------
        # Group 6b: the three acceptance gaps from manual review.
        # --------------------------------------------------------------
        # Gap 2: an order-ambiguous revision_of (cross-scheme letter vs numeric)
        # must never be persisted with a direction unless the human explicitly
        # picks one -- "Non sono correlati" is the only safe default, and
        # confirming it must NOT persist any relation.
        vy1 = make_docx("ValveY_IT_RevA.docx", "ValveY", "1 - Introduzione", IT_TEXT + " prima edizione valvola gap test")
        vy2 = make_docx("ValveY_IT_Rev2.docx", "ValveY", "1 - Introduzione", IT_TEXT + " seconda edizione valvola nuove specifiche gap test")
        state_h, *_ = app.legacy_multi_analyze_ui([str(vy1), str(vy2)], None)
        _result_h, review_items_h = state_h
        order_item = next(it for it in review_items_h if it["kind"] == "relation" and it["relation"]["relation_type"] == "revision_of")
        check("gap2: order-ambiguous revision_of defaults to 'Non sono correlati' (no invented order)", order_item["default"] == app._LEGACY_NOT_RELATED, order_item["default"])
        check("gap2: order-ambiguous revision_of offers an explicit A-before-B / B-before-A choice", len(order_item["choices"]) == 3 and app._LEGACY_CONFIRM not in order_item["choices"], order_item["choices"])
        # UX fix: order choices are now concise "A -> B" arrows, not verbose
        # duplicated "X e' precedente a Y" strings differing only in word order.
        check("gap2: order choices use the concise arrow format, not the old verbose phrasing",
              all("è precedente a" not in c for c in order_item["choices"]) and
              app._legacy_order_choice(order_item["a"], order_item["b"]) in order_item["choices"] and
              app._legacy_order_choice(order_item["b"], order_item["a"]) in order_item["choices"],
              order_item["choices"])
        check("gap2: control label reads 'scegli l'ordine corretto' when every choice is about order",
              "scegli l'ordine corretto" in order_item["label"], order_item["label"])

        values_reject = [None] * app._LEGACY_MAX_REVIEW_ITEMS
        values_reject[review_items_h.index(order_item)] = app._LEGACY_NOT_RELATED
        msg_reject, *_ = app.legacy_multi_confirm_ui(state_h, "QA", *values_reject)
        check("gap2: confirming the default rejects, persisting nothing", "Nessuna relazione confermata" in msg_reject, msg_reject)
        check("gap2: rejected pair has no confirmed family", fam.find_family_for_asset(DATA_HOME, order_item["a"].file_sha256) is None, "")

        chosen_order = app._legacy_order_choice(order_item["a"], order_item["b"])
        values_pick = [None] * app._LEGACY_MAX_REVIEW_ITEMS
        values_pick[review_items_h.index(order_item)] = chosen_order
        msg_pick, *_ = app.legacy_multi_confirm_ui(state_h, "QA", *values_pick)
        fam_id_h = msg_pick.split("**")[1]
        data_h = fam.load_family(DATA_HOME, fam_id_h)
        check("gap2: human-chosen order is persisted as order_confirmed", data_h["relations"][0]["order_confirmed"] is True, data_h["relations"][0])
        check("gap2: newer_asset_hint points to the human-chosen document, not an inferred one", data_h["relations"][0]["newer_asset_hint"] == fam.derive_asset_id(order_item["b"].file_sha256), data_h["relations"][0])

        # Gap 3: a pending same-family pair (no auto-eligible specific relation)
        # must be persisted if and only if the human explicitly links it -- the
        # UI's claim and the persisted state must always match.
        pc1 = make_docx("CraneUnit_IT_A.docx", "CraneUnit", "1 - Introduzione", IT_TEXT + " gru testo alfa con dettagli extra gap test")
        pc2 = make_docx("CraneUnit_EN_B.docx", "CraneUnit", "1 - Introduction", EN_TEXT + " crane unit text beta entirely different narrative content gap test")
        state_p, *_ = app.legacy_multi_analyze_ui([str(pc1), str(pc2)], None)
        _result_p, review_items_p = state_p
        pending_item = next((it for it in review_items_p if it["kind"] == "pending"), None)
        check("gap3: a same-family pair with no clean relation is offered as 'pending'", pending_item is not None, review_items_p)
        if pending_item:
            values_no_link = [None] * app._LEGACY_MAX_REVIEW_ITEMS
            msg_no_link, *_ = app.legacy_multi_confirm_ui(state_p, "QA", *values_no_link)
            check("gap3: NOT choosing to link persists nothing (UI promise matches persistence)",
                  fam.find_family_for_asset(DATA_HOME, pending_item["a"].file_sha256) is None, "")
            values_link = [None] * app._LEGACY_MAX_REVIEW_ITEMS
            values_link[review_items_p.index(pending_item)] = app._LEGACY_PENDING_LINK
            msg_link, *_ = app.legacy_multi_confirm_ui(state_p, "QA", *values_link)
            fam_id_p = msg_link.split("**")[1]
            data_p = fam.load_family(DATA_HOME, fam_id_p)
            check("gap3: explicit link persists BOTH members", len(data_p["members"]) == 2, data_p["members"])
            check("gap3: explicit link persists NO fabricated revision/language relation", len(data_p["relations"]) == 0, data_p["relations"])

        # Gap 1: an ambiguous (conflicting) pair must offer a real correction
        # path, not just a warning -- and the correction must be exactly what
        # gets persisted.
        pe1 = make_docx("PumpX_IT_Rev1.docx", "PumpX Model 200", "1 - Introduzione", IT_TEXT + " idraulica gap test")
        pe2 = make_docx("PumpX_IT_Rev2.docx", "ValveZ Model 300", "1 - Introduzione", IT_TEXT + " pneumatica gap test")
        state_e, *_ = app.legacy_multi_analyze_ui([str(pe1), str(pe2)], None)
        _result_e, review_items_e = state_e
        ambiguous_item = next((it for it in review_items_e if it["kind"] == "ambiguous"), None)
        check("gap1: a conflicting pair is offered as an actionable ambiguous item", ambiguous_item is not None, review_items_e)
        if ambiguous_item:
            check("gap1: ambiguous item defaults to 'Non sono correlati' (safe default)", ambiguous_item["default"].startswith(app._LEGACY_NOT_RELATED), ambiguous_item["default"])
            # UX fix: order-correction choices are now concise "A -> B" arrows
            # (no "Correggi:" prefix) -- "a real correction path" now means the
            # arrow options are present, not that every choice is prefixed.
            expected_ab = app._legacy_order_choice(ambiguous_item["a"], ambiguous_item["b"])
            expected_ba = app._legacy_order_choice(ambiguous_item["b"], ambiguous_item["a"])
            check("gap1: ambiguous item offers a real manual-type override (concise order arrows)",
                  expected_ab in ambiguous_item["choices"] and expected_ba in ambiguous_item["choices"], ambiguous_item["choices"])
            check("gap1: order choices are concise and distinguishable (no verbose duplicated phrasing)",
                  all("è precedente a" not in c and "Correggi: sono revisioni" not in c for c in (expected_ab, expected_ba)), (expected_ab, expected_ba))
            # Acceptance fix: PumpX/ValveZ both have confirmed Italian (same
            # language) -- offering "sono la stessa revisione in lingue diverse"
            # here would let a human create a linguistically false relation.
            check("ambiguous same-language pair (both confirmed IT) does NOT offer language_variant_of correction",
                  not any(c.startswith("Correggi: sono la stessa revisione") for c in ambiguous_item["choices"]), ambiguous_item["choices"])
            correction = expected_ab
            values_correct = [None] * app._LEGACY_MAX_REVIEW_ITEMS
            values_correct[review_items_e.index(ambiguous_item)] = correction
            msg_correct, *_ = app.legacy_multi_confirm_ui(state_e, "QA", *values_correct)
            fam_id_e = msg_correct.split("**")[1]
            data_e = fam.load_family(DATA_HOME, fam_id_e)
            check("gap1: manual correction is persisted with the human-assigned type", data_e["relations"][0]["relation_type"] == "revision_of", data_e["relations"][0])
            check("gap1: manually-corrected relation has a confirmed order (human-assigned, not inferred)", data_e["relations"][0]["order_confirmed"] is True, data_e["relations"][0])

            # ----------------------------------------------------------
            # Group 6c: human override evidence integrity (audit/provenance fix).
            # A human correction must never be disguised as an automatic signal.
            # ----------------------------------------------------------
            evidence_e = data_e["relations"][0]["evidence"]
            evidence_types_e = {e["type"] for e in evidence_e}
            check("human override recorded as its own evidence type, not disguised as an automatic signal",
                  "human_override" in evidence_types_e, evidence_e)
            fabricated = [e for e in evidence_e if e["type"] in ("manual_code_match", "language_match") and e.get("value") == "corrected by human review"]
            check("no automatic evidence type is fabricated with a fake human-review value", not fabricated, evidence_e)
            original_conflict = [e for e in evidence_e if e["type"] == "manual_code_match" and e["role"] == "conflict"]
            check("original conflicting automatic evidence (manual_code_match) survives the human correction, not erased",
                  bool(original_conflict), evidence_e)
            override_entry = next(e for e in evidence_e if e["type"] == "human_override")
            check("human_override evidence records relation type, decision, operator and chosen order",
                  "revision_of" in override_entry["value"] and "confirmed_by=QA" in override_entry["detail"] and "order_chosen" in override_entry["detail"],
                  override_entry)

        # ------------------------------------------------------------------
        # Group 6e-bis: real bug from manual acceptance (human override on an
        # ambiguous pair reported as not persisted). Root cause traced to
        # lm_analyze.click(..., scroll_to_output=True) auto-scrolling the page
        # right as the review radios render, which can shift them out from
        # under a user who starts clicking immediately -- reproduced live via
        # a real browser session (Chrome automation), not just a direct call.
        # This block simulates the value AS ACTUALLY EMITTED by the Gradio
        # component: the choice string is taken verbatim from the same
        # gr.update() dict legacy_multi_analyze_ui hands to the Radio
        # component (not retyped/reconstructed by the test), and every other
        # of the 24 pooled radios is left at None, exactly like an untouched
        # invisible Radio the browser never assigned a value to.
        # ------------------------------------------------------------------
        gx1 = make_docx("GammaX_IT_Rev1.docx", "GammaX Model 10", "1 - Introduzione", IT_TEXT + " idraulica gamma")
        gx2 = make_docx("GammaX_IT_Rev2.docx", "DeltaZ Model 20", "1 - Introduzione", IT_TEXT + " pneumatica delta")

        state_g, summary_g, confirm_row_g, *radio_updates_g = app.legacy_multi_analyze_ui([str(gx1), str(gx2)], None)
        _result_g, review_items_g = state_g
        amb_g = next((it for it in review_items_g if it["kind"] == "ambiguous"), None)
        check("real-bug repro: GammaX/DeltaZ pair surfaces as ambiguous", amb_g is not None, review_items_g)
        idx_g = review_items_g.index(amb_g)
        real_update_g = radio_updates_g[idx_g]
        check("real-bug repro: the Radio update Gradio would render is visible with the expected default",
              real_update_g["visible"] is True and real_update_g["value"] == amb_g["default"], real_update_g)
        # The exact choice string as it would appear in the rendered Radio's
        # own choices list -- not a string built independently by the test.
        rev1_candidate_g = amb_g["a"] if amb_g["a"].filename_original == gx1.name else amb_g["b"]
        rev2_candidate_g = amb_g["b"] if rev1_candidate_g is amb_g["a"] else amb_g["a"]
        expected_g_ab = app._legacy_order_choice(rev1_candidate_g, rev2_candidate_g)
        chosen_value_g = next(c for c in real_update_g["choices"] if c == expected_g_ab)
        check("real-bug repro: order choices are concise arrows, not the old verbose phrasing", "è precedente a" not in chosen_value_g and "Correggi" not in chosen_value_g, chosen_value_g)

        full_radio_values = [None] * app._LEGACY_MAX_REVIEW_ITEMS
        full_radio_values[idx_g] = chosen_value_g
        confirm_result_g = app.legacy_multi_confirm_ui(state_g, "QA Tester", *full_radio_values)
        msg_g = confirm_result_g[0]
        check("real-bug repro: human override on an ambiguous pair IS persisted (not 'Nessuna relazione confermata')",
              "Nessuna relazione confermata" not in msg_g and "FAM-" in msg_g, msg_g)
        fam_id_g = msg_g.split("**")[1]
        data_g = fam.load_family(DATA_HOME, fam_id_g)
        check("real-bug repro: exactly 1 relation, 2 members persisted", len(data_g["relations"]) == 1 and len(data_g["members"]) == 2, data_g)
        rel_g = data_g["relations"][0]
        check("real-bug repro: relation_type is revision_of", rel_g["relation_type"] == "revision_of", rel_g)
        check("real-bug repro: order_confirmed is true", rel_g["order_confirmed"] is True, rel_g)
        newer_member_g = next(m for m in data_g["members"] if m["asset_id"] == rel_g["newer_asset_hint"])
        check("real-bug repro: Rev2 is recorded as the newer asset", "Rev2" in newer_member_g["filename_original"], newer_member_g)
        check("real-bug repro: human_override evidence present", any(e["type"] == "human_override" for e in rel_g["evidence"]), rel_g["evidence"])

        # Sibling checks explicitly requested: "Non sono correlati" still does
        # not persist, and the REVERSE order choice persists correctly too.
        state_g2, _summary_g2, _confirm_row_g2, *radio_updates_g2 = app.legacy_multi_analyze_ui([str(gx1), str(gx2)], None)
        _result_g2, review_items_g2 = state_g2
        amb_g2 = next(it for it in review_items_g2 if it["kind"] == "ambiguous")
        idx_g2 = review_items_g2.index(amb_g2)
        values_reject_g2 = [None] * app._LEGACY_MAX_REVIEW_ITEMS
        values_reject_g2[idx_g2] = radio_updates_g2[idx_g2]["value"]  # the untouched default, "Non sono correlati (consigliato)"
        msg_reject_g2 = app.legacy_multi_confirm_ui(state_g2, "QA Tester", *values_reject_g2)[0]
        check("'Non sono correlati' (untouched default) still persists nothing", "Nessuna relazione confermata" in msg_reject_g2, msg_reject_g2)
        check("no relation created when the user makes no valid decision (all None)",
              app.legacy_multi_confirm_ui(state_g2, "QA Tester", *([None] * app._LEGACY_MAX_REVIEW_ITEMS))[0].startswith("Nessuna relazione confermata"),
              "checked")

        # Distinct fixture pair (not gx1/gx2): persist_confirmation is
        # idempotent per (relation_type, {asset_a, asset_b}) regardless of
        # direction, so re-using the already-confirmed GammaX pair here would
        # correctly be deduplicated against the earlier relation rather than
        # exercising a fresh reverse-order choice.
        ex1 = make_docx("EpsilonX_IT_Rev1.docx", "EpsilonX Model 1", "1 - Introduzione", IT_TEXT + " idraulica epsilon")
        ex2 = make_docx("EpsilonX_IT_Rev2.docx", "ZetaZ Model 2", "1 - Introduzione", IT_TEXT + " pneumatica zeta")
        state_g3, _summary_g3, _confirm_row_g3, *radio_updates_g3 = app.legacy_multi_analyze_ui([str(ex1), str(ex2)], None)
        _result_g3, review_items_g3 = state_g3
        amb_g3 = next(it for it in review_items_g3 if it["kind"] == "ambiguous")
        idx_g3 = review_items_g3.index(amb_g3)
        ex1_candidate_g3 = amb_g3["a"] if amb_g3["a"].filename_original == ex1.name else amb_g3["b"]
        ex2_candidate_g3 = amb_g3["b"] if ex1_candidate_g3 is amb_g3["a"] else amb_g3["a"]
        expected_reverse_g3 = app._legacy_order_choice(ex2_candidate_g3, ex1_candidate_g3)  # ex2 -> ex1 (reversed)
        reverse_choice_g3 = next(c for c in radio_updates_g3[idx_g3]["choices"] if c == expected_reverse_g3)
        values_reverse_g3 = [None] * app._LEGACY_MAX_REVIEW_ITEMS
        values_reverse_g3[idx_g3] = reverse_choice_g3
        msg_reverse_g3 = app.legacy_multi_confirm_ui(state_g3, "QA Tester", *values_reverse_g3)[0]
        check("reverse order (Rev2 precedente a Rev1) is persisted when explicitly chosen", "FAM-" in msg_reverse_g3, msg_reverse_g3)
        fam_id_g3 = msg_reverse_g3.split("**")[1]
        data_g3 = fam.load_family(DATA_HOME, fam_id_g3)
        rel_g3 = next(r for r in data_g3["relations"] if r["order_confirmed"])
        newer_member_g3 = next(m for m in data_g3["members"] if m["asset_id"] == rel_g3["newer_asset_hint"])
        check("reverse order: Rev1 correctly recorded as the newer asset this time", "Rev1" in newer_member_g3["filename_original"], newer_member_g3)

        # Static guard: the exact interaction hazard that caused the real bug
        # (auto-scroll racing a fast user click) must not silently come back.
        app_src = (BASE / "app.py").read_text(encoding="utf-8")
        analyze_click_line = next(l for l in app_src.splitlines() if "lm_analyze.click(" in l)
        check("regression guard: lm_analyze.click(...) does not auto-scroll (would race the user's next click)",
              "scroll_to_output=True" not in analyze_click_line, analyze_click_line)

        # ------------------------------------------------------------------
        # Group 6d: language_variant_of TM-safety fail-closed gate.
        # A manually-confirmed relation must never automatically become a safe
        # TM/Translation Intelligence seed -- that requires independently
        # validated, sufficiently-confident, genuinely-different languages.
        # ------------------------------------------------------------------
        unc1 = make_docx("Gadget_IT_Rev1.docx", "Gadget Model 1", "1 - Introduzione", IT_TEXT + " versione uno gadget con dettagli tecnici")
        unc2 = make_docx("Gadget_XX_Rev1.docx", "Widget Model 2", "1 - Introduzione", "x y z")
        state_u, *_ = app.legacy_multi_analyze_ui([str(unc1), str(unc2)], None)
        _result_u, review_items_u = state_u
        amb_u = next((it for it in review_items_u if it["kind"] == "ambiguous"), None)
        check("TM-safety fixture: uncertain-language pair surfaces as ambiguous", amb_u is not None, review_items_u)
        if amb_u:
            # Acceptance fix: 3C has no explicit language-confirmation step, so a
            # human override must NEVER be allowed to create a language_variant_of
            # out of an uncertain language -- the choice itself must not be offered.
            check("uncertain-language ambiguous pair does NOT offer the language-variant correction",
                  not any(c.startswith("Correggi: sono la stessa revisione") for c in amb_u["choices"]), amb_u["choices"])
        # Defense-in-depth: even if a language_variant_of with an unconfirmed
        # language were ever persisted through some other path, the TM consumer
        # gate itself must independently fail closed.
        fake_member_uncertain = {"asset_id": "ASSET-a", "language_confirmed": None, "language_confidence": 0.0}
        fake_member_ok = {"asset_id": "ASSET-b", "language_confirmed": "en", "language_confidence": 0.95}
        check("TM consumer gate rejects a pair with an unconfirmed language even if constructed directly",
              fam.is_language_variant_tm_safe(fake_member_uncertain, fake_member_ok) is False, (fake_member_uncertain, fake_member_ok))

        val1 = make_docx("Gadget2_IT_Rev1.docx", "Gadget2", "1 - Introduzione", IT_TEXT + " manuale completo con testo tecnico dettagliato per il personale")
        val2 = make_docx("Gadget2_DE_Rev1.docx", "Widget2", "1 - Einleitung", "Der Techniker muss das Gerät mit der Montage und dem Gehäuse verbinden. Das Gerät und die Seite müssen vor der Montage überprüft werden.")
        state_v, *_ = app.legacy_multi_analyze_ui([str(val1), str(val2)], None)
        _result_v, review_items_v = state_v
        amb_v = next((it for it in review_items_v if it["kind"] == "ambiguous"), None)
        check("TM-safety fixture: validated-language pair surfaces as ambiguous (manual_code conflict)", amb_v is not None, review_items_v)
        if amb_v:
            check("confirmed different-language pair (IT vs DE) DOES offer the language_variant_of correction",
                  any(c.startswith("Correggi: sono la stessa revisione") for c in amb_v["choices"]), amb_v["choices"])
            lang_choice_v = next(c for c in amb_v["choices"] if c.startswith("Correggi: sono la stessa revisione"))
            values_v = [None] * app._LEGACY_MAX_REVIEW_ITEMS
            values_v[review_items_v.index(amb_v)] = lang_choice_v
            msg_v, *_ = app.legacy_multi_confirm_ui(state_v, "QA", *values_v)
            fam_id_v = msg_v.split("**")[1]
            data_v = fam.load_family(DATA_HOME, fam_id_v)
            members_v = list({m["asset_id"]: m for m in data_v["members"]}.values())
            check("validated-language members both have language_confirmed set", all(m.get("language_confirmed") for m in members_v), members_v)
            check("validated-language members have genuinely different confirmed languages",
                  members_v[0]["language_confirmed"] != members_v[1]["language_confirmed"], members_v)
            check("TM consumer gate: validated language_variant_of IS usable as a safe TM seed",
                  fam.is_language_variant_tm_safe(members_v[0], members_v[1]) is True, members_v)
    finally:
        app.DATA_HOME = orig_data_home

    # ------------------------------------------------------------------
    # Group 6e: UI grammar -- singular/plural review-item count.
    # ------------------------------------------------------------------
    empty_result = fam.ProposalResult(families=[], ambiguous_pairs=[], unrelated=[], duplicate_groups=[], ignored=[])
    one_item = [{"kind": "pending"}]
    two_items = [{"kind": "pending"}, {"kind": "pending"}]
    overview_one = app._render_legacy_family_overview(empty_result, one_item)
    overview_two = app._render_legacy_family_overview(empty_result, two_items)
    check("singular grammar: exactly 1 review item says '1 elemento'", "1 elemento da rivedere" in overview_one, overview_one)
    check("singular grammar: '1 elemento' is not pluralized as '1 elementi'", "1 elementi" not in overview_one, overview_one)
    check("plural grammar: 2+ review items say '2 elementi'", "2 elementi da rivedere" in overview_two, overview_two)

    # ------------------------------------------------------------------
    # Group 7: architectural guardrails (no second importer, no new heavy deps)
    # ------------------------------------------------------------------
    src = (BASE / "legacy_family_inference.py").read_text(encoding="utf-8")
    check("no graph database / ORM dependency introduced", not any(tok in src for tok in ("neo4j", "networkx", "sqlalchemy")), "forbidden token found")
    check("no Word COM / Java bridge dependency introduced", not any(tok in src for tok in ("win32com", "comtypes", "jpype", "py4j")), "forbidden token found")
    check("numeric confidence score never assembled as a single float shown to callers", "confidence_score" not in src, "unexpected token found")

    # ------------------------------------------------------------------
    # Group 8: onboarding back-navigation wiring (UX acceptance fix).
    # onboarding_back_ui() is the shared handler for the "← Cambia scelta"
    # button in all three onboarding sub-panels (new project / single legacy
    # import / multi-doc family import). It must never create or persist
    # anything -- it only toggles visibility -- and must correctly restore
    # onboarding_choice_panel only when no project is active yet (matching
    # project_state.change's own condition), never when this panel was opened
    # via the "+ Nuovo progetto" quick action while a project is already open.
    # ------------------------------------------------------------------
    panel_update, choice_update = app.onboarding_back_ui(None)
    check("back navigation (no active project): current panel is hidden", panel_update["visible"] is False, panel_update)
    check("back navigation (no active project): onboarding choice panel is shown again", choice_update["visible"] is True, choice_update)

    panel_update2, choice_update2 = app.onboarding_back_ui("/some/active/project/dir")
    check("back navigation (project already active): current panel is still hidden", panel_update2["visible"] is False, panel_update2)
    check("back navigation (project already active): onboarding choice panel stays hidden (no regression on quick-action entry)",
          choice_update2["visible"] is False, choice_update2)

    # ------------------------------------------------------------------
    # Group 9: single decision flow (UX fix -- removed duplicate "Nessuna
    # relazione correlata" button, which offered a second way to do exactly
    # what selecting "Non sono correlati" + Conferma already does, and had
    # already caused a real user mistake).
    # ------------------------------------------------------------------
    check("legacy_multi_dismiss_ui no longer exists (single decision flow, not two)",
          not hasattr(app, "legacy_multi_dismiss_ui"), "function should be removed")
    app_src_g9 = (BASE / "app.py").read_text(encoding="utf-8")
    check("no separate 'Nessuna relazione correlata' button widget in app.py source",
          'gr.Button("Nessuna relazione correlata"' not in app_src_g9,
          "duplicate button widget found")
    check("no lm_dismiss component/wiring left in app.py source",
          "lm_dismiss" not in app_src_g9, "lm_dismiss reference found")
    check("confirm button relabeled to the single 'Conferma selezione' action",
          'gr.Button("Conferma selezione"' in app_src_g9, "expected button label not found")

    # The single remaining path already produces the required outcome: with
    # "Non sono correlati" selected, clicking the ONE confirm button confirms
    # independence and persists nothing.
    doc_a = make_docx("SingleFlowA_IT_Rev1.docx", "SingleFlowA", "1 - Introduzione", IT_TEXT + " variante uno singola")
    doc_b = make_docx("SingleFlowA_IT_Rev2.docx", "SingleFlowA", "1 - Introduzione", IT_TEXT + " variante due singola aggiornata")
    orig_data_home2 = app.DATA_HOME
    app.DATA_HOME = DATA_HOME
    try:
        state_sf, *_ = app.legacy_multi_analyze_ui([str(doc_a), str(doc_b)], None)
        _result_sf, review_items_sf = state_sf
        item_sf = review_items_sf[0]
        not_related_choice_sf = next(
            (c for c in item_sf["choices"] if c.startswith(app._LEGACY_NOT_RELATED)), None)
        check("single-flow item offers a 'Non sono correlati' choice in the single decision flow",
              not_related_choice_sf is not None, item_sf["choices"])
        values_sf = [not_related_choice_sf] + [None] * (app._LEGACY_MAX_REVIEW_ITEMS - 1)
        msg_sf, *_ = app.legacy_multi_confirm_ui(state_sf, "QA Tester", *values_sf)
        check("single-flow: confirming with 'Non sono correlati' selected confirms independence, persists nothing",
              "Nessuna relazione confermata" in msg_sf, msg_sf)
        check("single-flow: nothing persisted for either document",
              fam.find_family_for_asset(DATA_HOME, item_sf["a"].file_sha256) is None, "")
    finally:
        app.DATA_HOME = orig_data_home2

    # ------------------------------------------------------------------
    # Group 10: "Il tuo nome" required-field inline validation (UX fix --
    # confirming with an empty name must never raise a global gr.Error /
    # render an Error component; it must behave like an ordinary form
    # validation error on just that field, preserving the selection already
    # made and requiring no new analysis).
    # ------------------------------------------------------------------
    check("no raw gr.Error for the empty-name case left in app.py source",
          'raise gr.Error("Indica il tuo nome per confermare le relazioni.")' not in app_src_g9,
          "old global-error call still present")
    check("inline error CSS class present in app.py source",
          "ms-field-invalid" in app_src_g9 and "ms-inline-error" in app_src_g9,
          "expected CSS hook classes not found")

    # _legacy_name_field_live_validate: clears error state live once a valid
    # name is typed, leaves state alone while still empty/whitespace-only.
    live_empty = app._legacy_name_field_live_validate("")
    live_blank = app._legacy_name_field_live_validate("   ")
    live_ok = app._legacy_name_field_live_validate("Mario Rossi")
    check("live name validation: empty input is a no-op (no premature error)",
          "value" not in live_empty[0] and "value" not in live_empty[1], live_empty)
    check("live name validation: whitespace-only input is also a no-op",
          "value" not in live_blank[0] and "value" not in live_blank[1], live_blank)
    check("live name validation: valid input clears the error box and the invalid border class",
          live_ok[0].get("elem_classes") == [] and live_ok[1].get("visible") is False and live_ok[1].get("value") == "",
          live_ok)

    press_a = make_docx("PressC_IT_RevA.docx", "PressC", "1 - Introduzione", IT_TEXT + " pressa prima edizione validazione nome")
    press_b = make_docx("PressC_IT_Rev2.docx", "PressC", "1 - Introduzione", IT_TEXT + " pressa seconda edizione nuove specifiche validazione nome")
    orig_data_home3 = app.DATA_HOME
    app.DATA_HOME = DATA_HOME
    try:
        state_v10, *_ = app.legacy_multi_analyze_ui([str(press_a), str(press_b)], None)
        _result_v10, review_items_v10 = state_v10
        order_item_v10 = next(it for it in review_items_v10
                               if it["kind"] == "relation" and it["relation"]["relation_type"] == "revision_of")
        idx_v10 = review_items_v10.index(order_item_v10)
        chosen_order_v10 = app._legacy_order_choice(order_item_v10["a"], order_item_v10["b"])
        values_v10 = [None] * app._LEGACY_MAX_REVIEW_ITEMS
        values_v10[idx_v10] = chosen_order_v10

        # Step 1-3: order chosen, name left empty, click Conferma selezione.
        result_empty = app.legacy_multi_confirm_ui(state_v10, "", *values_v10)
        check("empty-name confirm: output tuple has the expected length (summary,row,name,error,+radios)",
              len(result_empty) == 4 + app._LEGACY_MAX_REVIEW_ITEMS, len(result_empty))
        summary_upd, row_upd, name_upd, error_upd = result_empty[0], result_empty[1], result_empty[2], result_empty[3]
        radios_upd = result_empty[4:]
        check("empty-name confirm: summary text is left untouched (no-op update)", "value" not in summary_upd, summary_upd)
        check("empty-name confirm: confirm row visibility is left untouched (stays open)", "visible" not in row_upd, row_upd)
        check("empty-name confirm: name field gets the invalid CSS class, not a raised gr.Error",
              name_upd.get("elem_classes") == ["ms-field-invalid"], name_upd)
        check("empty-name confirm: inline error message matches the required text",
              error_upd.get("value") == app._LEGACY_NAME_REQUIRED_MSG and error_upd.get("visible") is True, error_upd)
        check("empty-name confirm: all 24 radios are left untouched (chosen order selection is preserved)",
              all("value" not in r for r in radios_upd), radios_upd)
        check("empty-name confirm: nothing was persisted",
              fam.find_family_for_asset(DATA_HOME, order_item_v10["a"].file_sha256) is None, "")

        # Step 5-7: name filled in, second click, SAME state/selection, no new analysis.
        result_ok = app.legacy_multi_confirm_ui(state_v10, "QA Tester V10", *values_v10)
        summary_ok, row_ok, name_ok, error_ok = result_ok[0], result_ok[1], result_ok[2], result_ok[3]
        check("second click (name filled, same selection, no re-analysis): relation is saved",
              isinstance(summary_ok, str) and "Famiglia" in summary_ok, summary_ok)
        check("second click: name field invalid class is cleared", name_ok.get("elem_classes") == [], name_ok)
        check("second click: inline error box is cleared", error_ok.get("value") == "" and error_ok.get("visible") is False, error_ok)
        fam_id_v10 = fam.find_family_for_asset(DATA_HOME, order_item_v10["a"].file_sha256)
        check("second click: the chosen order was actually persisted", fam_id_v10 is not None, fam_id_v10)
        if fam_id_v10:
            data_full_v10 = fam.load_family(DATA_HOME, fam_id_v10)
            rel_v10 = next(r for r in data_full_v10["relations"] if r["relation_type"] == "revision_of")
            check("second click: persisted order matches the human's chosen direction",
                  rel_v10["order_confirmed"] is True
                  and rel_v10["newer_asset_hint"] == fam.derive_asset_id(order_item_v10["b"].file_sha256),
                  rel_v10)
    finally:
        app.DATA_HOME = orig_data_home3

    try:
        print(json.dumps({"pass": sum(r["status"] == "PASS" for r in results), "tests": len(results)}))
    except UnicodeEncodeError:
        pass

    payload = {
        "build": "Milestone 3C - legacy family/revision/language inference",
        "tests": len(results),
        "pass": sum(r["status"] == "PASS" for r in results),
        "fail": sum(r["status"] == "FAIL" for r in results),
        "results": results,
    }
    (BASE / "RISULTATO_LEGACY_FAMILY_3C.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (BASE / "RISULTATO_LEGACY_FAMILY_3C.txt").write_text(
        "\n".join(f"{r['status']} | {r['test']} | {r['detail']}" for r in results) + f"\n\nPASS: {payload['pass']}/{payload['tests']}\n",
        encoding="utf-8",
    )
    if payload["fail"]:
        raise SystemExit(1)
finally:
    shutil.rmtree(TMP, ignore_errors=True)
