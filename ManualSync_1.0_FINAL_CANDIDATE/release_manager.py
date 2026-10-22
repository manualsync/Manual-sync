from __future__ import annotations

import hashlib
import json
import shutil
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

import pandas as pd

from ccms_industrial import IndustrialRepository
from compliance_intelligence import ComplianceService
from publishing_bridges import export_publication_docx
from idml_hardening import patch_idml_slots
from release_readiness import assess_release_readiness, export_readiness_report
from product_version import PRODUCT_BUILD
from project_layout import ProjectLayout
from language_codes import canonical_lang, canonical_langs
from linguistic_safety import collect_cross_language_issues, cross_language_issue

RELEASE_BUILD = "0.28.0-compliance-release-snapshot-3-5c-20260901"


def _sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for b in iter(lambda: f.read(1024 * 1024), b""):
            h.update(b)
    return h.hexdigest()


TEST_FIXTURE_MARKERS = ("TRADUZIONE DI COLLAUDO",)


def _iter_docx_visible_text(doc):
    """Yield customer-visible text from paragraphs and table cells."""
    for p in doc.paragraphs:
        yield str(p.text or "")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield str(p.text or "")


def _validate_customer_docx(path: Path, *, release_mode: str = "production") -> dict[str, Any]:
    """Fail closed on publishing regressions and accidental test fixtures.

    ``production`` is strict: ManualSync's internal test markers are forbidden.
    ``test`` keeps all structural/customer-facing quality gates active but allows
    those explicit fixture markers so an internal benchmark can still be packaged.
    """
    from docx import Document

    mode = str(release_mode or "production").strip().lower()
    if mode not in {"production", "test"}:
        raise ValueError(f"Modalità release non valida: {release_mode}")

    doc = Document(str(path))
    heading_issues: list[str] = []
    source_ref_hits: list[str] = []
    fixture_hits: list[str] = []

    for p in doc.paragraphs:
        text = str(p.text or "")
        style = p.style.name if p.style else ""
        if style.startswith("Heading"):
            if any(ch in text for ch in ("\n", "\r", "\uFFED")):
                heading_issues.append(text)
            if text != " ".join(text.split()):
                heading_issues.append(text)
        if text.startswith("Source:"):
            source_ref_hits.append(text)

    for text in _iter_docx_visible_text(doc):
        if text.startswith("Source:") and text not in source_ref_hits:
            source_ref_hits.append(text)
        for marker in TEST_FIXTURE_MARKERS:
            if marker in text:
                fixture_hits.append(text[:240])
                break

    if heading_issues:
        raise RuntimeError(f"DOCX cliente contiene heading non normalizzati: {heading_issues[:3]}")
    if source_ref_hits:
        raise RuntimeError(f"DOCX cliente contiene source refs visibili: {source_ref_hits[:3]}")
    if mode == "production" and fixture_hits:
        raise RuntimeError(
            "Release PRODUZIONE bloccata: il DOCX contiene marker di collaudo "
            f"({TEST_FIXTURE_MARKERS[0]}). Usa una traduzione reale oppure crea una release di collaudo interno."
        )

    return {
        "release_mode": mode.upper(),
        "headings_checked": sum(1 for p in doc.paragraphs if p.style and p.style.name.startswith("Heading")),
        "source_refs_visible": 0,
        "test_fixture_hits": len(fixture_hits),
    }


def _ccms_db(project_dir: str | Path) -> Path:
    p = ProjectLayout(project_dir).find_ccms_db()
    if p is None:
        raise FileNotFoundError("Repository CCMS non trovato")
    return p




def _validate_production_linguistics(repo: IndustrialRepository, publication_id: str, target_lang: str, *, profile_id: str | None = None, dimensions: dict[str, Any] | None = None) -> dict[str, Any]:
    """Independent production-only language plausibility gate.

    Human approval and XLIFF structural QA are necessary but not sufficient. A
    production bundle must also fail closed when a target in a different language
    is simply a copy (or a strongly mismatched-language copy) of the source.
    """
    pub = repo.conn.execute("SELECT source_lang FROM publications WHERE publication_id=?", (publication_id,)).fetchone()
    if not pub:
        raise KeyError(publication_id)
    source_lang = canonical_lang(str(pub["source_lang"] or ""))
    target_lang = canonical_lang(str(target_lang or ""))
    if not target_lang or target_lang == source_lang:
        return {"checked": 0, "issues": []}

    source_frame, _ = repo.assemble_structured_publication(publication_id, profile_id=profile_id, dimensions=dimensions, target_lang=None)
    target_frame, _ = repo.assemble_structured_publication(publication_id, profile_id=profile_id, dimensions=dimensions, target_lang=target_lang)
    src_by_block = {str(r.get("block_id", "")): r for _, r in source_frame.iterrows()}
    issues: list[str] = []
    checked = 0

    # Topic titles are visible customer content and need the same safety gate.
    src_titles = source_frame[["topic_id", "topic_title"]].drop_duplicates(subset=["topic_id"]) if not source_frame.empty else pd.DataFrame()
    tgt_titles = target_frame[["topic_id", "topic_title"]].drop_duplicates(subset=["topic_id"]) if not target_frame.empty else pd.DataFrame()
    tgt_title_map = {str(r["topic_id"]): str(r["topic_title"] or "") for _, r in tgt_titles.iterrows()} if not tgt_titles.empty else {}
    if not src_titles.empty:
        for _, r in src_titles.iterrows():
            tid = str(r["topic_id"]); src = str(r["topic_title"] or ""); tgt = tgt_title_map.get(tid, "")
            reason = cross_language_issue(src, tgt, source_lang, target_lang)
            checked += 1
            if reason:
                issues.append(f"topic {tid}: {reason}: {src[:120]}")

    for _, tr in target_frame.iterrows():
        bid = str(tr.get("block_id", ""))
        sr = src_by_block.get(bid)
        if sr is None:
            continue
        pairs = collect_cross_language_issues(sr.get("content"), tr.get("content"), source_lang, target_lang, limit=8)
        checked += 1
        for issue in pairs:
            issues.append(f"block {bid}: {issue}")
            if len(issues) >= 25:
                break
        if len(issues) >= 25:
            break

    return {"checked": checked, "issues": issues, "source_lang": source_lang, "target_lang": target_lang}

def _write_provenance(frame: pd.DataFrame, path: Path, language: str) -> Path:
    cols = [
        "publication_position", "topic_id", "topic_title", "topic_title_translation_status",
        "block_id", "block_type", "block_position", "object_ref_id", "translation_status",
        "source_refs", "evidence_status", "status", "version",
    ]
    available = [c for c in cols if c in frame.columns]
    out = frame[available].copy() if not frame.empty else pd.DataFrame(columns=available)
    out.insert(0, "language", language)
    out.to_csv(path, index=False, encoding="utf-8-sig")
    return path


class ReleaseIdConflictError(ValueError):
    """Raised when an immutable release identifier has already been reserved."""


def _existing_release_ids(repo: IndustrialRepository, release_ids: list[str]) -> list[str]:
    if not release_ids:
        return []
    marks = ",".join("?" for _ in release_ids)
    rows = repo.conn.execute(
        f"SELECT release_id FROM publication_releases WHERE release_id IN ({marks}) ORDER BY release_id",
        tuple(release_ids),
    ).fetchall()
    return [str(r["release_id"]) for r in rows]


def create_release_bundle(project_dir: str | Path, publication_id: str, *, target_langs: Iterable[str] | str | None,
                          released_by: str, release_id: str | None = None, profile_id: str | None = None,
                          dimensions: dict[str, Any] | None = None, idml_template: str | Path | None = None,
                          slot_map: dict[str, str] | None = None, style_map: dict[str, str] | None = None,
                          release_mode: str = "production",
                          out_dir: str | Path | None = None) -> dict[str, Any]:
    """Create one controlled release bundle.

    R17 invariant: all customer-facing artifacts are generated and validated in a
    temporary staging area *before* any immutable release row is written to the
    CCMS. Snapshot rows are then inserted in one SQLite transaction and committed
    only after every language snapshot/provenance file and the ZIP have been built.
    A business-rule block therefore leaves no new ``publication_releases`` rows.

    3.5C: Compliance readiness is read TWICE, deliberately, at two
    different points with two different jobs:
    1. Once here, via assess_release_readiness(), as a fast advisory
       pre-check -- fails fast, before any expensive DOCX/IDML generation,
       if Compliance is obviously blocked.
    2. Once more, authoritatively, at the very start of the ATOMIC CCMS
       SNAPSHOT TRANSACTION below, AFTER this function's own
       `repo.conn.execute("BEGIN IMMEDIATE")` has acquired SQLite's write
       lock (RESERVED) on this project's database. That second read is the
       one actually used both for the final gate check and for the
       compliance_snapshot embedded in every language's release row --
       closing the TOCTOU window a single early read would leave open: no
       other writer can change a Compliance decision between "the state we
       gated on" and "the state we committed", because from BEGIN IMMEDIATE
       to this transaction's commit/rollback, this connection holds SQLite's
       only write lock on the file (verified empirically: a second
       connection's write blocks with "database is locked" for the whole
       window; see regressione_compliance_release_3_5c.py's
       TOCTOU/fault-injection test).

    The ComplianceService instance itself is constructed exactly ONCE, right
    after `repo` opens below -- BEFORE any transaction -- because its
    schema-ensure step (conn.executescript()) implicitly commits any
    pending transaction; constructing a second instance (or re-running
    _ensure_schema) inside the locked window would silently release the
    lock this fix depends on. The second, authoritative Compliance read
    reuses that already-constructed instance's plain read methods (SELECT
    only, never executescript), which is what makes it safe to run inside
    the lock.

    release_publication() itself independently re-derives and gates
    Compliance whenever it is called standalone (commit=True), so a direct
    call bypassing this function cannot skip the gate either -- see its own
    docstring. commit=False is reserved exclusively for this function's own
    per-language loop; it is the only real call site in the codebase and it
    always passes an already-gated compliance_snapshot from the locked read
    above (see the RuntimeError release_publication() raises if that
    contract is ever violated).
    """
    project_dir = Path(project_dir)
    release_mode = str(release_mode or "production").strip().lower()
    if release_mode not in {"production", "test"}:
        raise ValueError("release_mode deve essere 'production' oppure 'test'.")

    readiness = assess_release_readiness(
        project_dir, publication_id, target_langs=target_langs,
        profile_id=profile_id, dimensions=dimensions,
        require_linguistic_plausibility=(release_mode == "production"),
    )
    if not readiness["release_ready"]:
        blocked = readiness["matrix"][readiness["matrix"]["status"] == "BLOCKED"]
        reasons = [f"{r.language}: {r.issues}" for r in blocked.itertuples()]
        compliance = readiness.get("compliance") or {}
        if compliance.get("compliance_applicable") and not compliance.get("compliance_ready"):
            items = compliance.get("blocking_items") or []
            preview = "; ".join(f"{b['scope_id']}/{b['requirement_id']} ({b['status']})" for b in items[:10])
            reasons.append(f"Compliance: {len(items)} requisiti bloccanti su {compliance['active_scope_count']} scope attivi — {preview}")
        raise ValueError("Release bloccata: " + "; ".join(reasons))

    release_id = release_id or f"REL-{datetime.now().strftime('%Y%m%d-%H%M%S-%f')[:-3]}"
    out = Path(out_dir) if out_dir else Path(tempfile.mkdtemp(prefix="manualsync_release_"))
    out.mkdir(parents=True, exist_ok=True)

    langs = readiness["matrix"]["language"].astype(str).tolist()
    child_release_ids = [f"{release_id}-{lang}" for lang in langs]

    repo = IndustrialRepository(_ccms_db(project_dir))
    try:
        existing = _existing_release_ids(repo, child_release_ids)
        if existing:
            raise ReleaseIdConflictError(
                "Release ID già utilizzato o riservato da un tentativo precedente: "
                f"{release_id}. Gli snapshot immutabili esistenti non vengono sovrascritti. "
                "Scegli un nuovo Release ID."
            )

        # Constructed ONCE, here, before any transaction: ComplianceService's
        # schema-ensure step would implicitly commit a pending transaction if
        # constructed again later, inside the BEGIN IMMEDIATE window below.
        compliance_svc = ComplianceService(repo)

        # Everything is staged away from the persistent Releases directory. If a
        # preflight gate fails, TemporaryDirectory removes every transient file.
        with tempfile.TemporaryDirectory(prefix=f"manualsync_{release_id}_") as td:
            stage = Path(td)
            artifacts = stage / "ARTIFACTS"
            audit = stage / "AUDIT"
            artifacts.mkdir(parents=True, exist_ok=True)
            audit.mkdir(parents=True, exist_ok=True)

            readiness_files = export_readiness_report(readiness, audit)
            current_audit_paths: list[Path] = [Path(x) for x in readiness_files.values()]
            produced: list[dict[str, Any]] = []
            releases: list[dict[str, Any]] = []
            preflight: list[dict[str, Any]] = []

            # ---------- PRE-FLIGHT: zero database mutations ----------
            for lang in langs:
                target = None if lang == readiness["source_lang"] else lang
                rid = f"{release_id}-{lang}"

                validation = repo.validate_publication_for_release(
                    publication_id, profile_id=profile_id,
                    dimensions=dimensions, target_lang=target,
                )
                if not validation["PASS"]:
                    raise ValueError("Release bloccata: " + "; ".join(validation["issues"]))

                if release_mode == "production" and target:
                    linguistic = _validate_production_linguistics(
                        repo, publication_id, target, profile_id=profile_id, dimensions=dimensions
                    )
                    if linguistic["issues"]:
                        preview = " | ".join(linguistic["issues"][:5])
                        raise ValueError(
                            f"Release PRODUZIONE bloccata: controllo linguistico {linguistic['source_lang']}→{linguistic['target_lang']} fallito "
                            f"({len(linguistic['issues'])} anomalie rilevate). {preview}"
                        )

                frame = validation["frame"].copy()
                provenance_path = audit / f"CONTENT_PROVENANCE_{lang}_{release_id}.csv"
                _write_provenance(frame, provenance_path, lang)
                current_audit_paths.append(provenance_path)

                docx = artifacts / f"{publication_id}_{lang}_{release_id}.docx"
                export_publication_docx(
                    repo, publication_id, docx,
                    profile_id=profile_id, dimensions=dimensions,
                    target_lang=target, include_source_refs=False,
                )
                quality = _validate_customer_docx(docx, release_mode=release_mode)
                produced.append({
                    "language": lang, "type": "DOCX",
                    "path": f"ARTIFACTS/{docx.name}",
                    "sha256": _sha(docx), "bytes": docx.stat().st_size,
                })

                if idml_template:
                    idml = artifacts / f"{publication_id}_{lang}_{release_id}.idml"
                    patch_idml_slots(
                        repo, publication_id, idml_template, idml,
                        profile_id=profile_id, target_lang=target,
                        slot_map=slot_map or None, style_map=style_map or None,
                    )
                    produced.append({
                        "language": lang, "type": "IDML",
                        "path": f"ARTIFACTS/{idml.name}",
                        "sha256": _sha(idml), "bytes": idml.stat().st_size,
                    })

                preflight.append({
                    "language": lang, "target": target, "release_id": rid,
                    "frame": frame, "quality": quality,
                })

            # ---------- ATOMIC CCMS SNAPSHOT TRANSACTION ----------
            # release_publication(commit=False) lets all language snapshots share
            # one transaction. Any exception rolls the complete release attempt back.
            try:
                # BEGIN IMMEDIATE acquires SQLite's RESERVED write lock right
                # now (verified empirically: a concurrent connection's write
                # blocks with "database is locked" until this transaction
                # ends) -- closing the TOCTOU window between the Compliance
                # gate and the release commit. The re-read below is the
                # authoritative one: it is what actually gates this release
                # AND what gets frozen into every language's snapshot, so a
                # concurrent Compliance edit can never land "between" the
                # gate and the snapshot -- there is no longer a gap for it to
                # land in.
                repo.conn.execute("BEGIN IMMEDIATE")
                locked_compliance = compliance_svc.compliance_readiness_for_publication(publication_id)
                if locked_compliance.get("compliance_applicable") and not locked_compliance.get("compliance_ready"):
                    items = locked_compliance.get("blocking_items") or []
                    preview = "; ".join(f"{b['scope_id']}/{b['requirement_id']} ({b['status']})" for b in items[:10])
                    raise ValueError(
                        "Release bloccata: lo stato Compliance è cambiato (o è stato osservato per la prima volta come "
                        f"non pronto) nella finestra sotto lock precedente il commit — {len(items)} requisiti bloccanti "
                        f"su {locked_compliance['active_scope_count']} scope attivi. {preview}"
                    )

                for item in preflight:
                    lang = item["language"]
                    rid = item["release_id"]
                    target = item["target"]
                    frame = item["frame"]
                    rel = repo.release_publication(
                        publication_id, release_id=rid,
                        released_by=released_by or "Release Manager",
                        profile_id=profile_id, dimensions=dimensions,
                        target_lang=target, commit=False,
                        compliance_snapshot=locked_compliance,
                    )

                    snapshot_row = repo.conn.execute(
                        "SELECT snapshot_json FROM publication_releases WHERE release_id=?", (rid,)
                    ).fetchone()
                    if not snapshot_row:
                        raise RuntimeError(f"Snapshot CCMS non trovato dopo release: {rid}")
                    snapshot_path = audit / f"CCMS_SNAPSHOT_{lang}_{release_id}.json"
                    snapshot_path.write_text(str(snapshot_row["snapshot_json"]), encoding="utf-8")
                    if _sha(snapshot_path) != str(rel["sha256"]):
                        raise RuntimeError(f"Checksum snapshot CCMS incoerente: {rid}")
                    current_audit_paths.append(snapshot_path)

                    releases.append({
                        "language": lang, **rel,
                        "snapshot_path": f"AUDIT/{snapshot_path.name}",
                        "provenance_path": f"AUDIT/CONTENT_PROVENANCE_{lang}_{release_id}.csv",
                    })

                audit_files = []
                seen_audit: set[str] = set()
                for pth in sorted(current_audit_paths, key=lambda x: x.name):
                    if not pth.is_file() or pth.name in seen_audit:
                        continue
                    seen_audit.add(pth.name)
                    audit_files.append({
                        "path": f"AUDIT/{pth.name}",
                        "sha256": _sha(pth), "bytes": pth.stat().st_size,
                    })

                manifest = {
                    "build": PRODUCT_BUILD,
                    "module_build": RELEASE_BUILD,
                    "release_id": release_id,
                    "publication_id": publication_id,
                    "released_by": released_by,
                    "release_mode": release_mode.upper(),
                    "created_at": datetime.now(timezone.utc).isoformat(),
                    "profile_id": profile_id or "",
                    "source_lang": readiness["source_lang"],
                    "languages": readiness["matrix"].to_dict(orient="records"),
                    # Convenience summary only. The canonical, hash-verified
                    # Compliance record for this release is the "compliance"
                    # section inside each language's own
                    # publication_releases.snapshot_json/snapshot_sha256
                    # (see AUDIT/CCMS_SNAPSHOT_*.json) -- never this manifest.
                    "compliance_summary": {
                        "compliance_applicable": readiness["compliance"]["compliance_applicable"],
                        "compliance_ready": readiness["compliance"]["compliance_ready"],
                        "active_scope_count": readiness["compliance"]["active_scope_count"],
                        "counts": readiness["compliance"]["counts"],
                    },
                    "ccms_releases": releases,
                    "artifacts": produced,
                    "audit_files": audit_files,
                    "preflight_quality": [
                        {"language": x["language"], **x["quality"]} for x in preflight
                    ],
                }
                manifest_path = stage / "RELEASE_MANIFEST.json"
                manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

                release_zip = stage / f"ManualSync_{publication_id}_{release_id}_RELEASE.zip"
                with zipfile.ZipFile(release_zip, "w", compression=zipfile.ZIP_DEFLATED) as zf:
                    zf.write(manifest_path, arcname="RELEASE_MANIFEST.json")
                    for item in produced:
                        pth = stage / item["path"]
                        if not pth.exists():
                            raise RuntimeError(f"Artifact dichiarato ma assente: {item['path']}")
                        zf.write(pth, arcname=item["path"])
                    for item in audit_files:
                        pth = stage / item["path"]
                        if not pth.exists():
                            raise RuntimeError(f"Audit dichiarato ma assente: {item['path']}")
                        zf.write(pth, arcname=item["path"])

                # Package exists and all quality gates passed: now persist the DB
                # transaction. Failed production/test gates never reach this point.
                repo.conn.commit()
            except Exception:
                repo.conn.rollback()
                raise

            # Promote only a completed/committed release into the persistent folder.
            (out / "ARTIFACTS").mkdir(parents=True, exist_ok=True)
            (out / "AUDIT").mkdir(parents=True, exist_ok=True)
            for item in produced:
                src = stage / item["path"]
                dst = out / item["path"]
                shutil.copy2(src, dst)
            for item in audit_files:
                src = stage / item["path"]
                dst = out / item["path"]
                shutil.copy2(src, dst)
            final_manifest = out / "RELEASE_MANIFEST.json"
            shutil.copy2(manifest_path, final_manifest)
            final_zip = out / release_zip.name
            shutil.copy2(release_zip, final_zip)

            return {
                "release_zip": final_zip,
                "manifest": final_manifest,
                "artifacts": pd.DataFrame(produced),
                "readiness": readiness,
                "release_id": release_id,
            }
    finally:
        repo.close()

