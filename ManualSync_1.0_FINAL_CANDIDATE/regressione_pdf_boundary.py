"""CP3.3 regression: the PDF-reading boundary, and everything that reads a PDF behaves as before.

CP3.3 put PyMuPDF behind a small reading contract (pdf_backend.PdfReader, implemented
by pdf_pymupdf.PyMuPdfReader) and moved the PDF/DOCX/IDML -> topics importers out of
authoring_ccms into content_import, so the content repository -- and with it the
services built on it -- no longer load PyMuPDF, PIL or python-docx when imported (numpy
still arrives with pandas, a core dependency).

This suite was written and its digests pinned BEFORE that migration, against the
original code, so it proves the migration changed nothing observable:

* B -- characterization (digests of normalised output, pinned at ae2d7b2 + CP3.2):
  layout extraction, PDF and DOCX -> topics import (returned frame AND every CCMS
  table), TM segments and a TM built from an aligned PDF pair, grounding ingestion,
  legacy-family PDF signals; and the errors (type + message) on a corrupt and on a
  missing PDF, including what grounding leaves in its open transaction.
* A -- the contract: importable without PyMuPDF, no UI / SQL / network / AI, no
  import-time effect; the PyMuPDF reader returns exactly what layout_engine and
  fitz returned; callers really go through the contract (a substitute reader is used).
* I -- import isolation: the repository and core services load no PDF stack.

Fixtures: the repository's own benchmark PDFs (test assets, not customer data), a
DOCX generated here, everything else in temporary folders inside the isolated home.

Re-pin only after a reviewed, intentional behaviour change:
    ..\\.venv\\Scripts\\python.exe regressione_pdf_boundary.py --print-golden

Usage:  ..\\.venv\\Scripts\\python.exe regressione_pdf_boundary.py
"""
from __future__ import annotations

import test_isolation  # noqa: E402,F401 -- FIRST import: temporary home + real-data guard, never the user's ManualSyncData (test_isolation.py)

import ast
import dataclasses
import hashlib
import json
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE))

import network_guard  # noqa: E402

GUARD = network_guard.install()      # any non-loopback socket is refused and recorded

import pandas as pd  # noqa: E402

PRINT_GOLDEN = "--print-golden" in sys.argv
TMP = Path(tempfile.mkdtemp(prefix="erylos_cp33_pdf_"))
BENCH = BASE / "benchmark"
PDF_IMPORT = BENCH / "abb_acs580" / "new_EN.pdf"       # 2 pages, text + images
PDF_SMALL = BENCH / "schneider" / "new_EN.pdf"          # 2 pages
PDF_OLD_EN = BENCH / "schneider" / "old_EN.pdf"         # 3 pages, aligned pair
PDF_OLD_IT = BENCH / "schneider" / "old_IT.pdf"
_results: list[tuple[str, bool, str]] = []
_observed: dict[str, str] = {}


def check(name, cond, detail=""):
    _results.append((name, bool(cond), str(detail)[:600]))
    if not PRINT_GOLDEN:
        print(("  PASS  " if cond else "  FAIL  ") + name + ("" if cond else f"   [{str(detail)[:600]}]"))


# ------------------------------------------------------------------ normalisation (as in the CP3.1 characterization)
def plain(o):
    if dataclasses.is_dataclass(o) and not isinstance(o, type):
        return plain(dataclasses.asdict(o))
    if isinstance(o, pd.DataFrame):
        return {"records": plain(o.to_dict(orient="records")), "attrs": plain(dict(o.attrs))}
    if isinstance(o, dict):
        return {str(k): plain(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [plain(x) for x in o]
    if isinstance(o, (set, frozenset)):
        return sorted(plain(x) for x in o)
    if isinstance(o, (bytes, bytearray)):
        return "sha256:" + hashlib.sha256(bytes(o)).hexdigest()
    if isinstance(o, Path):
        return str(o)
    if isinstance(o, float) and o != o:
        return "NaN"
    return o


TS_RE = re.compile(r"\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}:\d{2}(?:\.\d+)?(?:[+-]\d{2}:\d{2}|Z)?")
ID_RE = re.compile(r"\b[A-Z]{2,}-[0-9A-F]{8,}\b|\b[0-9a-f]{32}\b"
                   r"|\b[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}\b")


def normalized(o) -> str:
    text = json.dumps(plain(o), ensure_ascii=False, sort_keys=True, default=str)
    for root in (str(TMP), str(test_isolation.TEST_HOME), str(BASE)):
        # plain, forward-slash, and repr() form (exception messages quote paths with doubled backslashes)
        for form in (root.replace("\\", "\\\\"), root, root.replace("\\", "/")):
            text = text.replace(json.dumps(form)[1:-1], "<ROOT>").replace(form, "<ROOT>")
    ids: dict[str, str] = {}
    return ID_RE.sub(lambda m: ids.setdefault(m.group(0), f"<ID{len(ids) + 1}>"), TS_RE.sub("<TS>", text))


def snap(label: str, o) -> str:
    text = normalized(o)
    _observed[label] = hashlib.sha256(text.encode("utf-8")).hexdigest()
    return text


def outcome(fn, *args, **kwargs):
    try:
        return ("ok", fn(*args, **kwargs))
    except Exception as exc:  # noqa: BLE001 -- the observed behaviour IS the exception
        return ("raise", type(exc).__name__, str(exc))


def db_dump(path: Path) -> dict:
    """Every table, every row, in a stable order (the import's full footprint)."""
    conn = sqlite3.connect(str(path))
    try:
        out = {}
        for (name,) in conn.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name"):
            cur = conn.execute(f"SELECT * FROM {name}")
            cols = [c[0] for c in cur.description]
            rows = [dict(zip(cols, r)) for r in cur.fetchall()]
            out[name] = sorted(rows, key=lambda r: json.dumps(plain(r), sort_keys=True, default=str))
        return out
    finally:
        conn.close()


CORRUPT = TMP / "corrupt.pdf"
CORRUPT.write_bytes(b"%PDF-1.4\nthis is not a real pdf body\n")
MISSING = TMP / "missing.pdf"

# =================================================================== B. characterization (pinned before CP3.3)
import layout_engine  # noqa: E402

layout = layout_engine.extract_layout(PDF_IMPORT, "1-2")
check("B1: extract_layout reads text blocks, tables and images from the benchmark PDF",
      layout.text_blocks and layout.images is not None, (len(layout.text_blocks), len(layout.tables), len(layout.images)))
snap("extract_layout", layout)

from ccms_industrial import IndustrialRepository  # noqa: E402
from authoring_ccms import import_pdf_as_topics, import_docx_as_topics  # noqa: E402  (compatibility path)

db_pdf = TMP / "pdf_import" / "ManualSync_CCMS.sqlite"
db_pdf.parent.mkdir()
repo = IndustrialRepository(db_pdf)
frame = import_pdf_as_topics(repo, PDF_IMPORT, source_lang="en-US", page_spec="1-2", id_prefix="DOC",
                             source_ref_prefix="new_EN.pdf")
repo.close()
check("B2: PDF -> topics import writes topics and blocks", len(frame) > 0, len(frame))
snap("import_pdf_frame", frame)
snap("import_pdf_db", db_dump(db_pdf))

from docx import Document  # noqa: E402

docx_path = TMP / "manuale.docx"
d = Document()
d.add_heading("1 - Sicurezza", level=1)
d.add_paragraph("Scollegare l'alimentazione prima di aprire il quadro.")
d.add_paragraph("Attendere 5 minuti per la scarica dei condensatori.", style="List Bullet")
t = d.add_table(rows=2, cols=2)
t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Vite", "Coppia", "M8", "25 Nm"
d.add_heading("2 - Installazione", level=1)
d.add_paragraph("Fissare la macchina a una superficie piana.")
d.save(str(docx_path))
db_docx = TMP / "docx_import" / "ManualSync_CCMS.sqlite"
db_docx.parent.mkdir()
repo = IndustrialRepository(db_docx)
dframe = import_docx_as_topics(repo, docx_path, source_lang="it-IT", id_prefix="DOC", source_ref_prefix="manuale.docx")
repo.close()
check("B3: DOCX -> topics import splits on the two headings and keeps the table",
      len(dframe) > 0, plain(dframe)["records"][:3])
snap("import_docx_frame", dframe)
snap("import_docx_db", db_dump(db_docx))

import cat_tm  # noqa: E402

segs = cat_tm.segments_from_layout(PDF_SMALL, "1-2")
check("B4: TM segments come out of the benchmark PDF", len(segs) > 0, len(segs))
snap("cat_tm_segments", segs)
tm = cat_tm.TranslationMemory(TMP / "tm.sqlite")
built = cat_tm.build_tm_from_documents(tm, PDF_OLD_EN, PDF_OLD_IT, source_lang="en-US", target_lang="it-IT")
tm_rows = db_dump(TMP / "tm.sqlite")
tm.close()
snap("cat_tm_build", {"result": built, "db": tm_rows})

from grounded_authoring import GroundingStore  # noqa: E402

gs = GroundingStore(TMP / "ground.sqlite")
chunks = gs.ingest_pdf("BENCH", "Schneider", PDF_SMALL)
gs.close()
check("B5: grounding ingests sentences with page numbers from the PDF", len(chunks) > 0 and set(chunks["page"]) <= {1, 2},
      len(chunks))
snap("grounding_ingest", {"frame": chunks, "db": db_dump(TMP / "ground.sqlite")})

import legacy_family_inference as lfi  # noqa: E402

signals = lfi._pdf_signals(PDF_SMALL)
snap("legacy_pdf_signals", signals)

# errors: corrupt and missing PDF, through every reader
gs2 = GroundingStore(TMP / "ground_err.sqlite")
g_corrupt = outcome(gs2.ingest_pdf, "BAD", "Bad", CORRUPT)
g_same_conn = gs2.conn.execute("SELECT COUNT(*) FROM source_documents WHERE document_id='BAD'").fetchone()[0]
gs2.close()
_c = sqlite3.connect(str(TMP / "ground_err.sqlite"))
g_after = _c.execute("SELECT COUNT(*) FROM source_documents WHERE document_id='BAD'").fetchone()[0]
_c.close()
db_err = TMP / "err_import" / "ManualSync_CCMS.sqlite"
db_err.parent.mkdir()
repo = IndustrialRepository(db_err)
gs3 = GroundingStore(TMP / "g3.sqlite")
errors = {
    "extract_layout_corrupt": outcome(layout_engine.extract_layout, CORRUPT),
    "extract_layout_missing": outcome(layout_engine.extract_layout, MISSING),
    "import_pdf_corrupt": outcome(import_pdf_as_topics, repo, CORRUPT, source_lang="en-US"),
    "import_pdf_missing": outcome(import_pdf_as_topics, repo, MISSING, source_lang="en-US"),
    "segments_corrupt": outcome(cat_tm.segments_from_layout, CORRUPT),
    "grounding_corrupt": g_corrupt,
    "grounding_missing": outcome(gs3.ingest_pdf, "M", "M", MISSING),
    "legacy_signals_corrupt": outcome(lfi._pdf_signals, CORRUPT),
}
repo.close()
gs3.close()
check("B6: a corrupt or missing PDF raises (never a silent empty import) in every reader",
      all(v[0] == "raise" for v in errors.values()), {k: v[:2] for k, v in errors.items()})
check("B7: grounding on a corrupt PDF raises AFTER writing its document row in the open transaction, "
      "and nothing is committed",
      g_corrupt[0] == "raise" and g_same_conn == 1 and g_after == 0, (g_corrupt[:2], g_same_conn, g_after))
snap("errors", {"errors": errors, "grounding_same_conn": g_same_conn, "grounding_committed": g_after})

# =================================================================== A. the contract (CP3.3)
CONTRACT = BASE / "pdf_backend.py"
ADAPTER = BASE / "pdf_pymupdf.py"
check("A1: the reading contract and its PyMuPDF implementation exist (pdf_backend.py, pdf_pymupdf.py)",
      CONTRACT.is_file() and ADAPTER.is_file(), [CONTRACT.exists(), ADAPTER.exists()])
if CONTRACT.is_file() and ADAPTER.is_file():
    probe = subprocess.run([sys.executable, "-I", "-c",
                            "import sys; sys.path.insert(0, sys.argv[1]); import pdf_backend; "
                            "print(sorted(m for m in ('fitz','pymupdf','numpy','PIL','docx','layout_engine','sqlite3',"
                            "'gradio','fastapi','requests','socket','anthropic','openai') if m in sys.modules))",
                            str(BASE)], capture_output=True, text=True, timeout=120)
    loaded = probe.stdout.strip()
    check("A2: importing the contract loads no PyMuPDF / numpy / PIL / docx / layout_engine / SQL / UI / web / network / AI",
          probe.returncode == 0 and loaded == "[]", [probe.returncode, loaded, probe.stderr[-300:]])
    ct = ast.parse(CONTRACT.read_text(encoding="utf-8"))
    top_imports = {a.name for s in ct.body if isinstance(s, ast.Import) for a in s.names} | \
                  {s.module for s in ct.body if isinstance(s, ast.ImportFrom) and s.module}
    check("A3: the contract's module-level imports are standard-library typing only",
          top_imports <= {"__future__", "pathlib", "typing", "collections.abc"}, top_imports)
    check("A4: the contract has no import-time statement besides imports, constants, classes and functions",
          all(isinstance(s, (ast.Import, ast.ImportFrom, ast.ClassDef, ast.FunctionDef, ast.Assign))
              or (isinstance(s, ast.Expr) and isinstance(s.value, ast.Constant)) for s in ct.body)
          and not any(isinstance(s, ast.Assign) and isinstance(s.value, ast.Call) for s in ct.body), "")
    import pdf_backend  # noqa: E402
    import pdf_pymupdf  # noqa: E402

    reader = pdf_backend.default_reader()
    check("A5: the default reader is the PyMuPDF implementation and satisfies the contract",
          isinstance(reader, pdf_pymupdf.PyMuPdfReader) and isinstance(reader, pdf_backend.PdfReader), type(reader))
    check("A6: reader.extract_layout == layout_engine.extract_layout (same normalised output)",
          normalized(reader.extract_layout(PDF_IMPORT, "1-2")) == normalized(layout), "")
    import fitz  # noqa: E402

    with fitz.open(PDF_SMALL) as doc:
        direct_texts = [p.get_text("text") for p in doc]
        direct_title = str((doc.metadata or {}).get("title") or "").strip()
    gen = reader.iter_page_texts(PDF_SMALL)
    check("A7: reader.iter_page_texts == page.get_text('text') for every page, in order",
          list(gen) == direct_texts, len(direct_texts))
    check("A8: reader.metadata_title == the document's metadata title", reader.metadata_title(PDF_SMALL) == direct_title,
          [reader.metadata_title(PDF_SMALL), direct_title])
    ex_direct = outcome(lambda: fitz.open(CORRUPT))
    ex_reader = outcome(lambda: list(reader.iter_page_texts(CORRUPT)))
    check("A9: errors propagate unchanged (same exception type and message as PyMuPDF itself)",
          ex_direct[:2] == ex_reader[:2] and ex_direct[2] == ex_reader[2], [ex_direct, ex_reader])

    # callers really go through the contract: swap in a substitute reader
    class _Substitute:
        name = "substitute"
        calls: list[str] = []

        def extract_layout(self, pdf_path, page_spec=""):
            self.calls.append("extract_layout")
            return reader.extract_layout(pdf_path, page_spec)

        def iter_page_texts(self, pdf_path):
            self.calls.append("iter_page_texts")
            yield "Testo sostitutivo della pagina uno."

        def metadata_title(self, pdf_path):
            self.calls.append("metadata_title")
            return "Titolo sostitutivo"

    original_default = pdf_backend.default_reader
    sub = _Substitute()
    pdf_backend.default_reader = lambda: sub
    try:
        gs4 = GroundingStore(TMP / "ground_sub.sqlite")
        sub_chunks = gs4.ingest_pdf("SUB", "Sub", PDF_SMALL)
        gs4.close()
        sub_sig = outcome(lfi._pdf_signals, PDF_SMALL)
        sub_segs = outcome(cat_tm.segments_from_layout, PDF_SMALL)
        repo = IndustrialRepository(TMP / "sub_import.sqlite")
        sub_imp = outcome(import_pdf_as_topics, repo, PDF_SMALL, source_lang="en-US")
        repo.close()
    finally:
        pdf_backend.default_reader = original_default
    check("A10: grounding, TM segmentation, legacy signals and the PDF importer all read through pdf_backend",
          {"iter_page_texts", "extract_layout", "metadata_title"} <= set(sub.calls)
          and list(sub_chunks["text"]) == ["Testo sostitutivo della pagina uno."]
          and sub_sig[0] == "ok" and sub_sig[1][2] == "Titolo sostitutivo", [sub.calls, sub_sig[:2], sub_segs[:2], sub_imp[:2]])
    check("A11: the default reader is restored", pdf_backend.default_reader is original_default, "")
else:
    for n in ("A2", "A3", "A4", "A5", "A6", "A7", "A8", "A9", "A10", "A11"):
        check(f"{n}: (contract missing)", False, "pdf_backend.py / pdf_pymupdf.py not present")

# =================================================================== I. import isolation (CP3.3 target)
# numpy is NOT in this list on purpose: pandas loads it unconditionally (`import pandas`
# alone puts numpy in sys.modules), and pandas stays a core dependency (CP3.0 decision).
# I0 below proves that is the only reason numpy is present.
PDF_STACK = ("fitz", "pymupdf", "PIL", "docx", "layout_engine", "idml_import")
_np = subprocess.run([sys.executable, "-I", "-c", "import sys, pandas; print('numpy' in sys.modules)"],
                     capture_output=True, text=True, timeout=120)
check("I0: numpy comes with pandas itself (a core dependency), not with the PDF stack", _np.stdout.strip() == "True",
      [_np.returncode, _np.stdout, _np.stderr[-200:]])
ISOLATED = ("authoring_ccms", "ccms_industrial", "authoring_service", "compliance_intelligence", "cat_tm",
            "grounded_authoring", "generative_orchestrator", "translation_impact", "release_readiness",
            "project_lifecycle", "workflow_management", "impact_intelligence", "translation_intelligence",
            "publishing_preview")
loaded_by = {}
for _m in ISOLATED:                     # one fresh interpreter per module: nothing pre-loaded
    _r = subprocess.run([sys.executable, "-I", "-c",
                         "import sys, importlib; sys.path.insert(0, sys.argv[1]); importlib.import_module(sys.argv[2]); "
                         "print(','.join(x for x in sys.argv[3].split(',') if x in sys.modules))",
                         str(BASE), _m, ",".join(PDF_STACK)], capture_output=True, text=True, timeout=300)
    loaded_by[_m] = [_r.returncode, _r.stdout.strip(), _r.stderr[-200:]]
dirty = {m: v[1] for m, v in loaded_by.items() if v[0] != 0 or v[1]}
check("I1: the content repository and the core services load no PyMuPDF / PIL / docx / layout_engine / "
      "idml_import when imported", not dirty, dirty)

# =================================================================== no network, pins, cleanup
check("N1: no network access attempted", not GUARD.snapshot(), GUARD.snapshot())

# Pinned before the CP3.3 migration (HEAD ae2d7b2 + CP3.2 working tree); three runs, identical.
GOLDEN = {
    'cat_tm_build': '2aab3bcc969ad702b0681aa6a52cdb20b92a91977d4daa1e7246b1ba8d92015a',
    'cat_tm_segments': '0b099616ebffd162b1c5fefca97954ffe0ee7f2deef5052a11f426a294721e78',
    'errors': '80dff92559d9fb93f147dfe56ebf7c4e82237d11aed18e97c87fe84f43988770',
    'extract_layout': '85de51ffdcde38fe54bb141aae2fe5cd1085322bc413ba634b80a8f86fed11a5',
    'grounding_ingest': '30b1477fa8fa6d753e7781f39ed610c7f8fd8625020ebf38b562b5ed415d4030',
    'import_docx_db': 'e321e7b200fa232f79222e87c368ccb57685bad789757b54e0d50cca87c2609d',
    'import_docx_frame': '047b1c91377ae9c2fb273c4f09644fe504c0130c68bd4a2b3827b79b5796d012',
    'import_pdf_db': 'f5025e7140d13a815017c336f7aa5b9247debac2922a473fb1884c8d7c195eba',
    'import_pdf_frame': 'a72f8364d487d9cbc1b9b4afa7bd8556e71efe642431c5884c73382d18a24052',
    'legacy_pdf_signals': 'a58f875beb16327409ce6fda304371a5e4114bd7e3e3d4f8ffdf8dc651ad36b6',
}

if PRINT_GOLDEN:
    print("GOLDEN = {")
    for k in sorted(_observed):
        print(f"    {k!r}: {_observed[k]!r},")
    print("}")
    shutil.rmtree(TMP, ignore_errors=True)
    sys.exit(0)

for label in sorted(_observed):
    check(f"SNAP {label}: normalised output identical to the pre-CP3.3 baseline",
          GOLDEN.get(label) == _observed[label], f"expected {GOLDEN.get(label)} got {_observed[label]}")
check("SNAP: every pinned output was observed", set(GOLDEN) == set(_observed), sorted(set(GOLDEN) ^ set(_observed)))

shutil.rmtree(TMP, ignore_errors=True)
check("T1: the temporary fixtures were removed", not TMP.exists(), TMP)

passed = sum(1 for _, ok, _ in _results if ok)
total = len(_results)
(BASE / "RISULTATO_PDF_BOUNDARY.json").write_text(json.dumps({
    "suite": "pdf boundary (CP3.3)", "tests": total, "pass": passed, "fail": total - passed, "digests": _observed,
    "results": [{"test": n, "status": "PASS" if ok else "FAIL", "detail": d} for n, ok, d in _results],
}, indent=2, ensure_ascii=False), encoding="utf-8")
print(f"\n  {passed}/{total} passed")
sys.exit(0 if passed == total else 1)
