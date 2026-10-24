"""CP5.2a regression: CHARACTERIZATION of the import orchestration in app.import_manual_ui.

import_manual_ui (app.py) is a UI callback that orchestrates the migration import: it
refuses a re-import once the publication holds reviewed content, RETIRES the previous
migration draft with DELETEs on publication_topics, topics, structured_tables and
figures (content_blocks, translations and versions go with them through ON DELETE
CASCADE), updates `publications`, copies the uploaded file into Sources/ and writes the
IDML review state. Until CP5.2b it issued that SQL itself; CP5.2b-1 moved the 11 SQL
calls, verbatim, into IndustrialRepository (migration_draft_objects,
retire_migration_draft, publication_exists, update_publication_header,
structured_table_count) and left the orchestration -- order, the two commits, every
exception -- where it was. This suite pinned the behaviour at HEAD d4f7c45, BEFORE
the move, quirks included; every observation below stayed identical across it:

* S01-S03   new import: PDF, DOCX, IDML (table, figure, orphan story -> review item);
* S04       PDF re-import over an IDML draft that carries translations: since HARDENING-02B
            REFUSED (gr.Error naming the translations), nothing touched -- what the retirement
            deletes when nothing protects the draft is pinned by S04x;
* S05/S06   DOCX over DOCX (publication title/language UPDATE), IDML over IDML
            (the non-destructive path: retired only after the new parse succeeded);
* S07       refusal: a topic no longer 'Migrated' (Approved, and also Draft) blocks the
            re-import -- since HARDENING-02B a gr.Error with the same text (a ValueError never
            reaches the browser: show_error=False), and nothing at all is copied (DLS-03);
* S08-S11   failures: corrupt PDF, empty PDF, corrupt DOCX run AFTER the old draft was
            retired (the publication is left without topics -- pinned, not fixed); a
            corrupt IDML leaves the old draft byte-for-byte untouched;
* S12       .doc through the converter seam (success and failure);
* S13/S14   the gr.Error guard and the argument defaults (id, language, title);
* S15       the guided onboarding (_guided_create_and_import) that reuses the same path;
* S16       (CP5.2b-0) the publication title/language asymmetry: the FIRST import goes through
            create_publication (title stripped, language canonicalized: "it-it" -> "it-IT"),
            a RE-import through the inline UPDATE (both stored as received) -- pinned so that
            moving that UPDATE cannot "unify" the two paths silently;
* R1-R4     (CP5.2b-1) the structure after the move: import_manual_ui issues no SQL, calls the
            five IndustrialRepository methods, the draft is still retired at the same two points,
            and those methods hold exactly the 11 original statements, in order, with the two
            commits (frozen copy below); a SQL call put back into import_manual_ui is caught;
* S04x      (HARDENING-02A) S04 without the translations: the retirement with no human work at all;
* S17a-k    (HARDENING-02A, B5 characterization) one kind of human work on a migration draft, then a
            re-import: a Studio block / title edit (PDF, DOCX), a block created in Studio, an open review
            (PDF vs IDML), an archived topic, an autosave with nothing changed, IDML review decisions
            (database + JSON, JSON only), the guided onboarding with the same / an empty name (DLS-02),
            a same-name file refused after approval (DLS-03), an audit row older than the draft (as the
            re-import leaves it, and seeded directly), an orphan autosave, compliance evidence. Pinned
            at 9ae4989 as they behaved then -- destructive. HARDENING-02B CHANGED THEM ON PURPOSE (B5):
            every draft that carries human work is now refused before any effect -- gr.Error naming
            what was found, database, files and Sources/ untouched -- while an autosave, rows older
            than the draft and a draft with no human work still let the re-import through;
* S18/S19   (HARDENING-02B) each detection kind on its own, its time filter, the decisions D1-D5
            (review/evidence filtered, evidence also by its link, a corrupt review JSON does not count,
            JSON decisions count for every format, an unreadable date counts) and a refusal that must
            not even create Sources/;
* R5/R6     (HARDENING-02B) the order guard -> Sources/ -> copy -> retire -> parse in the source of
            import_manual_ui; the two new IndustrialRepository methods only read; the five CP5.2b-1
            methods frozen byte-for-byte (SHA-256 of their source at 9ae4989).

For each scenario the WHOLE content database (every table, every row, in rowid
order), the project files (Sources/, review JSON, manifest) and the returned
message/DataFrames are captured before and after. Only UTC timestamps, the
temporary folder and what is drawn from uuid4 (the IDML caption-candidate key, the
TOPIC-/BLK-/TBL- ids of what a user creates in Studio -- each replaced consistently,
so the links stay pinned) are normalized; the ids of imported content, statuses,
texts, relations and counts are not. Golden digests were pinned after three
identical runs; section M proves they are sensitive. Nothing is written outside a
temporary folder.

Re-pin only after a reviewed, intentional behaviour change:
    ..\\.venv\\Scripts\\python.exe regressione_import_orchestration.py --print-golden

Usage:  ..\\.venv\\Scripts\\python.exe regressione_import_orchestration.py
"""
from __future__ import annotations

import test_isolation  # noqa: E402,F401 -- FIRST import: temporary home + real-data guard, never the user's ManualSyncData (test_isolation.py)

import contextlib
import copy
import hashlib
import json
import re
import shutil
import sqlite3
import sys
import tempfile
import zipfile
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE))

import network_guard  # noqa: E402

GUARD = network_guard.install()

import fitz  # noqa: E402
from docx import Document  # noqa: E402

import app  # noqa: E402
from ccms_industrial import IndustrialRepository  # noqa: E402
from legacy_doc_conversion import LegacyDocConversionError  # noqa: E402
from project_lifecycle import create_project  # noqa: E402

PRINT_GOLDEN = "--print-golden" in sys.argv
TMP = Path(tempfile.mkdtemp(prefix="erylos_cp52a_import_"))
INPUTS = TMP / "inputs"
INPUTS.mkdir()
_results: list[tuple[str, bool, str]] = []
_observed: dict[str, str] = {}


def check(name, cond, detail=""):
    _results.append((name, bool(cond), str(detail)[:700]))
    if not PRINT_GOLDEN:
        print(("  PASS  " if cond else "  FAIL  ") + name + ("" if cond else f"   [{str(detail)[:700]}]"))


# ------------------------------------------------------------ normalization: timestamps, the temp folder, random keys
_TS = re.compile(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:\.\d+)?(?:\+00:00|Z)?")
# idml_import draws the caption-candidate key from uuid4 (idml_import.py, candidate_key = uuid4().hex[:12]).
# It is written twice -- in the figure's source_refs and as the review item's item_key -- so it is
# replaced CONSISTENTLY (same key -> same placeholder): the figure <-> review-item link stays pinned.
_CAPTION_KEY = re.compile(r'(idml_caption_candidate_key=|\["item_key", ")([0-9a-f]{12})')   # plain(): dicts are [k, v] pairs
# HARDENING-02A: what a user creates in Studio gets its id from uuid4 too (authoring_service.py: TOPIC-/BLK-/TBL- +
# uuid4().hex[:10].upper()). Same consistent replacement, so an audit row that names a deleted block stays pinned to it.
# The ids of imported content (DOC-0001-B001, IDML-0001-2-F002, ...) never match and are never normalized.
_STUDIO_ID = re.compile(r"\b(TOPIC|BLK|TBL)-([0-9A-F]{10})\b")


def plain(o):
    if hasattr(o, "to_dict") and hasattr(o, "columns"):          # pandas DataFrame
        return {"columns": [str(c) for c in o.columns], "rows": plain(o.to_dict(orient="records"))}
    if isinstance(o, dict):
        return [[plain(k), plain(v)] for k, v in o.items()]
    if isinstance(o, (list, tuple)):
        return [plain(x) for x in o]
    if isinstance(o, float):
        return "NaN" if o != o else repr(o)
    if isinstance(o, Path):
        return str(o)
    return o


def normalized(o) -> str:
    text = json.dumps(plain(o), ensure_ascii=False, default=repr)
    for root in (str(TMP), str(test_isolation.TEST_HOME)):
        for form in (root.replace("\\", "\\\\"), root, root.replace("\\", "/")):
            text = text.replace(json.dumps(form)[1:-1], "<TMP>").replace(form, "<TMP>")
    keys: dict[str, str] = {}
    text = _CAPTION_KEY.sub(lambda m: m.group(1) + keys.setdefault(m.group(2), f"<CKEY{len(keys) + 1}>"), text)
    uids: dict[str, str] = {}
    text = _STUDIO_ID.sub(lambda m: m.group(1) + "-" + uids.setdefault(m.group(2), f"<UID{len(uids) + 1}>"), text)
    return _TS.sub("<TS>", text)


def digest(o) -> str:
    return hashlib.sha256(normalized(o).encode("utf-8")).hexdigest()


def snap(label, o):
    _observed[label] = digest(o)


# ------------------------------------------------------------ synthetic inputs (deterministic content)
def make_pdf(name: str, sections) -> Path:
    path = INPUTS / name
    doc = fitz.open()
    for title, paras in sections:
        page = doc.new_page(width=595, height=842)
        if title:
            page.insert_text((50, 80), title, fontsize=18, fontname="hebo")
        for i, p in enumerate(paras):
            page.insert_text((50, 120 + 20 * i), p, fontsize=10, fontname="helv")
    doc.save(str(path), garbage=4, deflate=True, no_new_id=True)
    doc.close()
    return path


def make_docx(name: str, chapters) -> Path:
    path = INPUTS / name
    d = Document()
    for title, para, table in chapters:
        d.add_heading(title, level=1)
        d.add_paragraph(para)
        if table:
            t = d.add_table(rows=len(table), cols=len(table[0]))
            for r, row in enumerate(table):
                for c, val in enumerate(row):
                    t.cell(r, c).text = val
    d.save(str(path))
    return path


_MIMETYPE = b"application/vnd.adobe.indesign-idml-package"


def _para(style, text, extra=""):
    return (f'<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/{style}"><CharacterStyleRange>{extra}'
            f'<Content>{text}</Content></CharacterStyleRange></ParagraphStyleRange>')


def make_idml(name: str, heading: str, body: str, *, table, orphan: bool) -> Path:
    path = INPUTS / name
    figure = ('<Rectangle Self="r1" ItemTransform="1 0 0 1 5 5"><Link LinkResourceURI="file:///C:/nonexistent/pump.png" '
              'StoredState="Normal"/><Image/></Rectangle>')
    cells = "".join(f'<Cell Name="{r}:{c}">{_para("Body Text", v)}</Cell>' for r, row in enumerate(table) for c, v in enumerate(row))
    story = (f'<?xml version="1.0" encoding="UTF-8"?><Story Self="s1">{_para("Heading 1", heading)}{_para("Body Text", body)}'
             f'{_para("Body Text", "", figure)}{_para("Caption", "Didascalia pompa idraulica")}'
             f'<Table Self="tbl_s1">{cells}</Table></Story>')
    stories = ["Stories/Story_s1.xml"] + (["Stories/Story_o1.xml"] if orphan else [])
    designmap = ('<?xml version="1.0" encoding="UTF-8"?><Document xmlns:idPkg="http://ns.adobe.com/AdobeInDesign/idml/1.0/packaging">'
                 '<idPkg:Spread src="Spreads/Spread_a.xml"/>' + "".join(f'<idPkg:Story src="{s}"/>' for s in stories) + '</Document>')
    parts = {"designmap.xml": designmap,
             "Spreads/Spread_a.xml": ('<?xml version="1.0" encoding="UTF-8"?><Spread Self="a">'
                                      '<TextFrame Self="tf1" ParentStory="s1" ItemTransform="1 0 0 1 0 0"/></Spread>'),
             "Stories/Story_s1.xml": story}
    if orphan:
        parts["Stories/Story_o1.xml"] = ('<?xml version="1.0" encoding="UTF-8"?><Story Self="o1">'
                                         f'{_para("Body Text", "Testo mai collegato a una pagina.")}</Story>')
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr(zipfile.ZipInfo("mimetype", (2020, 1, 1, 0, 0, 0)), _MIMETYPE, compress_type=zipfile.ZIP_STORED)
        for part, content in parts.items():
            zf.writestr(zipfile.ZipInfo(part, (2020, 1, 1, 0, 0, 0)), content.encode("utf-8"))
    return path


PDF_A = make_pdf("Manuale_A.pdf", [("1 Introduzione", ["Il presente manuale descrive il sistema idraulico.",
                                                         "Leggere con attenzione prima dell'uso."]),
                                     ("2 Manutenzione", ["Scollegare l'alimentazione prima di ogni intervento.",
                                                         "Pressione massima ammessa 12 bar."])])
PDF_B = make_pdf("Manuale_B.pdf", [("1 Introduzione", ["Revisione B del manuale."]),
                                     ("2 Uso previsto", ["Uso esclusivamente industriale."]),
                                     ("3 Smaltimento", ["Smaltire secondo le norme locali."])])
PDF_EMPTY = make_pdf("Vuoto.pdf", [("", [])])
PDF_CORRUPT = INPUTS / "Rotto.pdf"
PDF_CORRUPT.write_bytes(b"%PDF-1.4 questo non e' un PDF")
DOCX_A = make_docx("Manuale_A.docx", [("1 Introduzione", "Testo introduttivo del capitolo uno.",
                                        [["Parametro", "Valore"], ["Tensione", "24 V"]]),
                                       ("2 Manutenzione", "Scollegare l'alimentazione.", None)])
DOCX_B = make_docx("Manuale_B.docx", [("1 Panoramica", "Testo della revisione B.",
                                        [["Grandezza", "Limite"], ["Pressione", "16 bar"], ["Temperatura", "60 C"]])])
DOCX_CORRUPT = INPUTS / "Rotto.docx"
DOCX_CORRUPT.write_bytes(b"questo non e' un docx")
DOC_LEGACY = INPUTS / "Legacy.doc"
DOC_LEGACY.write_bytes(b"segnaposto .doc (convertitore simulato)")
IDML_A = make_idml("Manuale_A.idml", "1 Introduzione", "Il presente manuale descrive il sistema idraulico.",
                   table=[["Parametro", "Valore"], ["Pressione", "10 bar"]], orphan=True)
IDML_B = make_idml("Manuale_B.idml", "1 Panoramica", "Revisione B del pacchetto InDesign.",
                   table=[["Grandezza", "Limite"], ["Portata", "40 l/min"]], orphan=False)
IDML_CORRUPT = INPUTS / "Rotto.idml"
IDML_CORRUPT.write_bytes(b"PK questo non e' un IDML")
INPUT_SHA = {hashlib.sha256(p.read_bytes()).hexdigest(): p.name for p in INPUTS.iterdir()}


# ------------------------------------------------------------ observation: database, files, orphans
def new_project(name: str) -> Path:
    return create_project(TMP / "projects", project_name=name, source_lang="it-IT")


def db_rows(project: Path) -> dict:
    con = sqlite3.connect(project / "ManualSync_CCMS.sqlite")
    con.row_factory = sqlite3.Row
    out = {}
    try:
        for (name,) in con.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name").fetchall():
            out[name] = [dict(r) for r in con.execute(f"SELECT * FROM {name} ORDER BY rowid")]
    finally:
        con.close()
    return out


ORPHAN_LINKS = [  # (child table, child column, parent table, parent column)
    ("content_blocks", "topic_id", "topics", "topic_id"),
    ("publication_topics", "topic_id", "topics", "topic_id"),
    ("publication_topics", "publication_id", "publications", "publication_id"),
    ("topic_translations", "topic_id", "topics", "topic_id"),
    ("block_translations", "block_id", "content_blocks", "block_id"),
    ("structured_table_translations", "table_id", "structured_tables", "table_id"),
    ("figure_translations", "figure_id", "figures", "figure_id"),
    ("topic_versions", "topic_id", "topics", "topic_id"),
    ("block_versions", "block_id", "content_blocks", "block_id"),
    ("structured_table_versions", "table_id", "structured_tables", "table_id"),
    ("figure_versions", "figure_id", "figures", "figure_id"),
]


def orphans(project: Path) -> dict:
    """Rows whose parent is gone, plus tables/figures no block points at."""
    con = sqlite3.connect(project / "ManualSync_CCMS.sqlite")
    out = {}
    try:
        tables = {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        for child, col, parent, pcol in ORPHAN_LINKS:
            if child in tables and parent in tables:
                out[f"{child}.{col}->{parent}"] = [r[0] for r in con.execute(
                    f"SELECT {col} FROM {child} WHERE {col} NOT IN (SELECT {pcol} FROM {parent}) ORDER BY rowid")]
        for kind, table, idcol in (("table", "structured_tables", "table_id"), ("figure", "figures", "figure_id")):
            out[f"{table} not referenced by a block"] = [r[0] for r in con.execute(
                f"SELECT {idcol} FROM {table} WHERE {idcol} NOT IN "
                f"(SELECT object_ref_id FROM content_blocks WHERE block_type=?) ORDER BY rowid", (kind,))]
    finally:
        con.close()
    return out


def files(project: Path) -> list:
    """Every project file: databases by name, uploaded copies by WHICH input they are, JSON by content."""
    out = []
    for p in sorted(project.rglob("*")):
        rel = p.relative_to(project).as_posix()
        if p.is_dir():
            out.append([rel + "/", "dir"])
        elif p.suffix == ".sqlite":
            out.append([rel, "sqlite"])
        elif p.suffix == ".json":
            try:
                out.append([rel, json.loads(p.read_text(encoding="utf-8"))])
            except ValueError:                      # HARDENING-02B S18j/S18k: a corrupt review JSON, pinned by content
                out.append([rel, f"invalid JSON sha256:{hashlib.sha256(p.read_bytes()).hexdigest()}"])
        else:
            sha = hashlib.sha256(p.read_bytes()).hexdigest()
            out.append([rel, f"copy of {INPUT_SHA[sha]}" if sha in INPUT_SHA else f"sha256:{sha}"])
    return out


def state(project: Path) -> dict:
    return {"db": db_rows(project), "orphans": orphans(project), "files": files(project)}


def outcome(fn, *a, **kw):
    try:
        res = fn(*a, **kw)
        return ("ok", res)
    except Exception as exc:  # noqa: BLE001 -- the observed behaviour IS the exception
        return ("raise", type(exc).__name__, str(getattr(exc, "message", None) or exc))


def run(label: str, project: Path, *args):
    """import_manual_ui(project, *args): outcome + state before/after, pinned under `label`."""
    before = state(project)
    res = outcome(app.import_manual_ui, str(project), *[str(a) if isinstance(a, Path) else a for a in args])
    after = state(project)
    snap(label, {"before": before, "outcome": res, "after": after})
    return res, before, after


def count(st: dict, table: str) -> int:
    return len(st["db"].get(table, []))


def ids(st: dict, table: str, col: str) -> list:
    return [r[col] for r in st["db"].get(table, [])]


def message(res) -> str:
    return res[1][0] if res[0] == "ok" else res[2]


# HARDENING-02B: the text of every B5 refusal (one gr.Error) and of the topic-status refusal (D3: same text as the
# ValueError of migration_draft_objects, now a gr.Error). S04 pins the B5 sentence word for word; the other refusals
# build it from the kinds they expect, so a kind found or missed changes the check, not only the digest.
STATUS_REFUSAL = ("La pubblicazione contiene contenuti già revisionati/approvati. "
                  "Usa un nuovo Publication ID per una nuova migrazione.")


def refused(res, before: dict, after: dict, work: dict | None = None) -> bool:
    """A refusal with no effect at all: gr.Error ('Error') with the B5 message for `work` (the topic-status text when
    `work` is None), and the database, the orphans and every project file -- Sources/ included -- as before."""
    text = STATUS_REFUSAL if work is None else app._reimport_refused_message(work)
    return res == ("raise", "Error", text) and after == before


# ============================================================ S01-S03: new import
p01 = new_project("S01 PDF")
r01, b01, a01 = run("S01_pdf_new", p01, PDF_A, "it-IT", "PUB-MAIN", "Manuale PDF")
check("S01: a new PDF import creates the publication, one topic per section, one block per paragraph (status Migrated)",
      r01[0] == "ok" and ids(a01, "topics", "topic_id") == ["DOC-0001", "DOC-0002"] and count(a01, "content_blocks") == 4
      and {r["status"] for r in a01["db"]["topics"] + a01["db"]["content_blocks"]} == {"Migrated"}
      and [(r["publication_id"], r["title"], r["source_lang"]) for r in a01["db"]["publications"]]
      == [("PUB-MAIN", "Manuale PDF", "it-IT")], a01["db"].get("topics"))
check("S01b: the message is the PDF one, with the counts; topics/blocks DataFrames are returned",
      r01[0] == "ok" and message(r01).startswith("✅ Manuale importato come draft in **2 topic semantici** e **0 tabelle strutturate**")
      and len(r01[1][1]) == 2 and len(r01[1][2]) == 4, message(r01))
check("S01c: the uploaded PDF is copied into Sources/ byte-for-byte; nothing else is added to the project",
      ["Sources/Manuale_A.pdf", "copy of Manuale_A.pdf"] in a01["files"]
      and [f[0] for f in a01["files"] if f not in b01["files"]] == ["Sources/", "Sources/Manuale_A.pdf"],
      [f for f in a01["files"] if f not in b01["files"]])

p02 = new_project("S02 DOCX")
r02, b02, a02 = run("S02_docx_new", p02, DOCX_A, "it-IT", "PUB-MAIN", "Manuale DOCX")
check("S02: a new DOCX import creates topics, blocks and the structured table referenced by a table block",
      r02[0] == "ok" and count(a02, "topics") == 2 and ids(a02, "structured_tables", "table_id") == ["DOC-0001-T001"]
      and any(b["block_type"] == "table" and b["object_ref_id"] == "DOC-0001-T001" for b in a02["db"]["content_blocks"]),
      a02["db"].get("structured_tables"))

p03 = new_project("S03 IDML")
r03, b03, a03 = run("S03_idml_new", p03, IDML_A, "it-IT", "PUB-MAIN", "Manuale IDML")
review03 = [f for f in a03["files"] if f[0] == "ManualSync_IDML_Review_PUB-MAIN.json"]
check("S03: a new IDML import creates the section topic, the table, the figure (missing link) and their blocks",
      r03[0] == "ok" and ids(a03, "topics", "topic_id") == ["IDML-0001"] and count(a03, "structured_tables") == 1
      and count(a03, "figures") == 1 and message(r03).startswith("✅ Manuale InDesign importato come draft"), message(r03))
check("S03b: the orphan story and the caption candidate become two pending review items in "
      "ManualSync_IDML_Review_PUB-MAIN.json",
      review03 and [it.get("reason") for it in review03[0][1]] == ["orphan_story", "caption_candidate"]
      and all(it.get("decision") is None for it in review03[0][1]), review03)
_fig_key03 = [ref.split("=", 1)[1] for r in a03["db"].get("figures", []) for ref in json.loads(r["source_refs_json"])
              if ref.startswith("idml_caption_candidate_key=")]
check("S03d: the caption-candidate review item points at the figure through the same (random) key",
      review03 and _fig_key03 == [it["item_key"] for it in review03[0][1] if it.get("reason") == "caption_candidate"],
      [_fig_key03, review03])
check("S03c: the IDML is copied into Sources/ only after the successful import",
      ["Sources/Manuale_A.idml", "copy of Manuale_A.idml"] in a03["files"], a03["files"])

# ============================================================ S04: PDF re-import over an IDML draft carrying translations
p04 = new_project("S04 Reimport")
outcome(app.import_manual_ui, str(p04), str(IDML_A), "it-IT", "PUB-MAIN", "Bozza IDML")
repo = IndustrialRepository(p04 / "ManualSync_CCMS.sqlite")
try:
    _blk = repo.conn.execute("SELECT block_id FROM content_blocks WHERE block_type='paragraph' ORDER BY rowid").fetchone()[0]
    _tbl = repo.conn.execute("SELECT table_id FROM structured_tables").fetchone()[0]
    _fig = repo.conn.execute("SELECT figure_id FROM figures").fetchone()[0]
    repo.set_translation("IDML-0001", "en-US", "Translated body", target_title="Introduction")
    repo.set_block_translation(_blk, "en-US", "This manual describes the hydraulic system.")
    repo.set_table_translation(_tbl, "en-US", caption="Parameters")
    repo.set_figure_translation(_fig, "en-US", caption="Hydraulic pump", alt_text="pump")
finally:
    repo.close()
r04, b04, a04 = run("S04_pdf_reimport_over_idml_draft", p04, PDF_B, "it-IT", "PUB-MAIN", "Revisione B")
# BEHAVIOR INTENTIONALLY CHANGED (HARDENING-02B, B5): until 9ae4989 this re-import retired the draft and the cascade
# took the four translations. What the retirement does to a draft with no human work is pinned by S04x below.
check("S04: (HARDENING-02B) a PDF re-import over an IDML draft carrying 4 translations is REFUSED with a gr.Error that "
      "names them and the two ways out -- the exact text the user reads",
      r04 == ("raise", "Error", "Reimportazione bloccata: la bozza di questa pubblicazione contiene lavoro fatto dopo "
                                "l'importazione, che andrebbe perso — traduzioni (4). Non è stato copiato alcun file e la "
                                "bozza non è stata modificata. Per importare questo file usa un nuovo Publication ID, "
                                "oppure ripristina il file originale e continua a lavorare sulla bozza esistente."), r04)
check("S04b: ... and the refusal has no effect: database, orphans and every file as before -- the IDML draft and its "
      "4 translations intact, the PDF not copied into Sources/ (DLS-03)",
      refused(r04, b04, a04, {"translations": 4}) and ids(a04, "topics", "topic_id") == ["IDML-0001"]
      and all(count(a04, t) == 1 for t in ("topic_translations", "block_translations", "structured_table_translations",
                                         "figure_translations"))
      and not any(f[0] == "Sources/Manuale_B.pdf" for f in a04["files"]), [r04[:2], [f[0] for f in a04["files"]]])

# ============================================================ S05/S06: DOCX over DOCX, IDML over IDML
p05 = new_project("S05 DOCX twice")
outcome(app.import_manual_ui, str(p05), str(DOCX_A), "it-IT", "PUB-MAIN", "Prima")
r05, b05, a05 = run("S05_docx_reimport_title_and_language", p05, DOCX_B, "en-US", "PUB-MAIN", "")
check("S05: DOCX over DOCX: the old table is replaced; title falls back to the file stem, language is updated",
      r05[0] == "ok" and ids(b05, "structured_tables", "table_id") == ["DOC-0001-T001"]
      and [json.loads(r["rows_json"]) for r in a05["db"]["structured_tables"]]
      == [[["Grandezza", "Limite"], ["Pressione", "16 bar"], ["Temperatura", "60 C"]]]
      and [(r["title"], r["source_lang"]) for r in a05["db"]["publications"]] == [("Manuale_B", "en-US")],
      a05["db"]["publications"])

p06 = new_project("S06 IDML twice")
outcome(app.import_manual_ui, str(p06), str(IDML_A), "it-IT", "PUB-MAIN", "Prima")
r06, b06, a06 = run("S06_idml_reimport_over_idml", p06, IDML_B, "it-IT", "PUB-MAIN", "Seconda")
review06_before = [f[1] for f in b06["files"] if f[0] == "ManualSync_IDML_Review_PUB-MAIN.json"]
review06 = [f[1] for f in a06["files"] if f[0] == "ManualSync_IDML_Review_PUB-MAIN.json"]
check("S06: IDML over IDML: retired after the new parse; the review JSON is REPLACED by the new file's items "
      "(the old orphan item is gone, the caption candidate is a new item)",
      r06[0] == "ok" and [[it.get("reason") for it in r] for r in review06] == [["caption_candidate"]]
      and review06[0][0]["item_key"] not in {it["item_key"] for it in review06_before[0]}
      and [json.loads(r["rows_json"]) for r in a06["db"]["structured_tables"]]
      == [[["Grandezza", "Limite"], ["Portata", "40 l/min"]]], [review06_before, review06])

# ============================================================ S07: refusal once content is no longer a migration draft
for tag, label, status in (("S07a", "S07_refused_approved", "Approved"), ("S07b", "S07b_refused_draft", "Draft")):
    p07 = new_project(label)
    outcome(app.import_manual_ui, str(p07), str(PDF_A), "it-IT", "PUB-MAIN", "Originale")
    con = sqlite3.connect(p07 / "ManualSync_CCMS.sqlite")
    con.execute("UPDATE topics SET status=? WHERE topic_id='DOC-0002'", (status,))   # simulates reviewed content
    con.commit()
    con.close()
    r07, b07, a07 = run(label, p07, PDF_B, "it-IT", "PUB-MAIN", "Tentativo")
    # BEHAVIOR INTENTIONALLY CHANGED (HARDENING-02B): a ValueError until 9ae4989 (never shown: show_error=False),
    # and the refused PDF had already been copied into Sources/.
    check(f"{tag}: a topic in status {status!r} refuses the re-import -- since HARDENING-02B a gr.Error with the same "
          "text (D3), so the user sees it", r07 == ("raise", "Error", STATUS_REFUSAL), r07)
    check(f"{tag}b: the database is untouched by the refusal", a07["db"] == b07["db"], "")
    check(f"{tag}c: (DLS-03) ... and nothing is copied any more: every project file as before, no Manuale_B.pdf",
          refused(r07, b07, a07) and not any(f[0] == "Sources/Manuale_B.pdf" for f in a07["files"]),
          [f[0] for f in a07["files"]])
p07i = new_project("S07c IDML refused")
outcome(app.import_manual_ui, str(p07i), str(IDML_A), "it-IT", "PUB-MAIN", "Originale")
con = sqlite3.connect(p07i / "ManualSync_CCMS.sqlite")
con.execute("UPDATE topics SET status='Approved'")
con.commit()
con.close()
r07i, b07i, a07i = run("S07c_refused_idml", p07i, IDML_B, "it-IT", "PUB-MAIN", "Tentativo")
check("S07c: an IDML re-import is refused the same way (gr.Error since HARDENING-02B, D3), and nothing is copied",
      refused(r07i, b07i, a07i), r07i)

# ============================================================ S08-S11: failures during the re-import
p08 = new_project("S08 corrupt PDF")
outcome(app.import_manual_ui, str(p08), str(PDF_A), "it-IT", "PUB-MAIN", "Originale")
r08, b08, a08 = run("S08_corrupt_pdf_after_retire", p08, PDF_CORRUPT, "it-IT", "PUB-MAIN", "Rotto")
check("S08: a corrupt PDF raises the raw PyMuPDF error (not a gr.Error)", r08[0] == "raise" and r08[1] == "FileDataError", r08)
check("S08b (pinned quirk): the old draft was retired BEFORE the parse -- the publication is left with no topic",
      count(b08, "topics") == 2 and count(a08, "topics") == 0 and count(a08, "publication_topics") == 0
      and count(a08, "publications") == 1, {t: count(a08, t) for t in ("publications", "topics", "content_blocks")})
r08b, b08b, a08b = run("S08c_good_pdf_after_data_loss", p08, PDF_B, "it-IT", "PUB-MAIN", "Ripristino")
check("S08c: a publication with no linked topic accepts the next import (nothing to refuse)",
      r08b[0] == "ok" and count(a08b, "topics") == 3, message(r08b))

p09 = new_project("S09 empty PDF")
outcome(app.import_manual_ui, str(p09), str(PDF_A), "it-IT", "PUB-MAIN", "Originale")
r09, b09, a09 = run("S09_empty_pdf_after_retire", p09, PDF_EMPTY, "it-IT", "PUB-MAIN", "Vuoto")
check("S09: an empty PDF raises ValueError 'Nessun topic ricostruito dal PDF.' after the old draft was retired",
      r09 == ("raise", "ValueError", "Nessun topic ricostruito dal PDF.") and count(a09, "topics") == 0, r09)

p10 = new_project("S10 corrupt DOCX")
outcome(app.import_manual_ui, str(p10), str(DOCX_A), "it-IT", "PUB-MAIN", "Originale")
r10, b10, a10 = run("S10_corrupt_docx_after_retire", p10, DOCX_CORRUPT, "it-IT", "PUB-MAIN", "Rotto")
check("S10: a corrupt DOCX raises gr.Error with the Word message; the old draft is already gone",
      r10 == ("raise", "Error", "Il documento Word non è valido o è danneggiato. Verifica il file e riprova.")
      and count(b10, "topics") == 2 and count(a10, "topics") == 0 and count(a10, "structured_tables") == 0, r10)

p11 = new_project("S11 corrupt IDML")
outcome(app.import_manual_ui, str(p11), str(IDML_A), "it-IT", "PUB-MAIN", "Originale")
r11, b11, a11 = run("S11_corrupt_idml_keeps_draft", p11, IDML_CORRUPT, "it-IT", "PUB-MAIN", "Rotto")
check("S11: a corrupt IDML raises gr.Error with the InDesign message",
      r11[:2] == ("raise", "Error") and r11[2] == app._LEGACY_IDML_FRIENDLY_ERROR, r11)
check("S11b: ... and leaves the previous draft, its files and its review state exactly as they were", a11 == b11, "")

# ============================================================ S12: .doc through the converter seam
_orig_convert = app.convert_legacy_doc_to_docx


@contextlib.contextmanager
def _fake_convert_ok(_src):
    yield DOCX_A


def _fake_convert_fail(_src):
    raise LegacyDocConversionError("simulated converter failure")


try:
    app.convert_legacy_doc_to_docx = _fake_convert_ok
    p12 = new_project("S12 DOC")
    r12, b12, a12 = run("S12_doc_converted", p12, DOC_LEGACY, "it-IT", "PUB-MAIN", "Legacy")
    check("S12: a .doc goes through the converter and the DOCX pipeline; provenance names the .doc, not the temp file",
          r12[0] == "ok" and count(a12, "topics") == 2
          and all("Legacy.doc" in r["source_refs_json"] for r in a12["db"]["topics"])
          and ["Sources/Legacy.doc", "copy of Legacy.doc"] in a12["files"], [r["source_refs_json"] for r in a12["db"]["topics"]])
    app.convert_legacy_doc_to_docx = _fake_convert_fail
    r12b, b12b, a12b = run("S12b_doc_converter_failure", p12, DOC_LEGACY, "it-IT", "PUB-MAIN", "Legacy")
    check("S12b: a converter failure raises gr.Error with the .doc message -- after the old draft was retired",
          r12b == ("raise", "Error", app._LEGACY_DOC_FRIENDLY_ERROR) and count(a12b, "topics") == 0, r12b)
finally:
    app.convert_legacy_doc_to_docx = _orig_convert
check("S12c: the converter seam is restored", app.convert_legacy_doc_to_docx is _orig_convert, "")

# ============================================================ S13/S14: the guard and the defaults
p13 = new_project("S13 guard")
guard = {"no project": outcome(app.import_manual_ui, "", str(PDF_A), "it-IT", "PUB-MAIN", "x"),
         "no file": outcome(app.import_manual_ui, str(p13), None, "it-IT", "PUB-MAIN", "x")}
check("S13: without project or file -> gr.Error 'Apri un progetto e carica il manuale corrente.', nothing created",
      all(v == ("raise", "Error", "Apri un progetto e carica il manuale corrente.") for v in guard.values())
      and not (p13 / "Sources").exists(), guard)
snap("S13_guard", {"outcome": guard, "after": state(p13)})

p14 = new_project("S14 defaults")
r14, b14, a14 = run("S14_defaults", p14, PDF_A, None, None, None)
check("S14: publication id None -> PUB-MAIN, language None -> the project's source_lang, title None -> file stem",
      r14[0] == "ok" and [(r["publication_id"], r["title"], r["source_lang"]) for r in a14["db"]["publications"]]
      == [("PUB-MAIN", "Manuale_A", "it-IT")], a14["db"].get("publications"))
p14b = new_project("S14b padded id")
r14b, b14b, a14b = run("S14b_padded_publication_id", p14b, PDF_A, " en-US ", "  PUB-X  ", "Titolo")
check("S14b: publication id and language are stripped",
      [(r["publication_id"], r["source_lang"]) for r in a14b["db"]["publications"]] == [("PUB-X", "en-US")],
      a14b["db"].get("publications"))

# ============================================================ S15: guided onboarding -> same path
_orig_root = app.DATA_ROOT
try:
    app.DATA_ROOT = TMP / "onboarding_root"
    app.DATA_ROOT.mkdir()
    ob = outcome(app._guided_create_and_import, "Onboarding Test", str(PDF_A), "it-IT")
    ok15 = ob[0] == "ok"
    p15 = Path(ob[1][0]) if ok15 else None                  # `active`: the project folder create_project_ui made
    st15 = state(p15) if ok15 and p15.is_dir() else None
    check("S15a: the onboarding project is created under DATA_ROOT", ok15 and p15.parent == app.DATA_ROOT,
          ob[1][0] if ok15 else ob)
    check("S15: the guided onboarding creates the project and imports through import_manual_ui (same PUB-MAIN result)",
          ok15 and st15 is not None and ids(st15, "topics", "topic_id") == ["DOC-0001", "DOC-0002"]
          and [(r["publication_id"], r["title"]) for r in st15["db"]["publications"]] == [("PUB-MAIN", "Onboarding Test")],
          ob if not ok15 else st15["db"].get("publications"))
    check("S15b: the guided result is the import message followed by the next-step sentence",
          ok15 and ob[1][3].startswith("✅ Manuale importato come draft in **2 topic semantici**")
          and ob[1][3].endswith("Prossimo passo: rivedi il contenuto importato prima di procedere."), ob[1][3] if ok15 else ob)
    snap("S15_guided_onboarding", {"result": ob[1][3] if ok15 else ob, "topics": ob[1][4] if ok15 else None,
                                   "blocks": ob[1][5] if ok15 else None, "after": st15})
finally:
    app.DATA_ROOT = _orig_root
check("S15c: app.DATA_ROOT is restored", app.DATA_ROOT == _orig_root, "")

# ============================================================ S16 (CP5.2b-0): title / language, first import vs re-import
p16 = new_project("S16 asymmetry")
b16a = state(p16)
r16a = outcome(app.import_manual_ui, str(p16), str(PDF_A), "it-it", "PUB-MAIN", "  Titolo  ")
a16a = state(p16)
r16b = outcome(app.import_manual_ui, str(p16), str(PDF_B), "it-it", "PUB-MAIN", "  Titolo  ")
a16b = state(p16)


def _pub(st):
    return [(r["publication_id"], r["title"], r["source_lang"]) for r in st["db"].get("publications", [])]


check("S16: first import (create_publication): the title is stripped and 'it-it' is canonicalized to 'it-IT'",
      r16a[0] == "ok" and _pub(a16a) == [("PUB-MAIN", "Titolo", "it-IT")], _pub(a16a))
check("S16b (pinned asymmetry): re-import (inline UPDATE): title and language are stored as received",
      r16b[0] == "ok" and _pub(a16b) == [("PUB-MAIN", "  Titolo  ", "it-it")], _pub(a16b))
check("S16c: the topics are canonicalized on both paths (only the publication row differs)",
      {r["source_lang"] for r in a16a["db"]["topics"]} == {"it-IT"} == {r["source_lang"] for r in a16b["db"]["topics"]},
      [{r["source_lang"] for r in a16a["db"]["topics"]}, {r["source_lang"] for r in a16b["db"]["topics"]}])
snap("S16_publication_title_language_asymmetry",
     {"before": b16a, "first": {"outcome": r16a, "after": a16a}, "reimport": {"outcome": r16b, "after": a16b}})

# ============================================================ S04x / S17 (HARDENING-02A): human work, then a re-import
# B5. Each scenario imports a migration draft, does ONE kind of human work on it through the real entry point
# (AuthoringService, the review queue, the Studio autosave callback, the IDML review callback, ComplianceService, the
# guided onboarding) and re-imports. At 9ae4989 (HARDENING-02A) the guard read only the TOPIC status and a Studio save
# keeps 'Migrated', so most of this work was silently destroyed -- deleted, or re-created under the SAME id (PDF/DOCX:
# DOC-0001-B001 ...) or a NEW one (IDML, imported before the retirement: IDML-0001-2 ...). BEHAVIOR INTENTIONALLY
# CHANGED in HARDENING-02B: IndustrialRepository.migration_draft_human_work (+ the review JSON for IDML decisions)
# refuses every such re-import before any effect, with a gr.Error naming what it found. What still goes through is
# pinned as well: a draft with no human work (S04x), an autosave (S17f, S17j-bis), rows older than the draft (S17j2).
from authoring_service import AuthoringService  # noqa: E402
from compliance_intelligence import ComplianceService  # noqa: E402

ACTOR = "Redattore"
SRC_B001 = "Il presente manuale descrive il sistema idraulico."     # DOC-0001-B001 as imported from PDF_A
EDITED = "Testo corretto dal redattore."
SAVED = "Bozza salvata in locale (nessuna nuova versione creata)."   # studio_save_draft_ui on success


@contextlib.contextmanager
def repository(project: Path):
    repo = IndustrialRepository(project / "ManualSync_CCMS.sqlite")
    try:
        yield repo
    finally:
        repo.close()


def draft_project(name: str, source: Path, title: str = "Bozza") -> tuple[Path, dict]:
    """A new project holding one freshly imported migration draft, and its state."""
    project = new_project(name)
    outcome(app.import_manual_ui, str(project), str(source), "it-IT", "PUB-MAIN", title)
    return project, state(project)


def rerun(label: str, project: Path, draft: dict, action, *args, then=None):
    """run(), pinning also the draft as imported, the human work done on it and (`then`) a reading taken afterwards."""
    before = state(project)
    res = outcome(app.import_manual_ui, str(project), *[str(a) if isinstance(a, Path) else a for a in args])
    after = state(project)
    later = then() if then else None
    snap(label, {"draft": draft, "action": action, "before": before, "outcome": res, "after": after, "then": later})
    return res, before, after, later


def rows(st: dict, table: str, **match) -> list:
    return [r for r in st["db"].get(table, []) if all(r.get(k) == v for k, v in match.items())]


def versions(st: dict, table: str, col: str, object_id: str) -> list:
    return [r["version"] for r in rows(st, table, **{col: object_id})]


def actions(st: dict) -> list:
    return [(r["object_type"], r["object_id"], r["action"]) for r in st["db"].get("audit_events", [])]


def b001(st: dict) -> list:
    return rows(st, "content_blocks", block_id="DOC-0001-B001")


def review_items(st: dict) -> list:
    return [f[1] for f in st["files"] if f[0] == "ManualSync_IDML_Review_PUB-MAIN.json"][0]


def decisions(st: dict) -> list:
    return [(it["reason"], it["decision"], it["decided_by"]) for it in review_items(st)]


_LIVE = {"topic": ("topics", "topic_id"), "block": ("content_blocks", "block_id"),
         "table": ("structured_tables", "table_id"), "figure": ("figures", "figure_id")}


def dangling(st: dict) -> dict:
    """Human-work rows whose topic/block/table/figure no longer exists (they have no foreign key to it)."""
    live = {kind: set(ids(st, table, col)) for kind, (table, col) in _LIVE.items()}
    return {table: [(r[kcol], r["object_id"]) for r in st["db"].get(table, [])
                    if r[kcol] in live and r["object_id"] not in live[r[kcol]]]
            for table, kcol in (("audit_events", "object_type"), ("review_tasks", "object_type"),
                                ("authoring_drafts", "kind"), ("compliance_evidence", "evidence_kind"))}


NO_DANGLING = {"audit_events": [], "review_tasks": [], "authoring_drafts": [], "compliance_evidence": []}

# ------------------------------------------------------------ S04x: S04 without the translations (no human work)
p04x, d04x = draft_project("S04x Reimport no translations", IDML_A, "Bozza IDML")
r04x, b04x, a04x, _ = rerun("S04x_pdf_reimport_over_idml_draft_no_translations", p04x, d04x, None,
                            PDF_B, "it-IT", "PUB-MAIN", "Revisione B")
_TRANSLATIONS = ("topic_translations", "block_translations", "structured_table_translations", "figure_translations")
check("S04x: an IDML draft with NO translation (no human work at all) is retired by a PDF re-import -- topic, table, "
      "figure DELETEd, the new PDF topics replace them, blocks and version history cascaded, no orphan row",
      r04x[0] == "ok" and ids(b04x, "topics", "topic_id") == ["IDML-0001"]
      and ids(a04x, "topics", "topic_id") == ["DOC-0001", "DOC-0002", "DOC-0003"]
      and [r["topic_id"] for r in a04x["db"]["publication_topics"]] == ["DOC-0001", "DOC-0002", "DOC-0003"]
      and count(a04x, "structured_tables") == 0 and count(a04x, "figures") == 0 and count(b04x, "figure_versions") > 0
      and not any(r["topic_id"] == "IDML-0001" for r in a04x["db"]["content_blocks"])
      and not any(r["topic_id"] == "IDML-0001" for r in a04x["db"].get("topic_versions", []))
      and all(not v for v in a04x["orphans"].values()), a04x["orphans"])
check("S04x-b: no translation row before or after -- the case B5 keeps accepting",
      all(count(st, t) == 0 for st in (b04x, a04x) for t in _TRANSLATIONS), "")
check("S04x-c: before the re-import S04 and S04x differ ONLY by S04's four translation rows -- the translations alone "
      "turn this accepted retirement into S04's refusal",
      all(count(b04, t) == 1 for t in _TRANSLATIONS) and r04x[0] == "ok" and r04[:2] == ("raise", "Error")
      and normalized({k: v for k, v in b04x["db"].items() if k not in _TRANSLATIONS})
      == normalized({k: v for k, v in b04["db"].items() if k not in _TRANSLATIONS}), "")
check("S04x-d: publications is UPDATEd in place (same row, new title), not recreated",
      [(r["publication_id"], r["title"]) for r in a04x["db"]["publications"]] == [("PUB-MAIN", "Revisione B")]
      and a04x["db"]["publications"][0]["created_at"] == b04x["db"]["publications"][0]["created_at"],
      a04x["db"]["publications"])
check("S04x-e: the previous IDML source and its review JSON stay in the project next to the new PDF",
      ["Sources/Manuale_A.idml", "copy of Manuale_A.idml"] in a04x["files"]
      and ["Sources/Manuale_B.pdf", "copy of Manuale_B.pdf"] in a04x["files"]
      and any(f[0] == "ManualSync_IDML_Review_PUB-MAIN.json" for f in a04x["files"]), [f[0] for f in a04x["files"]])

# ------------------------------------------------------------ S17a: a Studio block edit (PDF, DOCX)
p17a, d17a = draft_project("S17a block edit", PDF_A)
with repository(p17a) as repo:
    act17a = outcome(AuthoringService(repo).save_block, "DOC-0001-B001", text=EDITED, expected_version=1, actor=ACTOR)
r17a, b17a, a17a, _ = rerun("S17a_block_edit_then_pdf_reimport", p17a, d17a, act17a, PDF_A, "it-IT", "PUB-MAIN", "Bozza")
check("S17a: a Studio save is human work -- DOC-0001-B001 goes to v2 with the new text (block_versions [1, 2]) and an "
      "AUTHORING_SAVE audit row, yet block and topics stay 'Migrated'",
      act17a[0] == "ok" and [(r["version"], r["text_content"], r["status"]) for r in b001(b17a)] == [(2, EDITED, "Migrated")]
      and versions(b17a, "block_versions", "block_id", "DOC-0001-B001") == [1, 2]
      and actions(b17a) == [("block", "DOC-0001-B001", "AUTHORING_SAVE")]
      and {r["status"] for r in b17a["db"]["topics"]} == {"Migrated"}, act17a)
# BEHAVIOR INTENTIONALLY CHANGED (HARDENING-02B): until 9ae4989 the block came back at v1 with the source text.
check("S17a-b: (B5) the PDF re-import is now REFUSED -- 'contenuti modificati e salvati (1)' and the AUTHORING_SAVE "
      "row -- with no effect: the edit stays at v2, nothing is copied",
      refused(r17a, b17a, a17a, {"versions": 1, "audit": 1})
      and [(r["version"], r["text_content"]) for r in b001(a17a)] == [(2, EDITED)]
      and versions(a17a, "block_versions", "block_id", "DOC-0001-B001") == [1, 2], [r17a, b001(a17a)])
check("S17a-c: the AUTHORING_SAVE row is still the only audit row and names a live block (nothing dangling)",
      actions(a17a) == actions(b17a) == [("block", "DOC-0001-B001", "AUTHORING_SAVE")] and dangling(a17a) == NO_DANGLING,
      dangling(a17a))

p17a2, d17a2 = draft_project("S17a2 block edit DOCX", DOCX_A)
_para17a2 = [r["block_id"] for r in d17a2["db"]["content_blocks"] if r["block_type"] == "paragraph"][0]
with repository(p17a2) as repo:
    act17a2 = outcome(AuthoringService(repo).save_block, _para17a2, text=EDITED, expected_version=1, actor=ACTOR)
r17a2, b17a2, a17a2, _ = rerun("S17a2_block_edit_then_docx_reimport", p17a2, d17a2, act17a2,
                               DOCX_A, "it-IT", "PUB-MAIN", "Bozza")
check("S17a2: (B5) DOCX is refused like PDF, with no effect: the edited paragraph stays at v2 with the edited text",
      act17a2[0] == "ok" and refused(r17a2, b17a2, a17a2, {"versions": 1, "audit": 1})
      and versions(b17a2, "block_versions", "block_id", _para17a2) == [1, 2]
      and [(r["version"], r["text_content"]) for r in rows(a17a2, "content_blocks", block_id=_para17a2)] == [(2, EDITED)],
      [r17a2, rows(a17a2, "content_blocks", block_id=_para17a2)])

# ------------------------------------------------------------ S17b: a Studio title edit
p17b, d17b = draft_project("S17b title edit", PDF_A)
with repository(p17b) as repo:
    act17b = outcome(AuthoringService(repo).save_topic, "DOC-0001", title="Introduzione e avvertenze",
                     expected_version=1, actor=ACTOR)
r17b, b17b, a17b, _ = rerun("S17b_title_edit_then_pdf_reimport", p17b, d17b, act17b, PDF_A, "it-IT", "PUB-MAIN", "Bozza")
check("S17b: a title edit takes topic DOC-0001 to v2 (topic_versions [1, 2]) and leaves it 'Migrated' -- the very "
      "status the guard reads",
      act17b[0] == "ok" and [(r["version"], r["title"], r["status"]) for r in rows(b17b, "topics", topic_id="DOC-0001")]
      == [(2, "Introduzione e avvertenze", "Migrated")]
      and versions(b17b, "topic_versions", "topic_id", "DOC-0001") == [1, 2], act17b)
check("S17b-b: (B5) the re-import is REFUSED with no effect -- the edited title stays at v2",
      refused(r17b, b17b, a17b, {"versions": 1, "audit": 1})
      and [(r["version"], r["title"]) for r in rows(a17b, "topics", topic_id="DOC-0001")] == [(2, "Introduzione e avvertenze")]
      and actions(a17b) == [("topic", "DOC-0001", "AUTHORING_SAVE")], [r17b, rows(a17b, "topics", topic_id="DOC-0001")])

# ------------------------------------------------------------ S17c: a block created in Studio
p17c, d17c = draft_project("S17c Studio block", PDF_A)
with repository(p17c) as repo:
    act17c = outcome(AuthoringService(repo).create_block_in_topic, "DOC-0001", "paragraph", actor=ACTOR,
                     text="Avvertenza aggiunta dal redattore.")
r17c, b17c, a17c, _ = rerun("S17c_studio_block_then_pdf_reimport", p17c, d17c, act17c, PDF_A, "it-IT", "PUB-MAIN", "Bozza")
_new17c = act17c[1]["block_id"] if act17c[0] == "ok" else ""
check("S17c: a block created in Studio gets a random id (BLK-<10 hex>) and status 'Draft' in a topic still 'Migrated'",
      bool(re.fullmatch(r"BLK-[0-9A-F]{10}", _new17c))
      and [(r["topic_id"], r["status"]) for r in rows(b17c, "content_blocks", block_id=_new17c)] == [("DOC-0001", "Draft")]
      and [r["status"] for r in rows(b17c, "topics", topic_id="DOC-0001")] == ["Migrated"]
      and [a[2] for a in actions(b17c)] == ["AUTHORING_CREATE", "ADD_BLOCK"], [act17c, actions(b17c)])
check("S17c-b: (B5) the re-import is REFUSED -- the 'Draft' block and its two audit rows (AUTHORING_CREATE, ADD_BLOCK) "
      "-- with no effect: the Studio block is still in its topic",
      refused(r17c, b17c, a17c, {"status": 1, "audit": 2})
      and [(r["topic_id"], r["status"]) for r in rows(a17c, "content_blocks", block_id=_new17c)] == [("DOC-0001", "Draft")],
      [r17c, ids(a17c, "content_blocks", "block_id")])
check("S17c-c: nothing is left dangling any more (until 9ae4989 the AUTHORING_CREATE row was)",
      actions(a17c) == actions(b17c) and dangling(a17c) == NO_DANGLING, dangling(a17c))

# ------------------------------------------------------------ S17d: an open review on a block -- PDF vs IDML
p17d, d17d = draft_project("S17d review PDF", PDF_A)
with repository(p17d) as repo:
    act17d = outcome(repo.request_review, "block", "DOC-0001-B001", requested_by=ACTOR, assigned_to="Revisore")
r17d, b17d, a17d, _ = rerun("S17d_review_block_then_pdf_reimport", p17d, d17d, act17d, PDF_A, "it-IT", "PUB-MAIN", "Bozza")
check("S17d: a review request puts the BLOCK 'In Review' with an OPEN task while its topic stays 'Migrated' -- the "
      "topic-status guard alone (until 9ae4989) let the PDF re-import through",
      act17d == ("ok", 1) and [r["status"] for r in b001(b17d)] == ["In Review"]
      and [(r["object_id"], r["status"]) for r in b17d["db"]["review_tasks"]] == [("DOC-0001-B001", "OPEN")]
      and [r["status"] for r in b17d["db"]["topics"]] == ["Migrated", "Migrated"], act17d)
check("S17d-b: (B5) PDF: REFUSED -- status, REQUEST_REVIEW row and review task -- with no effect: block 'In Review', "
      "task OPEN",
      refused(r17d, b17d, a17d, {"status": 1, "audit": 1, "review": 1})
      and [r["status"] for r in b001(a17d)] == ["In Review"]
      and [(r["object_id"], r["status"]) for r in a17d["db"]["review_tasks"]] == [("DOC-0001-B001", "OPEN")],
      [r17d, b001(a17d)])

p17d2, d17d2 = draft_project("S17d2 review IDML", IDML_A)
with repository(p17d2) as repo:
    act17d2 = outcome(repo.request_review, "block", "IDML-0001-B001", requested_by=ACTOR, assigned_to="Revisore")
r17d2, b17d2, a17d2, _ = rerun("S17d2_review_block_then_idml_reimport", p17d2, d17d2, act17d2,
                               IDML_A, "it-IT", "PUB-MAIN", "Bozza")
check("S17d2: (B5) IDML: REFUSED the same way BEFORE the new file is imported -- no IDML-0001-2 draft, no new assets, "
      "the review JSON untouched; the OPEN task still names a live block (until 9ae4989 it was left dangling)",
      act17d2 == ("ok", 1) and refused(r17d2, b17d2, a17d2, {"status": 1, "audit": 1, "review": 1})
      and ids(a17d2, "topics", "topic_id") == ["IDML-0001"] and dangling(a17d2) == NO_DANGLING, [r17d2, dangling(a17d2)])

# ------------------------------------------------------------ S17e: an archived topic
p17e, d17e = draft_project("S17e archived topic", PDF_A)
with repository(p17e) as repo:
    act17e = outcome(AuthoringService(repo).archive_topic, "DOC-0002", actor=ACTOR)
r17e, b17e, a17e, _ = rerun("S17e_archived_topic_then_pdf_reimport", p17e, d17e, act17e, PDF_A, "it-IT", "PUB-MAIN", "Bozza")
check("S17e: archiving sets archived_at; the topic stays 'Migrated' and linked to the publication",
      act17e == ("ok", None) and [(bool(r["archived_at"]), r["status"]) for r in rows(b17e, "topics", topic_id="DOC-0002")]
      == [(True, "Migrated")] and "DOC-0002" in [r["topic_id"] for r in b17e["db"]["publication_topics"]], act17e)
check("S17e-b: (B5) the re-import is REFUSED -- the archive and its ARCHIVE row -- with no effect: DOC-0002 stays "
      "archived (until 9ae4989 the archive was undone)",
      refused(r17e, b17e, a17e, {"archived": 1, "audit": 1})
      and [bool(r["archived_at"]) for r in rows(a17e, "topics", topic_id="DOC-0002")] == [True], [r17e])

# ------------------------------------------------------------ S17f: an autosave with no human change
p17f, d17f = draft_project("S17f autosave", PDF_A)
act17f = outcome(app.studio_save_draft_ui, str(p17f), "block", "DOC-0001-B001", "paragraph", "", 1,
                 "", SRC_B001, "", "", None, "", "", ACTOR)                       # a blur with nothing changed
r17f, b17f, a17f, _ = rerun("S17f_autosave_no_change_then_pdf_reimport", p17f, d17f, act17f,
                            PDF_A, "it-IT", "PUB-MAIN", "Bozza")
check("S17f: a Studio blur with nothing changed still writes an autosave row (dirty=1, the unchanged text); no version, "
      "no audit row, no status change",
      act17f == ("ok", SAVED)
      and [(r["kind"], r["object_id"], r["base_version"], json.loads(r["payload_json"]), r["dirty"])
           for r in b17f["db"]["authoring_drafts"]] == [("block", "DOC-0001-B001", 1, {"text": SRC_B001}, 1)]
      and b17f["db"]["content_blocks"] == d17f["db"]["content_blocks"] and not b17f["db"]["audit_events"], act17f)
check("S17f-b: the re-import is ACCEPTED; the autosave row survives untouched, keyed on the re-created block",
      r17f[0] == "ok" and a17f["db"]["authoring_drafts"] == b17f["db"]["authoring_drafts"]
      and dangling(a17f) == NO_DANGLING, message(r17f))

# ------------------------------------------------------------ S17g: IDML review decisions (database + JSON, JSON only)
p17g, d17g = draft_project("S17g caption confirmed", IDML_A)
act17g = outcome(app.idml_review_confirm_ui, str(p17g), "QA",
                 *[app._IDML_CONFIRM_CAPTION if it["reason"] == "caption_candidate" else app._IDML_ACCEPT_EXCLUSION
                   for it in review_items(d17g)])
r17g, b17g, a17g, _ = rerun("S17g_idml_caption_confirmed_then_idml_reimport", p17g, d17g, act17g,
                            IDML_A, "it-IT", "PUB-MAIN", "Bozza")
check("S17g: confirming the caption is human work on BOTH sides: figure v2 with the caption (figure_versions [1, 2]) "
      "and explicit decisions in the review JSON (decided_by 'QA')",
      act17g[0] == "ok" and [(r["version"], r["caption"]) for r in b17g["db"]["figures"]] == [(2, "Didascalia pompa idraulica")]
      and [r["version"] for r in b17g["db"]["figure_versions"]] == [1, 2]
      and decisions(b17g) == [("orphan_story", "accepted_exclusion", "QA"), ("caption_candidate", "confirm_association", "QA")],
      [b17g["db"]["figures"], decisions(b17g)])
check("S17g-b: (B5) the IDML re-import is REFUSED -- the figure saved at v2 and the 2 explicit JSON decisions -- before "
      "the new file is imported: figure, caption, decisions and review key untouched (until 9ae4989 all were lost)",
      refused(r17g, b17g, a17g, {"versions": 1, "idml_decisions": 2})
      and [(r["figure_id"], r["version"], r["caption"]) for r in a17g["db"]["figures"]]
      == [("IDML-0001-F002", 2, "Didascalia pompa idraulica")] and decisions(a17g) == decisions(b17g), [r17g, decisions(a17g)])

p17g2, d17g2 = draft_project("S17g2 JSON-only decisions", IDML_A)
act17g2 = outcome(app.idml_review_confirm_ui, str(p17g2), "QA",
                  *[app._IDML_KEEP_SEPARATE if it["reason"] == "caption_candidate" else app._IDML_ACCEPT_EXCLUSION
                    for it in review_items(d17g2)])
r17g2, b17g2, a17g2, _ = rerun("S17g2_idml_json_only_decisions_then_idml_reimport", p17g2, d17g2, act17g2,
                               IDML_A, "it-IT", "PUB-MAIN", "Bozza")
check("S17g2: 'keep separate' and 'accept exclusion' are explicit decisions that live ONLY in the review JSON -- the "
      "database is untouched (figure still v1, no audit row)",
      act17g2[0] == "ok" and b17g2["db"] == d17g2["db"]
      and decisions(b17g2) == [("orphan_story", "accepted_exclusion", "QA"), ("caption_candidate", "keep_separate", "QA")],
      decisions(b17g2))
check("S17g2-b: what tells an explicit decision from the file the import writes: decision set, decided_by and "
      "decided_at filled -- in the generated file every item is None / '' / ''",
      all(it["decision"] is None and it["decided_by"] == "" and it["decided_at"] == "" for it in review_items(d17g2))
      and all(it["decision"] and it["decided_by"] and it["decided_at"] for it in review_items(b17g2)), review_items(d17g2))
check("S17g2-c: (B5) the IDML re-import is REFUSED on the 2 JSON-only decisions alone, with no effect: the decisions "
      "stay in the review JSON (until 9ae4989 they were replaced by undecided items)",
      refused(r17g2, b17g2, a17g2, {"idml_decisions": 2})
      and decisions(a17g2) == [("orphan_story", "accepted_exclusion", "QA"), ("caption_candidate", "keep_separate", "QA")],
      [r17g2, decisions(a17g2)])

# ------------------------------------------------------------ S17h: the guided onboarding (DLS-02), same / empty name
_orig_root17 = app.DATA_ROOT
try:
    app.DATA_ROOT = TMP / "onboarding_root_17h"
    app.DATA_ROOT.mkdir()
    g17h1 = outcome(app._guided_create_and_import, "Manuale Pompa", str(PDF_A), "it-IT")
    p17h = Path(g17h1[1][0])
    d17h = state(p17h)
    with repository(p17h) as repo:
        act17h = outcome(AuthoringService(repo).save_block, "DOC-0001-B001", text=EDITED, expected_version=1, actor=ACTOR)
    b17h = state(p17h)
    g17h2 = outcome(app._guided_create_and_import, "Manuale Pompa", str(PDF_B), "it-IT")
    a17h = state(p17h)
    dirs17h = sorted(p.name for p in app.DATA_ROOT.iterdir() if p.is_dir())
    snap("S17h_guided_same_name_reimport",
         {"first": [g17h1[0], g17h1[1][0], g17h1[1][3]], "draft": d17h, "action": act17h, "before": b17h,
          "second": [g17h2[0], g17h2[1][0], g17h2[1][3]] if g17h2[0] == "ok" else g17h2, "after": a17h, "projects": dirs17h})

    app.DATA_ROOT = TMP / "onboarding_root_17h2"
    app.DATA_ROOT.mkdir()
    g17h3 = outcome(app.guided_confirm_language_ui, "", str(PDF_A), "Italiano (it-IT)")
    p17h2 = Path(g17h3[1][0])
    d17h2 = state(p17h2)
    with repository(p17h2) as repo:
        act17h2 = outcome(AuthoringService(repo).save_block, "DOC-0001-B001", text=EDITED, expected_version=1, actor=ACTOR)
    b17h2 = state(p17h2)
    g17h4 = outcome(app.guided_confirm_language_ui, "   ", str(PDF_A), "Italiano (it-IT)")
    a17h2 = state(p17h2)
    dirs17h2 = sorted(p.name for p in app.DATA_ROOT.iterdir() if p.is_dir())
    snap("S17h2_guided_blank_name_reimport",
         {"first": [g17h3[0], g17h3[1][0], g17h3[1][3], g17h3[1][-1]], "draft": d17h2, "action": act17h2, "before": b17h2,
          "second": [g17h4[0], g17h4[1][0], g17h4[1][3], g17h4[1][-1]] if g17h4[0] == "ok" else g17h4,
          "after": a17h2, "projects": dirs17h2})
finally:
    app.DATA_ROOT = _orig_root17
check("S17h: (B5) the guided onboarding run again with the SAME name still reopens the existing project (no second "
      "folder), but its re-import is now REFUSED: the gr.Error reaches the user from the guided path and the draft, the "
      "Studio edit (v2) and Sources/ are untouched",
      g17h1[0] == "ok" and act17h[0] == "ok" and refused(g17h2, b17h, a17h, {"versions": 1, "audit": 1})
      and dirs17h == [p17h.name] and [(r["version"], r["text_content"]) for r in b001(a17h)] == [(2, EDITED)],
      [g17h2, dirs17h, b001(a17h)])
check("S17h-b: (DLS-02, residual) the refusal does not say that an existing project was reopened -- create_project_ui's "
      "'esiste già' message is still discarded", g17h2[0] == "raise" and "esiste già" not in g17h2[2], g17h2)
check("S17h2: (B5) an EMPTY name (blank, then spaces) still falls back to the file stem and lands in the same project, "
      "whose re-import over the edit is now REFUSED the same way, with no effect",
      g17h3[0] == "ok" and g17h3[1][-1] == "Manuale_A" and act17h2[0] == "ok"
      and refused(g17h4, b17h2, a17h2, {"versions": 1, "audit": 1}) and dirs17h2 == [p17h2.name]
      and [(r["version"], r["text_content"]) for r in b001(a17h2)] == [(2, EDITED)], [g17h4, dirs17h2, b001(a17h2)])
check("S17h-c: app.DATA_ROOT is restored", app.DATA_ROOT == _orig_root17, "")

# ------------------------------------------------------------ S17i: a second revision with the SAME file name (DLS-03)
(INPUTS / "rev1").mkdir()
(INPUTS / "rev2").mkdir()
PDF_C1 = make_pdf("rev1/Manuale_C.pdf", [("1 Sicurezza", ["Revisione 1: indossare i guanti."])])
PDF_C2 = make_pdf("rev2/Manuale_C.pdf", [("1 Sicurezza", ["Revisione 2: indossare guanti e occhiali."])])
for _pdf in (PDF_C1, PDF_C2):                    # files() names an uploaded copy by WHICH input it is
    INPUT_SHA[hashlib.sha256(_pdf.read_bytes()).hexdigest()] = _pdf.relative_to(INPUTS).as_posix()
p17i, d17i = draft_project("S17i same-name file", PDF_C1)
with repository(p17i) as repo:
    act17i = outcome(AuthoringService(repo).bulk_approve_publication, "PUB-MAIN", actor="Revisore")
r17i, b17i, a17i, _ = rerun("S17i_same_name_file_refused_overwrites_source", p17i, d17i, act17i,
                            PDF_C2, "it-IT", "PUB-MAIN", "Revisione 2")


def source_copy(st: dict, rel: str) -> list:
    return [f[1] for f in st["files"] if f[0] == rel]


# The label keeps the 9ae4989 name; BEHAVIOR INTENTIONALLY CHANGED (HARDENING-02B, DLS-03): the source is no longer
# overwritten.
check("S17i: after a human approval the next revision under the SAME file name is refused (gr.Error since "
      "HARDENING-02B, D3) and the database is untouched",
      act17i == ("ok", {"topics_approved": 1, "review_tasks_closed": 0}) and r17i == ("raise", "Error", STATUS_REFUSAL)
      and a17i["db"] == b17i["db"], [act17i, r17i])
check("S17i-b (DLS-03): ... and the refused file no longer overwrites Sources/Manuale_C.pdf: the approved content keeps "
      "the copy of the revision it came from (until 9ae4989 it was replaced by the refused one)",
      refused(r17i, b17i, a17i) and source_copy(a17i, "Sources/Manuale_C.pdf") == ["copy of rev1/Manuale_C.pdf"],
      source_copy(a17i, "Sources/Manuale_C.pdf"))

# ------------------------------------------------------------ S17j: an audit row older than the draft it now names
p17j, d17j = draft_project("S17j audit before draft", PDF_A)
with repository(p17j) as repo:
    act17j = outcome(AuthoringService(repo).save_block, "DOC-0001-B001", text=EDITED, expected_version=1, actor=ACTOR)
# D7 (HARDENING-02B): kept, re-pinned. Until 9ae4989 the intermediate re-import below was destructive and left the
# AUTHORING_SAVE row OLDER than the block it named; it is refused now, so that state can no longer be produced through
# the application -- S17j2 seeds it directly, as a project re-imported before B5 holds it.
mid17j = outcome(app.import_manual_ui, str(p17j), str(PDF_A), "it-IT", "PUB-MAIN", "Bozza")
r17j, b17j, a17j, _ = rerun("S17j_audit_before_draft_then_pdf_reimport", p17j, d17j,
                            {"edit": act17j, "intermediate_reimport": mid17j}, PDF_B, "it-IT", "PUB-MAIN", "Revisione B")
_save_at17j = [r["created_at"] for r in b17j["db"]["audit_events"] if r["action"] == "AUTHORING_SAVE"]
check("S17j: (D7) the intermediate re-import after the edit is REFUSED (version + audit), so the AUTHORING_SAVE row "
      "stays NEWER than the block it names and the edit stays at v2",
      act17j[0] == "ok" and mid17j == ("raise", "Error", app._reimport_refused_message({"versions": 1, "audit": 1}))
      and len(_save_at17j) == 1 and _save_at17j[0] >= b001(b17j)[0]["created_at"]
      and versions(b17j, "block_versions", "block_id", "DOC-0001-B001") == [1, 2], [mid17j, _save_at17j, b001(b17j)])
check("S17j-b: the next re-import (another revision) is refused the same way, with no effect",
      refused(r17j, b17j, a17j, {"versions": 1, "audit": 1})
      and actions(a17j) == [("block", "DOC-0001-B001", "AUTHORING_SAVE")], [r17j, actions(a17j)])

# The same stale row, written directly (the shape S17j leaves behind, as in a project re-imported before B5): it does not
# depend on the destructive re-import S17j needs to produce it, so it still describes that state once that changes.
p17j2, d17j2 = draft_project("S17j2 seeded audit before draft", PDF_A)
con = sqlite3.connect(p17j2 / "ManualSync_CCMS.sqlite")
con.execute("INSERT INTO audit_events(object_type,object_id,action,actor,detail_json,created_at) VALUES(?,?,?,?,?,?)",
            ("block", "DOC-0001-B001", "AUTHORING_SAVE", ACTOR,
             json.dumps({"from_version": 1, "to_version": 2}, sort_keys=True), "2000-01-01T00:00:00+00:00"))
con.commit()
con.close()
r17j2, b17j2, a17j2, _ = rerun("S17j2_seeded_audit_before_draft_then_pdf_reimport", p17j2, d17j2,
                               "seeded: AUTHORING_SAVE on DOC-0001-B001 dated 2000-01-01", PDF_B, "it-IT", "PUB-MAIN",
                               "Revisione B")
check("S17j2: a seeded AUTHORING_SAVE row dated before the draft (block still v1, never edited in this draft) does "
      "not stop the re-import: ACCEPTED, the row survives",
      r17j2[0] == "ok" and actions(a17j2) == actions(b17j2) == [("block", "DOC-0001-B001", "AUTHORING_SAVE")]
      and b17j2["db"]["audit_events"][0]["created_at"] < b001(b17j2)[0]["created_at"]
      and versions(b17j2, "block_versions", "block_id", "DOC-0001-B001") == [1]
      and [r["text_content"] for r in b001(a17j2)] == ["Revisione B del manuale."], [message(r17j2), actions(a17j2)])

# ------------------------------------------------------------ S17j-bis: an orphan autosave (IDML ids alternate)
p17jb, d17jb = draft_project("S17j-bis orphan autosave", IDML_A)
act17jb = outcome(app.studio_save_draft_ui, str(p17jb), "block", "IDML-0001-B001", "paragraph", "", 1,
                  "", "Bozza mai salvata dal redattore.", "", "", None, "", "", ACTOR)
r17jb, b17jb, a17jb, _ = rerun("S17jbis_orphan_autosave_after_idml_reimport", p17jb, d17jb, act17jb,
                               IDML_A, "it-IT", "PUB-MAIN", "Bozza")
r17jb2, _b17jb2, a17jb2, _ = rerun("S17jbis2_orphan_autosave_rebound_by_next_idml_reimport", p17jb, a17jb, None,
                                   IDML_A, "it-IT", "PUB-MAIN", "Bozza")
check("S17j-bis: an unsaved autosave does not stop an IDML re-import; the new draft gets new ids, so the autosave row "
      "is left DANGLING on IDML-0001-B001",
      act17jb == ("ok", SAVED) and r17jb[0] == "ok" and ids(a17jb, "topics", "topic_id") == ["IDML-0001-2"]
      and dangling(a17jb) == {**NO_DANGLING, "authoring_drafts": [("block", "IDML-0001-B001")]}, dangling(a17jb))
check("S17j-bis-b: the next IDML re-import alternates back to IDML-0001 (IDML-0001-2 is taken), so the same stale "
      "autosave is silently RE-BOUND to a new block whose version matches its base_version",
      r17jb2[0] == "ok" and ids(a17jb2, "topics", "topic_id") == ["IDML-0001"] and dangling(a17jb2) == NO_DANGLING
      and [(r["object_id"], r["base_version"], json.loads(r["payload_json"])) for r in a17jb2["db"]["authoring_drafts"]]
      == [("IDML-0001-B001", 1, {"text": "Bozza mai salvata dal redattore."})]
      and versions(a17jb2, "block_versions", "block_id", "IDML-0001-B001") == [1], [ids(a17jb2, "topics", "topic_id"),
                                                                                     dangling(a17jb2)])

# ------------------------------------------------------------ S17k: compliance evidence on a draft block
def _evidence_on_b001(project: Path):
    with repository(project) as repo:
        cs = ComplianceService(repo)
        cs.create_instrument("INS-1", "2006/42/CE", actor="Compliance", title="Direttiva macchine")
        cs.create_requirement("REQ-1", "Istruzioni per l'uso", actor="Compliance", instrument_id="INS-1")
        cs.create_scope("SC-1", "PUB-MAIN", actor="Compliance")
        cs.add_scope_reference("SC-1", "instrument", "INS-1", actor="Compliance")
        cs.set_applicability("SC-1", "REQ-1", "APPLIES", actor="Compliance")
        cs.add_evidence("EV-1", "block", actor="Compliance", object_id="DOC-0001-B001", object_version=1)
        cs.link_evidence_requirement("EV-1", "REQ-1", actor="Compliance")
        cs.set_assessment("SC-1", "REQ-1", "COVERED", actor="Compliance")
        return cs.evidence_for_requirement("REQ-1")


def _evidence_judgment(project: Path):
    with repository(project) as repo:
        return ComplianceService(repo).evidence_for_requirement("REQ-1")


p17k, d17k = draft_project("S17k compliance evidence", PDF_A)
act17k = outcome(_evidence_on_b001, p17k)
r17k, b17k, a17k, j17k = rerun("S17k_compliance_evidence_then_pdf_reimport", p17k, d17k, act17k,
                               PDF_B, "it-IT", "PUB-MAIN", "Revisione B", then=lambda: _evidence_judgment(p17k))
check("S17k: block evidence (DOC-0001-B001 v1, the imported text) backs a COVERED assessment",
      act17k[0] == "ok" and [(e["object_id"], e["object_version"], e["is_current"]) for e in act17k[1]]
      == [("DOC-0001-B001", 1, True)] and [r["status"] for r in b17k["db"]["compliance_assessment"]] == ["COVERED"], act17k)
check("S17k-b: (B5) a PDF re-import of ANOTHER revision is REFUSED on the evidence, with no effect: the block keeps "
      "the assessed text (until 9ae4989 it was replaced)",
      refused(r17k, b17k, a17k, {"evidence": 1}) and [(r["version"], r["text_content"]) for r in b001(a17k)]
      == [(1, SRC_B001)], [r17k, b001(a17k)])
check("S17k-c: ... so the evidence that judges itself CURRENT (same id, version 1) and the COVERED assessment still "
      "describe the text that was assessed",
      [(e["object_id"], e["object_version"], e["is_current"]) for e in j17k] == [("DOC-0001-B001", 1, True)]
      and [r["status"] for r in a17k["db"]["compliance_assessment"]] == ["COVERED"]
      and a17k["db"]["compliance_evidence"] == b17k["db"]["compliance_evidence"] and dangling(a17k) == NO_DANGLING, j17k)

# ============================================================ S18/S19 (HARDENING-02B): each kind alone, its filter, D1-D5
# What S17 cannot isolate: work that changes no version and no status (only the audit trail shows it), the time filter
# of every kind that has one (rows older than the draft, as a project re-imported before B5 holds them -- seeded,
# because the application can no longer produce them), the decisions D1-D5, and a refusal in a project with no
# Sources/ yet. Seeded timestamps are deliberately far in the past ("2000-...") or taken now, after the import.
import compliance_view  # noqa: E402
from datetime import datetime, timezone  # noqa: E402

LONG_AGO = "2000-01-01T00:00:00+00:00"


def seed(project: Path, sql: str, params=()) -> None:
    con = sqlite3.connect(project / "ManualSync_CCMS.sqlite")
    try:
        con.execute(sql, params)
        con.commit()
    finally:
        con.close()


def write_review(project: Path, items) -> None:
    (project / "ManualSync_IDML_Review_PUB-MAIN.json").write_text(
        items if isinstance(items, str) else json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")


# ------------------------------------------------------------ S18a-c: audit only (no version, no status)
p18a, d18a = draft_project("S18a reorder blocks", PDF_A)
with repository(p18a) as repo:
    act18a = outcome(AuthoringService(repo).reorder_blocks, "DOC-0001", ["DOC-0001-B002", "DOC-0001-B001"],
                     expected_block_ids=["DOC-0001-B001", "DOC-0001-B002"], actor=ACTOR)
r18a, b18a, a18a, _ = rerun("S18a_reorder_blocks_then_pdf_reimport", p18a, d18a, act18a, PDF_A, "it-IT", "PUB-MAIN", "Bozza")
check("S18a: (K5) reordering blocks changes no version and no status -- only the REORDER_BLOCKS row on the topic shows "
      "it, and that alone refuses the re-import, with no effect",
      act18a[0] == "ok" and {r["version"] for r in b18a["db"]["content_blocks"] + b18a["db"]["topics"]} == {1}
      and {r["status"] for r in b18a["db"]["content_blocks"]} == {"Migrated"}
      and actions(b18a) == [("topic", "DOC-0001", "REORDER_BLOCKS")] and refused(r18a, b18a, a18a, {"audit": 1}),
      [act18a, actions(b18a), r18a])

p18b, d18b = draft_project("S18b reorder topics", PDF_A)
with repository(p18b) as repo:
    act18b = outcome(AuthoringService(repo).reorder_publication_topics, "PUB-MAIN", ["DOC-0002", "DOC-0001"],
                     expected_topic_ids=["DOC-0001", "DOC-0002"], actor=ACTOR)
r18b, b18b, a18b, _ = rerun("S18b_reorder_topics_then_pdf_reimport", p18b, d18b, act18b, PDF_A, "it-IT", "PUB-MAIN", "Bozza")
check("S18b: (K5, publication level) reordering the topics leaves only a REORDER_TOPICS row on the PUBLICATION -- "
      "counted from the creation of the draft, it refuses the re-import, with no effect",
      act18b[0] == "ok" and {r["version"] for r in b18b["db"]["topics"]} == {1}
      and actions(b18b) == [("publication", "PUB-MAIN", "REORDER_TOPICS")] and refused(r18b, b18b, a18b, {"audit": 1})
      and [r["topic_id"] for r in sorted(a18b["db"]["publication_topics"], key=lambda r: r["position"])]
      == ["DOC-0002", "DOC-0001"], [act18b, actions(b18b), r18b])

p18c, d18c = draft_project("S18c archive and restore", PDF_A)
with repository(p18c) as repo:
    _auth18c = AuthoringService(repo)
    act18c = [outcome(_auth18c.archive_topic, "DOC-0002", actor=ACTOR), outcome(_auth18c.restore_topic, "DOC-0002", actor=ACTOR)]
r18c, b18c, a18c, _ = rerun("S18c_archive_restore_then_pdf_reimport", p18c, d18c, act18c, PDF_A, "it-IT", "PUB-MAIN", "Bozza")
check("S18c: (K5) a topic archived and restored is back to archived_at NULL at v1 -- the ARCHIVE and RESTORE rows alone "
      "refuse the re-import, with no effect",
      [a[0] for a in act18c] == ["ok", "ok"] and [r["archived_at"] for r in rows(b18c, "topics", topic_id="DOC-0002")] == [None]
      and {r["version"] for r in b18c["db"]["topics"]} == {1}
      and [a[2] for a in actions(b18c)] == ["ARCHIVE", "RESTORE"] and refused(r18c, b18c, a18c, {"audit": 2}),
      [act18c, actions(b18c), r18c])

# ------------------------------------------------------------ S18d/e: review tasks and their time filter (D1)
_REVIEW_SQL = ("INSERT INTO review_tasks(object_type,object_id,requested_by,assigned_to,status,comment,created_at) "
               "VALUES('block','DOC-0001-B001',?,'Revisore','OPEN','',?)")
p18d, d18d = draft_project("S18d stale review task", PDF_A)
seed(p18d, _REVIEW_SQL, (ACTOR, LONG_AGO))
r18d, b18d, a18d, _ = rerun("S18d_seeded_stale_review_then_pdf_reimport", p18d, d18d, "seeded: OPEN review task dated 2000",
                            PDF_A, "it-IT", "PUB-MAIN", "Bozza")
check("S18d: (K6, D1) an OPEN review task OLDER than the block it names (a project re-imported before B5) does not "
      "count: ACCEPTED, the task survives", r18d[0] == "ok" and a18d["db"]["review_tasks"] == b18d["db"]["review_tasks"]
      and [r["created_at"] for r in b18d["db"]["review_tasks"]] == [LONG_AGO], [r18d[:2], a18d["db"]["review_tasks"]])

p18e, d18e = draft_project("S18e fresh review task", PDF_A)
seed(p18e, _REVIEW_SQL, (ACTOR, datetime.now(timezone.utc).isoformat()))
r18e, b18e, a18e, _ = rerun("S18e_seeded_fresh_review_then_pdf_reimport", p18e, d18e, "seeded: OPEN review task dated now",
                            PDF_A, "it-IT", "PUB-MAIN", "Bozza")
check("S18e: (K6) the same task written after the import -- block status untouched, no audit row -- refuses on its own",
      [r["status"] for r in b001(b18e)] == ["Migrated"] and not b18e["db"]["audit_events"]
      and refused(r18e, b18e, a18e, {"review": 1}), r18e)


# ------------------------------------------------------------ S18f/g: compliance evidence, its filter and its link (D1)
def _evidence_via_ui(project: Path, *requirements: str):
    """The Compliance view's own path: link_ccms_object_as_evidence keys the evidence by kind+id+version
    (EV-block-DOC-0001-B001-v1) and REUSES an existing record -- the case K7's link date exists for."""
    with repository(project) as repo:
        cs = ComplianceService(repo)
        for rid in requirements:
            if not repo.conn.execute("SELECT 1 FROM compliance_requirement WHERE requirement_id=?", (rid,)).fetchone():
                cs.create_requirement(rid, f"Requisito {rid}", actor="Compliance")
        return [compliance_view.link_ccms_object_as_evidence(cs, "block", "DOC-0001-B001", 1, rid, actor="Compliance")
                for rid in requirements]


def _backdate_evidence(project: Path) -> None:
    seed(project, "UPDATE compliance_evidence SET added_at=?", (LONG_AGO,))
    seed(project, "UPDATE compliance_evidence_link SET linked_at=?", (LONG_AGO,))


p18f, d18f = draft_project("S18f stale evidence", PDF_A)
act18f = outcome(_evidence_via_ui, p18f, "REQ-1")
_backdate_evidence(p18f)
r18f, b18f, a18f, _ = rerun("S18f_seeded_stale_evidence_then_pdf_reimport", p18f, d18f, act18f, PDF_A, "it-IT", "PUB-MAIN",
                            "Bozza")
check("S18f: (K7, D1) evidence added AND linked before the block was created (the S17k state of a project re-imported "
      "before B5) does not count: ACCEPTED",
      act18f == ("ok", ["EV-block-DOC-0001-B001-v1"]) and r18f[0] == "ok"
      and [r["added_at"] for r in b18f["db"]["compliance_evidence"]] == [LONG_AGO], [act18f, r18f[:2]])

p18g, d18g = draft_project("S18g stale evidence new link", PDF_A)
act18g = [outcome(_evidence_via_ui, p18g, "REQ-1")]
_backdate_evidence(p18g)
act18g.append(outcome(_evidence_via_ui, p18g, "REQ-2"))                          # a NEW link, after the import
r18g, b18g, a18g, _ = rerun("S18g_stale_evidence_new_link_then_pdf_reimport", p18g, d18g, act18g, PDF_A, "it-IT",
                            "PUB-MAIN", "Bozza")
check("S18g: (K7, D1) the same old evidence record linked AGAIN after the import (same deterministic id, reused) counts "
      "by its link date: REFUSED, with no effect",
      act18g == [("ok", ["EV-block-DOC-0001-B001-v1"])] * 2
      and [r["added_at"] for r in b18g["db"]["compliance_evidence"]] == [LONG_AGO]
      and len(b18g["db"]["compliance_evidence_link"]) == 2 and refused(r18g, b18g, a18g, {"evidence": 1}), [act18g, r18g])

# ------------------------------------------------------------ S18h-k: IDML review JSON -- every format (D4), filter, D2
p18h, d18h = draft_project("S18h JSON decisions PDF reimport", IDML_A)
act18h = outcome(app.idml_review_confirm_ui, str(p18h), "QA",
                 *[app._IDML_KEEP_SEPARATE if it["reason"] == "caption_candidate" else app._IDML_ACCEPT_EXCLUSION
                   for it in review_items(d18h)])
r18h, b18h, a18h, _ = rerun("S18h_idml_json_decisions_then_pdf_reimport", p18h, d18h, act18h, PDF_B, "it-IT", "PUB-MAIN",
                            "Revisione B")
check("S18h: (K8, D4) JSON-only IDML decisions count for a PDF upload too -- it would retire the draft they describe: "
      "REFUSED, with no effect (the PDF is not copied)",
      act18h[0] == "ok" and refused(r18h, b18h, a18h, {"idml_decisions": 2})
      and not any(f[0] == "Sources/Manuale_B.pdf" for f in a18h["files"]), r18h)

p18i, d18i = draft_project("S18i stale JSON decisions", IDML_A)
write_review(p18i, [{**it, "decision": "accepted_exclusion" if it["reason"] == "orphan_story" else "keep_separate",
                     "decided_by": "QA", "decided_at": LONG_AGO} for it in review_items(d18i)])
r18i, b18i, a18i, _ = rerun("S18i_seeded_stale_json_decisions_then_idml_reimport", p18i, d18i,
                            "seeded: both review items decided by QA in 2000", IDML_A, "it-IT", "PUB-MAIN", "Bozza")
check("S18i: (K8 filter) decisions dated BEFORE the draft was created (a review file left by an earlier draft) do not "
      "count: ACCEPTED, and the IDML re-import rewrites the file with undecided items",
      decisions(b18i) == [("orphan_story", "accepted_exclusion", "QA"), ("caption_candidate", "keep_separate", "QA")]
      and r18i[0] == "ok" and decisions(a18i) == [("orphan_story", None, ""), ("caption_candidate", None, "")],
      [r18i[:2], decisions(a18i)])

_CORRUPT = "{questo non è JSON"
p18j, d18j = draft_project("S18j corrupt JSON no work", IDML_A)
write_review(p18j, _CORRUPT)
r18j, b18j, a18j, _ = rerun("S18j_corrupt_json_no_work_then_idml_reimport", p18j, d18j, "seeded: corrupt review JSON",
                            IDML_A, "it-IT", "PUB-MAIN", "Bozza")
check("S18j: (D2) a corrupt review JSON alone is not human work: the IDML re-import -- the documented repair -- is "
      "ACCEPTED and writes a valid file again",
      [f[1] for f in b18j["files"] if f[0] == "ManualSync_IDML_Review_PUB-MAIN.json"][0].startswith("invalid JSON")
      and r18j[0] == "ok" and decisions(a18j) == [("orphan_story", None, ""), ("caption_candidate", None, "")],
      [r18j[:2], [f for f in a18j["files"] if f[0].endswith(".json")]])

p18k, d18k = draft_project("S18k corrupt JSON with edit", IDML_A)
with repository(p18k) as repo:
    act18k = outcome(AuthoringService(repo).save_block, "IDML-0001-B001", text=EDITED, expected_version=1, actor=ACTOR)
write_review(p18k, _CORRUPT)
r18k, b18k, a18k, _ = rerun("S18k_corrupt_json_with_edit_then_idml_reimport", p18k, d18k, act18k, IDML_A, "it-IT",
                            "PUB-MAIN", "Bozza")
check("S18k: (D2) ... but work saved in the database is still protected next to a corrupt JSON: REFUSED on the edit "
      "alone, with no effect (the corrupt file stays as it was)",
      act18k[0] == "ok" and refused(r18k, b18k, a18k, {"versions": 1, "audit": 1}), r18k)

# ------------------------------------------------------------ S18l: a date that cannot be read (D5)
p18l, d18l = draft_project("S18l unreadable date", PDF_A)
seed(p18l, "INSERT INTO audit_events(object_type,object_id,action,actor,detail_json,created_at) VALUES(?,?,?,?,?,?)",
     ("block", "DOC-0001-B001", "AUTHORING_SAVE", ACTOR, "{}", "data illeggibile"))
r18l, b18l, a18l, _ = rerun("S18l_unreadable_audit_date_then_pdf_reimport", p18l, d18l,
                            "seeded: AUTHORING_SAVE on DOC-0001-B001 dated 'data illeggibile'", PDF_A, "it-IT",
                            "PUB-MAIN", "Bozza")
check("S18l: (D5) an audit row whose date cannot be read counts as human work (fail-closed): REFUSED, with no effect",
      refused(r18l, b18l, a18l, {"audit": 1}), r18l)

# ------------------------------------------------------------ S19: a refusal never creates Sources/ (DLS-03)
p19, d19 = draft_project("S19 no Sources", PDF_A)
with repository(p19) as repo:
    act19 = outcome(AuthoringService(repo).save_block, "DOC-0001-B001", text=EDITED, expected_version=1, actor=ACTOR)
shutil.rmtree(p19 / "Sources")                       # a project whose Sources/ is gone (or was never created)
r19, b19, a19, _ = rerun("S19_refusal_creates_no_sources", p19, d19, act19, PDF_B, "it-IT", "PUB-MAIN", "Revisione B")
check("S19: (DLS-03) a B5 refusal in a project with no Sources/ folder does not create it",
      not (p19 / "Sources").exists() and refused(r19, b19, a19, {"versions": 1, "audit": 1})
      and not any(f[0].startswith("Sources") for f in a19["files"]), [r19, [f[0] for f in a19["files"]]])

p19b = new_project("S19b Studio publication")
with repository(p19b) as repo:                       # written by hand, never imported: a 'Draft' topic
    act19b = [outcome(repo.create_topic, "MANUAL-0001", "Capitolo scritto a mano", "Testo scritto a mano.",
                      source_lang="it-IT", status="Draft"),
              outcome(repo.create_publication, "PUB-MAIN", "Pubblicazione scritta a mano", ["MANUAL-0001"],
                      source_lang="it-IT")]
r19b, b19b, a19b, _ = rerun("S19b_status_refusal_creates_no_sources", p19b, state(p19b), act19b, PDF_A, "it-IT",
                            "PUB-MAIN", "Import")
check("S19b: (DLS-03, D3) the topic-status refusal -- a publication written by hand ('Draft' topic), never imported -- "
      "does not create Sources/ either",
      [a[0] for a in act19b] == ["ok", "ok"] and not (p19b / "Sources").exists() and refused(r19b, b19b, a19b),
      [act19b, r19b])

# ------------------------------------------------------------ direct: the two new repository methods only read
with repository(p17a) as repo:
    _changes = repo.conn.total_changes
    _work17a = repo.migration_draft_human_work("PUB-MAIN")
    _since17a = repo.migration_draft_created_at("PUB-MAIN")
    _none = (repo.migration_draft_human_work("PUB-NONE"), repo.migration_draft_created_at("PUB-NONE"))
    _read_only = repo.conn.total_changes == _changes and not repo.conn.in_transaction
with repository(p04x) as repo:
    _work04x = repo.migration_draft_human_work("PUB-MAIN")
check("S18m: migration_draft_human_work reads only (no change, no open transaction) and returns only the kinds found; "
      "{} for an untouched draft and for an unknown publication; migration_draft_created_at is the earliest topic "
      "creation (None without topics)",
      _read_only and _work17a == {"versions": 1, "audit": 1} and _work04x == {} and _none == ({}, None)
      and _since17a == min(r["created_at"] for r in a17a["db"]["topics"]), [_work17a, _work04x, _none, _since17a])

# ============================================================ M: the digests see what matters, and only that
# HARDENING-02B: M1-M6 mutate the RETIREMENT of an IDML draft by a PDF re-import -- S04x since S04 is refused.
base04 = {"draft": d04x, "action": None, "before": b04x, "outcome": r04x, "after": a04x, "then": None}
ref = digest(base04)


def mutated(fn) -> str:
    m = copy.deepcopy(base04)
    fn(m)
    return digest(m)


check("M0: the M1-M6 base is the pinned S04x observation",
      ref == _observed.get("S04x_pdf_reimport_over_idml_draft_no_translations"), "")
check("M1: a different status on one topic changes the digest",
      mutated(lambda m: m["after"]["db"]["topics"][0].update(status="Approved")) != ref, "")
check("M2: one row more or less (e.g. a block the retirement forgot) changes the digest",
      mutated(lambda m: m["after"]["db"]["content_blocks"].pop()) != ref, "")
check("M3: a different relation (publication_topics position) changes the digest",
      mutated(lambda m: m["after"]["db"]["publication_topics"][0].update(position=99)) != ref, "")
check("M4: a different text changes the digest (content is never normalized)",
      mutated(lambda m: m["after"]["db"]["content_blocks"][0].update(text_content="altro")) != ref, "")
check("M5: a leftover file or a changed review item changes the digest",
      mutated(lambda m: m["after"]["files"].append(["Sources/x.pdf", "sha256:0"])) != ref, "")
check("M6: only the timestamps are normalized -- changing created_at alone leaves the digest unchanged",
      mutated(lambda m: m["after"]["db"]["topics"][0].update(created_at="2000-01-01T00:00:00+00:00")) == ref, "")
base03 = {"before": b03, "outcome": r03, "after": a03}
ref03 = digest(base03)


def _replace_everywhere(o, old: str, new: str):
    """The same observation with one random key swapped for another in EVERY place it appears."""
    if hasattr(o, "to_dict") and hasattr(o, "columns"):
        return o.map(lambda v: v.replace(old, new) if isinstance(v, str) else v)
    if isinstance(o, dict):
        return {k: _replace_everywhere(v, old, new) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return type(o)(_replace_everywhere(v, old, new) for v in o)
    return o.replace(old, new) if isinstance(o, str) else o


_key03 = [it["item_key"] for f in a03["files"] if f[0] == "ManualSync_IDML_Review_PUB-MAIN.json" for it in f[1]
          if it.get("reason") == "caption_candidate"][0]
m7a = _replace_everywhere(base03, _key03, "0123456789ab")
m7b = copy.deepcopy(base03)
[it for f in m7b["after"]["files"] if f[0] == "ManualSync_IDML_Review_PUB-MAIN.json" for it in f[1]
 if it.get("reason") == "caption_candidate"][0]["item_key"] = "0123456789ab"   # the review item alone
check("M7: another random caption key, used consistently, leaves the digest unchanged (it is run-dependent)",
      digest(m7a) == ref03, "")
check("M7b: ... but a review item that no longer points at its figure changes the digest (the link is pinned)",
      digest(m7b) != ref03, "")

base16 = {"before": b16a, "first": {"outcome": r16a, "after": a16a}, "reimport": {"outcome": r16b, "after": a16b}}
m8 = copy.deepcopy(base16)
m8["reimport"]["after"]["db"]["publications"][0].update(title="Titolo", source_lang="it-IT")
check("M8: a re-import that stripped the title and canonicalized the language like the first import "
      "(the two paths 'unified') changes the S16 digest", digest(m8) != digest(base16), "")

base17c = {"draft": d17c, "action": act17c, "before": b17c, "outcome": r17c, "after": a17c, "then": None}
ref17c = digest(base17c)
m9a = _replace_everywhere(base17c, _new17c, "BLK-0123456789")
m9b = copy.deepcopy(base17c)
[r for r in m9b["after"]["db"]["audit_events"] if r["object_type"] == "block"][0]["object_id"] = "BLK-0123456789"
check("M9: another random Studio id (BLK-<10 hex>), used consistently, leaves the S17c digest unchanged "
      "(it is run-dependent)", ref17c == _observed.get("S17c_studio_block_then_pdf_reimport") and _new17c != "BLK-0123456789"
      and digest(m9a) == ref17c, "")
check("M9b: ... but an audit row that no longer names the block it was written for changes the digest",
      digest(m9b) != ref17c, "")
base17a = {"draft": d17a, "action": act17a, "before": b17a, "outcome": r17a, "after": a17a, "then": None}
m10 = copy.deepcopy(base17a)
b001(m10["after"])[0].update(version=1, text_content=SRC_B001)
check("M10: a re-import that DELETED the Studio edit again (block back at v1 with the source text, as until 9ae4989) "
      "changes the S17a digest",
      digest(base17a) == _observed.get("S17a_block_edit_then_pdf_reimport") and digest(m10) != digest(base17a), "")
base17k = {"draft": d17k, "action": act17k, "before": b17k, "outcome": r17k, "after": a17k, "then": j17k}
m11 = copy.deepcopy(base17k)
m11["then"][0]["is_current"] = False
check("M11: evidence judged NOT current after the re-import changes the S17k digest",
      digest(base17k) == _observed.get("S17k_compliance_evidence_then_pdf_reimport") and digest(m11) != digest(base17k), "")

# ============================================================ R (CP5.2b-1): the SQL lives in IndustrialRepository
import ast  # noqa: E402

_SQL_CALLS = {"execute", "executemany", "executescript", "read_sql_query", "read_sql", "read_sql_table"}
# The 11 statements of app.import_manual_ui at f5d9f48 (before the move), as source text, in order.
ORIGINAL_SQL = [
    ("execute", '"SELECT 1 FROM publications WHERE publication_id=?"'),
    ("execute", '"SELECT t.topic_id,t.status FROM topics t JOIN publication_topics pt ON pt.topic_id=t.topic_id WHERE pt.publication_id=?"'),
    ("execute", 'f"SELECT object_ref_id FROM content_blocks WHERE topic_id IN ({marks}) AND block_type=\'table\' AND object_ref_id<>\'\'"'),
    ("execute", 'f"SELECT object_ref_id FROM content_blocks WHERE topic_id IN ({marks}) AND block_type=\'figure\' AND object_ref_id<>\'\'"'),
    ("execute", '"DELETE FROM publication_topics WHERE publication_id=?"'),
    ("executemany", '"DELETE FROM topics WHERE topic_id=?"'),
    ("executemany", '"DELETE FROM structured_tables WHERE table_id=?"'),
    ("executemany", '"DELETE FROM figures WHERE figure_id=?"'),
    ("execute", '"SELECT 1 FROM publications WHERE publication_id=?"'),
    ("execute", '"UPDATE publications SET title=?,source_lang=? WHERE publication_id=?"'),
    ("execute", '"SELECT COUNT(*) FROM structured_tables"'),
]
REPO_METHODS = ["migration_draft_objects", "retire_migration_draft", "publication_exists",
                "update_publication_header", "structured_table_count"]          # in the order of ORIGINAL_SQL
APP_SRC = (BASE / "app.py").read_text(encoding="utf-8")
CCMS_SRC = (BASE / "ccms_industrial.py").read_text(encoding="utf-8")


def _function(src: str, name: str, cls: str | None = None) -> ast.FunctionDef:
    body = ast.parse(src).body
    if cls:
        body = [n for n in body if isinstance(n, ast.ClassDef) and n.name == cls][0].body
    return [n for n in body if isinstance(n, ast.FunctionDef) and n.name == name][0]


def sql_calls(fn: ast.FunctionDef, src: str) -> list[tuple[str, str]]:
    calls = [n for n in ast.walk(fn) if isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute)
             and n.func.attr in _SQL_CALLS and n.args]
    return [(n.func.attr, ast.get_source_segment(src, n.args[0])) for n in sorted(calls, key=lambda n: (n.lineno, n.col_offset))]


def commits(fn: ast.FunctionDef) -> int:
    return sum(1 for n in ast.walk(fn) if isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute) and n.func.attr == "commit")


IMU = _function(APP_SRC, "import_manual_ui")
check("R1: import_manual_ui issues no SQL itself any more (no execute*/read_sql* call in its body)",
      sql_calls(IMU, APP_SRC) == [] and commits(IMU) == 0, [sql_calls(IMU, APP_SRC), commits(IMU)])
_repo_calls = {n.func.attr for n in ast.walk(IMU) if isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute)
               and isinstance(n.func.value, ast.Name) and n.func.value.id == "repo"}
check("R2: the five IndustrialRepository methods exist and import_manual_ui calls each of them",
      all(callable(getattr(IndustrialRepository, m, None)) for m in REPO_METHODS) and set(REPO_METHODS) <= _repo_calls,
      sorted(_repo_calls))
_imu_src = ast.get_source_segment(APP_SRC, IMU)
_retire_calls = [m.start() for m in re.finditer(r"(?<!def )_retire_old_draft\(\)", _imu_src)]   # calls, not the def
check("R2b: the draft is still retired at the same two points -- before the PDF/DOCX/DOC parse, after the IDML import",
      len(_retire_calls) == 2
      and _retire_calls[0] < _imu_src.index("if suffix in ('.docx','.doc'):") < _imu_src.index("imported=import_idml_as_topics(")
      < _retire_calls[1] < _imu_src.index("shutil.copy2(pdf,copied)\n            # Always replace the review state"),
      _retire_calls)
_moved = [c for m in REPO_METHODS for c in sql_calls(_function(CCMS_SRC, m, "IndustrialRepository"), CCMS_SRC)]
check("R3: the five methods hold exactly the 11 original statements, same text, same call kind, same order",
      _moved == ORIGINAL_SQL, _moved)
_commits = {m: commits(_function(CCMS_SRC, m, "IndustrialRepository")) for m in REPO_METHODS}
check("R3b: the two commits are where they were -- one after the retirement, one after the publication UPDATE",
      _commits == {"migration_draft_objects": 0, "retire_migration_draft": 1, "publication_exists": 0,
                   "update_publication_header": 1, "structured_table_count": 0}, _commits)
check("R3c: the refusal is still the same ValueError text",
      "raise ValueError(\"La pubblicazione contiene contenuti già revisionati/approvati. Usa un nuovo Publication ID "
      "per una nuova migrazione.\")" in ast.get_source_segment(CCMS_SRC, _function(CCMS_SRC, "migration_draft_objects",
                                                                                  "IndustrialRepository")), "")
_anchor = "    repo=IndustrialRepository(_ccms_db(project_dir))\n"
_mut = _imu_src.replace(_anchor, _anchor + "    repo.conn.execute(\"DELETE FROM topics\")\n") if _imu_src.count(_anchor) == 1 else ""
check("R4: a SQL call put back into import_manual_ui is caught by R1",
      bool(_mut) and sql_calls(ast.parse(_mut).body[0], _mut) == [("execute", '"DELETE FROM topics"')], bool(_mut))

# ============================================================ R5/R6 (HARDENING-02B): order, read-only guard, frozen CP5.2b-1
_COPY = "shutil.copy2(pdf,copied)"


def guard_order(src: str) -> list[int]:
    """Positions, in import_manual_ui, of: the topic-status guard, the human-work guard, its refusal, Sources/, the
    PDF/DOCX/DOC copy, the first retirement, the PDF/DOCX/DOC parse, the IDML import -- they must ascend."""
    retire = [m.start() for m in re.finditer(r"(?<!def )_retire_old_draft\(\)", src)]
    return [src.find("repo.migration_draft_objects(pub)"), src.find("repo.migration_draft_human_work(pub)"),
            src.find("raise gr.Error(_reimport_refused_message(work))"), src.find("srcdir.mkdir(exist_ok=True)"),
            src.find(_COPY), retire[0] if retire else -1, src.find("if suffix in ('.docx','.doc'):"),
            src.find("imported=import_idml_as_topics(")]


def order_ok(src: str) -> bool:
    pos = guard_order(src)
    return -1 not in pos and pos == sorted(pos) and src.count("srcdir.mkdir(") == 1 and src.count(_COPY) == 2


check("R5: import_manual_ui runs GUARD (topic status, human work, refusal) -> Sources/ -> COPY -> RETIRE -> PARSE, in "
      "that order in its source; Sources/ is created once, the upload copied at the two expected places",
      order_ok(_imu_src), guard_order(_imu_src))
_copy_line = "            " + _COPY + "\n"                 # the PDF/DOCX/DOC copy comes first, the IDML one second
_mut5 = (_imu_src.replace(_copy_line, "", 1).replace(_anchor, _anchor + "    " + _COPY + "\n", 1)
         if _imu_src.count(_copy_line) == 2 else "")
check("R5b: a copy moved back before the guard (the 9ae4989 order) is caught by R5",
      bool(_mut5) and not order_ok(_mut5) and _mut5.count(_COPY) == 2, bool(_mut5))
# SHA-256 of the source of each CP5.2b-1 method at 9ae4989 (git show 9ae4989:.../ccms_industrial.py).
FROZEN_CP52B1 = {
    "migration_draft_objects": "83fb292d3bd06d3fbdc46d5eb8689e9fc0c3eb50b13ab77fd23539a59a719676",
    "retire_migration_draft": "d5e50cb84b0816a281405b90d26d0adc47cddcdae1957e36cf12fd5aa489e19d",
    "publication_exists": "ec04a6d5072f6fada68b0903b4719691ebeefb18b9647e6fd5efd349352aa07b",
    "update_publication_header": "e061d173880857aa73931abf1e10d18a337d66dc1c168ce144c29877a6c39247",
    "structured_table_count": "46b85bb43ee5da6d0d6a2087572e4641fa881d2b404a4ab0b066d351493d66a0",
}
_frozen_now = {m: hashlib.sha256(ast.get_source_segment(CCMS_SRC, _function(CCMS_SRC, m, "IndustrialRepository"))
                                 .encode("utf-8")).hexdigest() for m in REPO_METHODS}
check("R6: the five CP5.2b-1 methods are byte-identical to 9ae4989 (SHA-256 of their source: docstrings, comments, "
      "SQL, commits)", _frozen_now == FROZEN_CP52B1, {m: h for m, h in _frozen_now.items() if FROZEN_CP52B1.get(m) != h})
_B5_METHODS = ("migration_draft_human_work", "migration_draft_created_at")
_WRITE_SQL = re.compile(r"\b(INSERT|UPDATE|DELETE|REPLACE|CREATE|DROP|ALTER|PRAGMA)\b", re.I)
_b5_src = [ast.get_source_segment(CCMS_SRC, _function(CCMS_SRC, m, "IndustrialRepository")) for m in _B5_METHODS]
check("R6b: the two B5 methods (and the draft query they share) issue no write and no commit; import_manual_ui raises "
      "the refusal itself, the repository never imports gradio",
      not any(_WRITE_SQL.search(s) for s in [*_b5_src, IndustrialRepository._MIGRATION_DRAFT])
      and all(commits(_function(CCMS_SRC, m, "IndustrialRepository")) == 0 for m in _B5_METHODS)
      and not re.search(r"^\s*(import gradio|from gradio)", CCMS_SRC, re.M), "")

check("N1: no network access attempted", not GUARD.snapshot(), GUARD.snapshot())

# Pinned at HEAD d4f7c45 (CP5.2a) after three identical runs; S16 added at 02652e5 (CP5.2b-0), same procedure;
# S04x/S17 added on f9eb474 (HARDENING-02A, B5 characterization), same procedure, the 21 digests above unchanged.
# HARDENING-02B (B5) on 9ae4989, same procedure: BEHAVIOR INTENTIONALLY CHANGED for S04, S07, S07b, S07c, S17a, S17a2,
# S17b, S17c, S17d, S17d2, S17e, S17g, S17g2, S17h, S17h2, S17i, S17j, S17k (refused before any effect; D3 gr.Error;
# DLS-03 no copy); S18a-S19b added; the other 22 digests byte-identical to 9ae4989.
GOLDEN = {
    'S01_pdf_new': '0a6cefe0be16b1b58a52e36d528b3ccb3969a562eff9cec2284a289de98a972b',
    'S02_docx_new': '1d19b535de0b302151cf915b985875d1f6768e2937bc0b83a4f68bc876d9b54a',
    'S03_idml_new': 'f652f258996ab1ad9b681351d35d8ae9f729125344a83d271105e8ab8c214a4f',
    'S04_pdf_reimport_over_idml_draft': 'bec7c263334b6f81e4d7ce50cd092f1622b24109ea3d47f8d9b96223df245308',
    'S05_docx_reimport_title_and_language': 'b01b967d636ed275d5a2b44834563b60751c7f474d0e75d160543863460af7f2',
    'S06_idml_reimport_over_idml': 'aadad84354c2eeb341a63284149bbee7d371094a70914329fd65a692091cccc2',
    'S07_refused_approved': 'a3fade7672cc182ec305adabbc88443548937a8e49847c0393cfea5efd6087a7',
    'S07b_refused_draft': '4fa41fbefd8b7cef9c7e1b6cf099cb9ca945f1579e4c9ec14c527449f0bc53d2',
    'S07c_refused_idml': '67aba658668e028cc178d9646c205ff5fdccd72d9b689193845c1fed412be343',
    'S08_corrupt_pdf_after_retire': '4325345e1b0c174c4331d2b6f1039c2bad095fb8cf315259b66d4eaf9e4f87e5',
    'S08c_good_pdf_after_data_loss': '8e885fc43dc4e212e7cb6445484b81452975cb49913438f77fdb20f03a34b864',
    'S09_empty_pdf_after_retire': 'c8e396683cfd653cf1c7b417cfb0c40c5d68925da4a443c2476324c2968a8485',
    'S10_corrupt_docx_after_retire': 'f485916d53cec81132a7246c547ed5eef0a25b4ae3b24304def3030ae02dc3e3',
    'S11_corrupt_idml_keeps_draft': 'ed5d15d735ac980e4da25b51fdd71802632fd4d2642438abd18fb0d9340a74e3',
    'S12_doc_converted': 'c978c52ea464560005ce950f215e4122fa283255a7f6ed66bebe8ed881ee4ae6',
    'S12b_doc_converter_failure': '3a67d63dbcf5db30f1998bcbf8454fb85f9e18b504168adc348fed23909ec306',
    'S13_guard': '0c2e51718b0b94c63563ee6250ffee880a2e972af4a1dd13a244052be5e24e49',
    'S14_defaults': '841893688e8f78f75cad7a9c67e8d8974d769e2809c0b18e6ba100ef3954aea6',
    'S14b_padded_publication_id': '54a14ba292b3ce8fe48f3193adf14e5e5c6ad7f8ca23440bb33168c8d46dc1b4',
    'S15_guided_onboarding': 'e0a6f69c73a601c59dc42ff64f72abe91e44fb33db329877b3cb50aafb114f0c',
    'S16_publication_title_language_asymmetry': '0d33da9ba132522f1aaf3ce96af62967b8c4a237e78899fd605b1689166c04ba',
    'S04x_pdf_reimport_over_idml_draft_no_translations': '4490ce2535a1753d184526e5c6aaa031e5fff6be81263fa3063127510ff86d28',
    'S17a2_block_edit_then_docx_reimport': '9532915ab0e68bc6bfd32fe26a08e5b2668666d020ed0ea9906aae020bc09db7',
    'S17a_block_edit_then_pdf_reimport': '724ca6b47eb598d72093953bf672513bc55ea61a0f68c2df56771bee27b2a67a',
    'S17b_title_edit_then_pdf_reimport': '6c1178e8f3cdde373635874ef671dfdb50e8a58b5833e71baa1dca6d39449aeb',
    'S17c_studio_block_then_pdf_reimport': 'b6ca02f6bc0ead12e7841766dca60aa09870c927d213ff8bcf00c091639b35da',
    'S17d2_review_block_then_idml_reimport': '4b9e7e71f9eb740ed0418759b886b36ab0affc39eeb535dd75b5da3941ad71b6',
    'S17d_review_block_then_pdf_reimport': '7c6874c9fe69bb1000750c65838dcb31fd802e08b7192139e9064092b30bdca3',
    'S17e_archived_topic_then_pdf_reimport': 'e75fa255ab8aa964d268e2662429b1ad4cd2f5a3e3f7416e982c6f813c6d583f',
    'S17f_autosave_no_change_then_pdf_reimport': 'b6cdaaf160ef3fb3268d37c3dd503f9f8fee08ec8caa8878db474e56c32a7b5b',
    'S17g2_idml_json_only_decisions_then_idml_reimport': 'cf872522190de276ba13d46faebf7464a39aed140ab6cef003aa9a5f2564649f',
    'S17g_idml_caption_confirmed_then_idml_reimport': '7c54c849ec53773d99f21197c6e33549a5102bc74e0a47819e5d23b929dc7143',
    'S17h2_guided_blank_name_reimport': '9662036b1c12fcb5fa0166b67908ebba178bebf692a690894c789b672221249e',
    'S17h_guided_same_name_reimport': '79a15053bde49f81f5ce81cd52cb39993becf37861091eaf8b6cb1eebf2ade6e',
    'S17i_same_name_file_refused_overwrites_source': '135b1c118d2057125cf6b9769e86bfa3463e1d137a50f87a70070d8ee402fa6a',
    'S17j_audit_before_draft_then_pdf_reimport': 'f09d46bb26a8d83314fce0516a944110c32797f947e094ca202d4e4419dca76e',
    'S17j2_seeded_audit_before_draft_then_pdf_reimport': '4870e41777e5cde5f8a2f89c3834ceb68f2e747d5c9fc33f971f21afb9c4be6c',
    'S17jbis2_orphan_autosave_rebound_by_next_idml_reimport': '7acf8e1dd31eb69152f108804e9a8d48ec0193b0828e4b6ecf74a5246ec3ee8e',
    'S17jbis_orphan_autosave_after_idml_reimport': 'cb7e0ed8708f9c4e606dab64b2d9ac1dc0b945ad480552ab0355967cdd5105c1',
    'S17k_compliance_evidence_then_pdf_reimport': '8cc4dc4217e8dd68e3dfad1e21ffe7339d72f6d311713e58c8bda6cd1f92561b',
    'S18a_reorder_blocks_then_pdf_reimport': '4df13c33a3a07af0037c1865859fdbd52086fbd56e380bbbcda1eff9837e5135',
    'S18b_reorder_topics_then_pdf_reimport': 'a27bc4da6fbb1bcf30088e20cf0b30b0410a4c2f00859e30a3730b91745045f2',
    'S18c_archive_restore_then_pdf_reimport': 'c0e23585520e2db2fcc9057b3114fa9d4b609555c76269f3c28da698a3879b88',
    'S18d_seeded_stale_review_then_pdf_reimport': '03d11b9794f08764380da29b174a463f92eb7b82039e768281e27035a8723a58',
    'S18e_seeded_fresh_review_then_pdf_reimport': '18d92e8d180db9e22909a117ab229edb61ffe76c869c044cb3394c50ee516e73',
    'S18f_seeded_stale_evidence_then_pdf_reimport': 'dcb65e5d01788d1b5536afc7354b60a6e3d8035e188a43c6c483485fed174f8c',
    'S18g_stale_evidence_new_link_then_pdf_reimport': '6510d68bbab98c5faeb5da8a7a4e1b9a087281f3ecfbd5692e99a9de99015b70',
    'S18h_idml_json_decisions_then_pdf_reimport': 'ae11e450254488fefd594763dc6b60698bf5626633c32586db542bf81b000fff',
    'S18i_seeded_stale_json_decisions_then_idml_reimport': 'ae8d079cab58de2c61225c90cce2fe3a719a53aa0fd6cddf24ef1850e460ea59',
    'S18j_corrupt_json_no_work_then_idml_reimport': 'd4963af5ff426bc00b2cbb407f4eccb29dd12ed7b3a17079386545a912269581',
    'S18k_corrupt_json_with_edit_then_idml_reimport': 'b74f4f424cde666b74223a86adb7b509e6c093919f93645a8b6a5859c27c4270',
    'S18l_unreadable_audit_date_then_pdf_reimport': '98580f8351a839a579134dc7421c29c486a464814e4fd6b27a349bb30e2e12ff',
    'S19_refusal_creates_no_sources': '4d065a496b8c303be93712cc331189d07939c8ef274051d11b01188ccbb556a6',
    'S19b_status_refusal_creates_no_sources': '2f916e5ffd3a567b845f9d11394001ecbc5a82996e0686219a0ac989798c3024',
}

if PRINT_GOLDEN:
    print("GOLDEN = {")
    for k in sorted(_observed):
        print(f"    {k!r}: {_observed[k]!r},")
    print("}")
    shutil.rmtree(TMP, ignore_errors=True)
    sys.exit(0)

for label in sorted(_observed):
    check(f"SNAP {label}: identical to the CP5.2a baseline", GOLDEN.get(label) == _observed[label],
          f"expected {GOLDEN.get(label)} got {_observed[label]}")
check("SNAP: every pinned output was observed", set(GOLDEN) == set(_observed), sorted(set(GOLDEN) ^ set(_observed)))

shutil.rmtree(TMP, ignore_errors=True)
check("T1: the temporary fixtures were removed", not TMP.exists(), TMP)

passed = sum(1 for _, ok, _ in _results if ok)
total = len(_results)
(BASE / "RISULTATO_IMPORT_ORCHESTRATION.json").write_text(json.dumps({
    "suite": "import orchestration (CP5.2a)", "tests": total, "pass": passed, "fail": total - passed, "digests": _observed,
    "results": [{"test": n, "status": "PASS" if ok else "FAIL", "detail": d} for n, ok, d in _results],
}, indent=2, ensure_ascii=False), encoding="utf-8")
print(f"\n  {passed}/{total} passed")
sys.exit(0 if passed == total else 1)
