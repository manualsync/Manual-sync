from __future__ import annotations

import hashlib
import json
import shutil
import sqlite3
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document

from ccms_industrial import IndustrialRepository
from translation_work_package import generate_translation_work_packages
from translation_roundtrip import commit_translation_roundtrip
from release_readiness import assess_release_readiness
from release_manager import create_release_bundle

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'regressione_desktop_r12_out'
if OUT.exists():
    shutil.rmtree(OUT)
OUT.mkdir(parents=True)

results=[]
def check(name, cond, detail=''):
    results.append({'name':name,'status':'PASS' if cond else 'FAIL','detail':str(detail)})
    if not cond:
        raise AssertionError(f'{name}: {detail}')

def sha(path: Path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

# 1) Backward-compatible schema migration for existing project DBs.
legacy = OUT/'legacy_schema.sqlite'
conn=sqlite3.connect(legacy)
conn.execute('''CREATE TABLE topic_translations (
    topic_id TEXT NOT NULL,
    target_lang TEXT NOT NULL,
    source_version INTEGER NOT NULL,
    target_body TEXT NOT NULL,
    status TEXT NOT NULL DEFAULT 'Draft',
    updated_at TEXT NOT NULL,
    PRIMARY KEY(topic_id,target_lang)
)''')
conn.commit(); conn.close()
repo=IndustrialRepository(legacy)
cols={r[1] for r in repo.conn.execute('PRAGMA table_info(topic_translations)')}
check('Legacy schema migrates target_title', 'target_title' in cols, sorted(cols))
repo.close()

# 2) Build a structured publication.
project=OUT/'project'; project.mkdir()
db=project/'ManualSync_CCMS.sqlite'
repo=IndustrialRepository(db)
repo.create_topic('T1','Electrical ratings','Legacy body.',source_lang='en-US',status='Approved',source_refs=['Spec#1'])
repo.create_publication('PUB1','Demo Drive Manual',['T1'],source_lang='en-US')
repo.add_block('B1','T1','paragraph',position=1,text='Install the drive vertically.',status='Approved',source_refs=['Spec#1'])
# Simulate an old R11-approved block translation with no localized topic title.
repo.set_block_translation('B1','it-IT','Installare il convertitore verticalmente.',status='Reviewed')
repo.close()

pre=assess_release_readiness(project,'PUB1',target_langs=['it-IT'])
check('Readiness blocks untranslated topic title', not pre['release_ready'], pre['matrix'].to_dict(orient='records'))
check('Topic-title blocker is explicit', any(str(x.get('block_type'))=='topic_title' for x in pre['blocking_items'].to_dict(orient='records')), pre['blocking_items'].to_dict(orient='records'))

# 3) Translation package must include topic title as first-class source unit.
repo=IndustrialRepository(db)
wp=generate_translation_work_packages(repo,'PUB1',kind='publication',target_langs=['it-IT'],out_dir=OUT/'workpkg')
repo.close()
langzip=Path(wp['language_zips'][0])
with zipfile.ZipFile(langzip) as z:
    xlf_name=[n for n in z.namelist() if n.endswith('_WORK_ONLY.xlf')][0]
    raw=z.read(xlf_name)
xlf=OUT/'returned_it.xlf'; xlf.write_bytes(raw)
ns={'x':'urn:oasis:names:tc:xliff:document:2.0'}
tree=ET.parse(xlf); root=tree.getroot()
translations={
    'Electrical ratings':'Valori nominali elettrici',
    'Install the drive vertically.':'Installare il convertitore verticalmente.',
}
seen=[]
for unit in root.findall('.//x:unit',ns):
    seg=unit.find('./x:segment',ns); src=seg.find('./x:source',ns); tgt=seg.find('./x:target',ns)
    source=''.join(src.itertext()) if src is not None else ''
    seen.append(source)
    target=translations.get(source, source)
    tgt.clear(); tgt.text=target; seg.set('state','translated')
tree.write(xlf,encoding='utf-8',xml_declaration=True)
check('Work package includes topic title', 'Electrical ratings' in seen, seen)

# 4) Round-trip persists title translation and block translation.
repo=IndustrialRepository(db)
commit=commit_translation_roundtrip(repo,project,langzip,xlf,human_approved=True,reviewer='Regression R12',out_dir=OUT/'roundtrip')
tr=repo.conn.execute("SELECT target_title,target_body,status FROM topic_translations WHERE topic_id='T1' AND target_lang='it-IT'").fetchone()
check('Topic title persisted after approved round-trip', tr and tr['target_title']=='Valori nominali elettrici' and tr['status']=='Reviewed', dict(tr) if tr else None)
frame,summary=repo.assemble_structured_publication('PUB1',target_lang='it-IT')
check('Localized title appears in assembled publication', frame.iloc[0]['topic_title']=='Valori nominali elettrici', frame[['topic_title','topic_title_translation_status']].to_dict(orient='records'))
repo.close()

ready=assess_release_readiness(project,'PUB1',target_langs=['it-IT'])
check('Readiness passes after title + content approval', ready['release_ready'], ready['matrix'].to_dict(orient='records'))

# 5) Controlled release package hardening.
rel=create_release_bundle(project,'PUB1',target_langs=['it-IT'],released_by='Regression R12',release_id='REL-R12',out_dir=OUT/'release')
zip_path=Path(rel['release_zip'])
check('Release ZIP created', zip_path.exists() and zip_path.stat().st_size>0, zip_path)
with zipfile.ZipFile(zip_path) as z:
    bad=z.testzip(); names=set(z.namelist())
    check('Release ZIP CRC valid', bad is None, bad)
    manifest=json.loads(z.read('RELEASE_MANIFEST.json'))
    check('Manifest artifact paths are archive-relative', all(a['path'].startswith('ARTIFACTS/') and a['path'] in names for a in manifest['artifacts']), manifest['artifacts'])
    check('Immutable CCMS snapshots included', all(r['snapshot_path'] in names for r in manifest['ccms_releases']), manifest['ccms_releases'])
    check('Per-language provenance included', all(r['provenance_path'] in names for r in manifest['ccms_releases']), manifest['ccms_releases'])
    # snapshot file bytes must hash to the release snapshot hash
    for r in manifest['ccms_releases']:
        check(f"Snapshot checksum matches {r['language']}", hashlib.sha256(z.read(r['snapshot_path'])).hexdigest()==r['sha256'], r)
    for a in manifest['artifacts']:
        check(f"Artifact checksum matches {a['language']}", hashlib.sha256(z.read(a['path'])).hexdigest()==a['sha256'], a)
    for a in manifest['audit_files']:
        check(f"Audit checksum matches {a['path']}", hashlib.sha256(z.read(a['path'])).hexdigest()==a['sha256'], a)

it_doc=OUT/'release'/'ARTIFACTS'/'PUB1_it-IT_REL-R12.docx'
doc=Document(it_doc)
headings=[p.text for p in doc.paragraphs if p.style and p.style.name.startswith('Heading')]
all_text='\n'.join(p.text for p in doc.paragraphs)
check('Italian DOCX heading is localized', 'Valori nominali elettrici' in headings and 'Electrical ratings' not in headings, headings)
check('Customer DOCX hides source audit lines', 'Source:' not in all_text, all_text[:1000])
check('Customer DOCX hides internal engine build', '0.18.0-word-idml-template-bridge' not in all_text and 'desktop-r12' not in all_text, all_text[:500])
check('DOCX metadata is commercialized', doc.core_properties.author=='ManualSync' and doc.core_properties.comments=='', {'author':doc.core_properties.author,'comments':doc.core_properties.comments,'created':str(doc.core_properties.created)})

summary={
    'build':'1.0.0-desktop-r12-translation-completeness-release-hardening-20260817',
    'tests':len(results),
    'pass':sum(r['status']=='PASS' for r in results),
    'fail':sum(r['status']=='FAIL' for r in results),
    'results':results,
}
(ROOT/'RISULTATO_REGRESSIONE_DESKTOP_R12.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
(ROOT/'RISULTATO_REGRESSIONE_DESKTOP_R12.txt').write_text('\n'.join(f"{r['status']} | {r['name']} | {r['detail']}" for r in results)+f"\n\nTOTAL {summary['pass']}/{summary['tests']} PASS\n",encoding='utf-8')
print(json.dumps({'tests':summary['tests'],'pass':summary['pass'],'fail':summary['fail']},indent=2))
