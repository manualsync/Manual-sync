from __future__ import annotations

import html
import json
import re
import shutil
import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Any
from uuid import uuid4

import pandas as pd
from docx import Document as DocxDocument

from layout_engine import extract_layout
from language_codes import canonical_lang
from idml_import import parse_idml, materialize_figure_assets

AUTHORING_BUILD = "0.17.0-structured-authoring-ccms-foundation-20260810"

COMPONENT_TOKEN_RE = re.compile(r"\{\{component:([A-Za-z0-9_.:-]+)\}\}")


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def _load_json(value: str | None, default):
    try:
        return json.loads(value or "")
    except Exception:
        return default


class ContentRepository:
    """Local structured-content repository.

    This is intentionally a foundation, not a full DITA/DocBook implementation.
    It establishes the core CCMS invariants ManualSync needs before adding richer
    schemas and publishing bridges: reusable topics/components, immutable version
    history, variant conditions, translations, publication assemblies and impact
    tracing.
    """

    def __init__(self, path: str | Path):
        self.path = Path(path)
        self.conn = sqlite3.connect(self.path)
        self.conn.row_factory = sqlite3.Row
        self.conn.execute("PRAGMA foreign_keys=ON")
        self.conn.execute("PRAGMA journal_mode=WAL")
        self._schema()

    def _schema(self) -> None:
        self.conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS topics (
                topic_id TEXT PRIMARY KEY,
                title TEXT NOT NULL,
                body TEXT NOT NULL,
                source_lang TEXT NOT NULL DEFAULT 'en-US',
                status TEXT NOT NULL DEFAULT 'Draft',
                version INTEGER NOT NULL DEFAULT 1,
                include_variants_json TEXT NOT NULL DEFAULT '[]',
                exclude_variants_json TEXT NOT NULL DEFAULT '[]',
                source_refs_json TEXT NOT NULL DEFAULT '[]',
                evidence_status TEXT NOT NULL DEFAULT 'SOURCE_GROUNDED',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS topic_versions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                topic_id TEXT NOT NULL,
                version INTEGER NOT NULL,
                title TEXT NOT NULL,
                body TEXT NOT NULL,
                status TEXT NOT NULL,
                include_variants_json TEXT NOT NULL,
                exclude_variants_json TEXT NOT NULL,
                source_refs_json TEXT NOT NULL,
                evidence_status TEXT NOT NULL,
                created_at TEXT NOT NULL,
                UNIQUE(topic_id, version),
                FOREIGN KEY(topic_id) REFERENCES topics(topic_id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS components (
                component_id TEXT PRIMARY KEY,
                title TEXT NOT NULL,
                body TEXT NOT NULL,
                source_lang TEXT NOT NULL DEFAULT 'en-US',
                version INTEGER NOT NULL DEFAULT 1,
                source_refs_json TEXT NOT NULL DEFAULT '[]',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS component_versions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                component_id TEXT NOT NULL,
                version INTEGER NOT NULL,
                title TEXT NOT NULL,
                body TEXT NOT NULL,
                source_refs_json TEXT NOT NULL,
                created_at TEXT NOT NULL,
                UNIQUE(component_id, version),
                FOREIGN KEY(component_id) REFERENCES components(component_id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS topic_translations (
                topic_id TEXT NOT NULL,
                target_lang TEXT NOT NULL,
                source_version INTEGER NOT NULL,
                target_title TEXT NOT NULL DEFAULT '',
                target_body TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'Draft',
                updated_at TEXT NOT NULL,
                PRIMARY KEY(topic_id, target_lang),
                FOREIGN KEY(topic_id) REFERENCES topics(topic_id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS publications (
                publication_id TEXT PRIMARY KEY,
                title TEXT NOT NULL,
                source_lang TEXT NOT NULL DEFAULT 'en-US',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS publication_topics (
                publication_id TEXT NOT NULL,
                topic_id TEXT NOT NULL,
                position INTEGER NOT NULL,
                PRIMARY KEY(publication_id, topic_id),
                FOREIGN KEY(publication_id) REFERENCES publications(publication_id) ON DELETE CASCADE,
                FOREIGN KEY(topic_id) REFERENCES topics(topic_id) ON DELETE CASCADE
            );
            CREATE INDEX IF NOT EXISTS ix_pubtopic_topic ON publication_topics(topic_id);
            CREATE INDEX IF NOT EXISTS ix_translation_lang ON topic_translations(target_lang,status);
            """
        )
        # R12 schema migration: older projects have topic_translations without target_title.
        topic_translation_cols = {str(r[1]) for r in self.conn.execute("PRAGMA table_info(topic_translations)")}
        if "target_title" not in topic_translation_cols:
            self.conn.execute("ALTER TABLE topic_translations ADD COLUMN target_title TEXT NOT NULL DEFAULT ''")
        # 4B.1 schema migration: archived_at is a separate axis from the editorial
        # `status` lifecycle (Draft/In Review/Approved) -- archiving a topic must
        # never be confused with a review/approval state. NULL = active (the
        # default/common case for every pre-existing row).
        topic_cols = {str(r[1]) for r in self.conn.execute("PRAGMA table_info(topics)")}
        if "archived_at" not in topic_cols:
            self.conn.execute("ALTER TABLE topics ADD COLUMN archived_at TEXT")
        self.conn.commit()

    def close(self) -> None:
        self.conn.close()

    def create_topic(self, topic_id: str, title: str, body: str, *, source_lang: str = "en-US",
                     status: str = "Draft", include_variants: Iterable[str] | None = None,
                     exclude_variants: Iterable[str] | None = None, source_refs: Iterable[str] | None = None,
                     evidence_status: str = "SOURCE_GROUNDED") -> str:
        topic_id = str(topic_id).strip()
        if not topic_id:
            raise ValueError("topic_id obbligatorio")
        now = _now()
        include = sorted(set(str(x).strip() for x in (include_variants or []) if str(x).strip()))
        exclude = sorted(set(str(x).strip() for x in (exclude_variants or []) if str(x).strip()))
        refs = sorted(set(str(x).strip() for x in (source_refs or []) if str(x).strip()))
        self.conn.execute(
            """INSERT INTO topics(topic_id,title,body,source_lang,status,version,include_variants_json,exclude_variants_json,source_refs_json,evidence_status,created_at,updated_at)
               VALUES(?,?,?,?,?,1,?,?,?,?,?,?)""",
            (topic_id, title.strip(), body.strip(), canonical_lang(source_lang), status, _json(include), _json(exclude), _json(refs), evidence_status, now, now),
        )
        self.conn.execute(
            """INSERT INTO topic_versions(topic_id,version,title,body,status,include_variants_json,exclude_variants_json,source_refs_json,evidence_status,created_at)
               VALUES(?,?,?,?,?,?,?,?,?,?)""",
            (topic_id, 1, title.strip(), body.strip(), status, _json(include), _json(exclude), _json(refs), evidence_status, now),
        )
        self.conn.commit()
        return topic_id

    def update_topic(self, topic_id: str, *, title: str | None = None, body: str | None = None,
                     status: str | None = None, include_variants: Iterable[str] | None = None,
                     exclude_variants: Iterable[str] | None = None, source_refs: Iterable[str] | None = None,
                     evidence_status: str | None = None) -> int:
        row = self.conn.execute("SELECT * FROM topics WHERE topic_id=?", (topic_id,)).fetchone()
        if not row:
            raise KeyError(topic_id)
        version = int(row["version"]) + 1
        values = {
            "title": row["title"] if title is None else title.strip(),
            "body": row["body"] if body is None else body.strip(),
            "status": row["status"] if status is None else status,
            "include": _load_json(row["include_variants_json"], []) if include_variants is None else sorted(set(include_variants)),
            "exclude": _load_json(row["exclude_variants_json"], []) if exclude_variants is None else sorted(set(exclude_variants)),
            "refs": _load_json(row["source_refs_json"], []) if source_refs is None else sorted(set(source_refs)),
            "evidence": row["evidence_status"] if evidence_status is None else evidence_status,
        }
        now = _now()
        self.conn.execute(
            """UPDATE topics SET title=?,body=?,status=?,version=?,include_variants_json=?,exclude_variants_json=?,source_refs_json=?,evidence_status=?,updated_at=? WHERE topic_id=?""",
            (values["title"], values["body"], values["status"], version, _json(values["include"]), _json(values["exclude"]), _json(values["refs"]), values["evidence"], now, topic_id),
        )
        self.conn.execute(
            """INSERT INTO topic_versions(topic_id,version,title,body,status,include_variants_json,exclude_variants_json,source_refs_json,evidence_status,created_at)
               VALUES(?,?,?,?,?,?,?,?,?,?)""",
            (topic_id, version, values["title"], values["body"], values["status"], _json(values["include"]), _json(values["exclude"]), _json(values["refs"]), values["evidence"], now),
        )
        self.conn.commit()
        return version

    def create_component(self, component_id: str, title: str, body: str, *, source_lang: str = "en-US",
                         source_refs: Iterable[str] | None = None) -> str:
        now = _now()
        refs = sorted(set(str(x).strip() for x in (source_refs or []) if str(x).strip()))
        self.conn.execute(
            "INSERT INTO components(component_id,title,body,source_lang,version,source_refs_json,created_at,updated_at) VALUES(?,?,?,?,1,?,?,?)",
            (component_id, title.strip(), body.strip(), canonical_lang(source_lang), _json(refs), now, now),
        )
        self.conn.execute(
            "INSERT INTO component_versions(component_id,version,title,body,source_refs_json,created_at) VALUES(?,?,?,?,?,?)",
            (component_id, 1, title.strip(), body.strip(), _json(refs), now),
        )
        self.conn.commit()
        return component_id

    def update_component(self, component_id: str, *, title: str | None = None, body: str | None = None,
                         source_refs: Iterable[str] | None = None) -> int:
        row = self.conn.execute("SELECT * FROM components WHERE component_id=?", (component_id,)).fetchone()
        if not row:
            raise KeyError(component_id)
        version = int(row["version"]) + 1
        title_v = row["title"] if title is None else title.strip()
        body_v = row["body"] if body is None else body.strip()
        refs = _load_json(row["source_refs_json"], []) if source_refs is None else sorted(set(source_refs))
        now = _now()
        self.conn.execute("UPDATE components SET title=?,body=?,version=?,source_refs_json=?,updated_at=? WHERE component_id=?",
                          (title_v, body_v, version, _json(refs), now, component_id))
        self.conn.execute("INSERT INTO component_versions(component_id,version,title,body,source_refs_json,created_at) VALUES(?,?,?,?,?,?)",
                          (component_id, version, title_v, body_v, _json(refs), now))
        self.conn.commit()
        return version

    def set_translation(self, topic_id: str, target_lang: str, target_body: str, *, target_title: str | None = None, status: str = "Draft") -> None:
        target_lang = canonical_lang(target_lang)
        topic = self.conn.execute("SELECT version FROM topics WHERE topic_id=?", (topic_id,)).fetchone()
        if not topic:
            raise KeyError(topic_id)
        existing = self.conn.execute("SELECT target_title FROM topic_translations WHERE topic_id=? AND REPLACE(target_lang,'_','-')=? COLLATE NOCASE", (topic_id, target_lang)).fetchone()
        title_value = (str(existing["target_title"] or "") if existing and target_title is None else str(target_title or "")).strip()
        now = _now()
        # Collapse any legacy casing/separator aliases before writing the canonical tag.
        self.conn.execute("DELETE FROM topic_translations WHERE topic_id=? AND REPLACE(target_lang,'_','-')=? COLLATE NOCASE", (topic_id, target_lang))
        self.conn.execute(
            """INSERT INTO topic_translations(topic_id,target_lang,source_version,target_title,target_body,status,updated_at)
               VALUES(?,?,?,?,?,?,?)
               ON CONFLICT(topic_id,target_lang) DO UPDATE SET source_version=excluded.source_version,target_title=excluded.target_title,target_body=excluded.target_body,status=excluded.status,updated_at=excluded.updated_at""",
            (topic_id, target_lang, int(topic["version"]), title_value, target_body.strip(), status, now),
        )
        self.conn.commit()

    def set_topic_title_translation(self, topic_id: str, target_lang: str, target_title: str, *, status: str = "Draft") -> None:
        target_lang = canonical_lang(target_lang)
        topic = self.conn.execute("SELECT version FROM topics WHERE topic_id=?", (topic_id,)).fetchone()
        if not topic:
            raise KeyError(topic_id)
        existing = self.conn.execute("SELECT target_body FROM topic_translations WHERE topic_id=? AND REPLACE(target_lang,'_','-')=? COLLATE NOCASE", (topic_id, target_lang)).fetchone()
        body_value = str(existing["target_body"] or "") if existing else ""
        self.set_translation(topic_id, target_lang, body_value, target_title=target_title, status=status)

    def create_publication(self, publication_id: str, title: str, topic_ids: Iterable[str], *, source_lang: str = "en-US") -> str:
        ids = [str(x).strip() for x in topic_ids if str(x).strip()]
        if not ids:
            raise ValueError("publication senza topic")
        now = _now()
        self.conn.execute("INSERT INTO publications(publication_id,title,source_lang,created_at,updated_at) VALUES(?,?,?,?,?)",
                          (publication_id, title.strip(), canonical_lang(source_lang), now, now))
        for pos, topic_id in enumerate(ids, start=1):
            if not self.conn.execute("SELECT 1 FROM topics WHERE topic_id=?", (topic_id,)).fetchone():
                raise KeyError(topic_id)
            self.conn.execute("INSERT INTO publication_topics(publication_id,topic_id,position) VALUES(?,?,?)", (publication_id, topic_id, pos))
        self.conn.commit()
        return publication_id

    def replace_publication_topics(self, publication_id: str, topic_ids: Iterable[str]) -> None:
        ids = [str(x).strip() for x in topic_ids if str(x).strip()]
        self.conn.execute("DELETE FROM publication_topics WHERE publication_id=?", (publication_id,))
        for pos, topic_id in enumerate(ids, start=1):
            self.conn.execute("INSERT INTO publication_topics(publication_id,topic_id,position) VALUES(?,?,?)", (publication_id, topic_id, pos))
        self.conn.execute("UPDATE publications SET updated_at=? WHERE publication_id=?", (_now(), publication_id))
        self.conn.commit()

    def topic_table(self, *, include_archived: bool = False) -> pd.DataFrame:
        sql = "SELECT * FROM topics"
        if not include_archived:
            sql += " WHERE archived_at IS NULL"
        sql += " ORDER BY topic_id"
        rows = []
        for r in self.conn.execute(sql):
            rows.append({
                "topic_id": r["topic_id"], "title": r["title"], "status": r["status"], "version": r["version"],
                "include_variants": " | ".join(_load_json(r["include_variants_json"], [])),
                "exclude_variants": " | ".join(_load_json(r["exclude_variants_json"], [])),
                "source_refs": " | ".join(_load_json(r["source_refs_json"], [])),
                "evidence_status": r["evidence_status"],
                "components": " | ".join(COMPONENT_TOKEN_RE.findall(r["body"] or "")),
            })
        return pd.DataFrame(rows)

    def component_table(self) -> pd.DataFrame:
        return pd.read_sql_query("SELECT component_id,title,source_lang,version,source_refs_json,updated_at FROM components ORDER BY component_id", self.conn)

    def _variant_allowed(self, topic: sqlite3.Row, variant: str | None) -> bool:
        include = set(_load_json(topic["include_variants_json"], []))
        exclude = set(_load_json(topic["exclude_variants_json"], []))
        v = str(variant or "").strip()
        if include and v not in include:
            return False
        if v and v in exclude:
            return False
        return True

    def resolve_components(self, body: str) -> str:
        def repl(match: re.Match) -> str:
            cid = match.group(1)
            row = self.conn.execute("SELECT body FROM components WHERE component_id=?", (cid,)).fetchone()
            return row["body"] if row else f"[MISSING COMPONENT: {cid}]"
        return COMPONENT_TOKEN_RE.sub(repl, body or "")

    def assemble_publication(self, publication_id: str, *, variant: str | None = None,
                             target_lang: str | None = None) -> tuple[pd.DataFrame, dict]:
        target_lang = canonical_lang(target_lang) if target_lang else None
        pub = self.conn.execute("SELECT * FROM publications WHERE publication_id=?", (publication_id,)).fetchone()
        if not pub:
            raise KeyError(publication_id)
        rows = []
        stale_translations = 0
        missing_translations = 0
        for link in self.conn.execute("SELECT * FROM publication_topics WHERE publication_id=? ORDER BY position", (publication_id,)):
            topic = self.conn.execute("SELECT * FROM topics WHERE topic_id=?", (link["topic_id"],)).fetchone()
            if not self._variant_allowed(topic, variant):
                continue
            body = topic["body"]
            title = topic["title"]
            language = topic["source_lang"]
            translation_status = "SOURCE"
            if target_lang and target_lang != canonical_lang(topic["source_lang"]):
                tr = self.conn.execute("SELECT * FROM topic_translations WHERE topic_id=? AND REPLACE(target_lang,'_','-')=? COLLATE NOCASE", (topic["topic_id"], target_lang)).fetchone()
                if tr:
                    body = tr["target_body"]
                    title = str(tr["target_title"] or topic["title"])
                    language = target_lang
                    if int(tr["source_version"]) != int(topic["version"]):
                        stale_translations += 1
                        translation_status = "STALE"
                    else:
                        translation_status = tr["status"]
                else:
                    missing_translations += 1
                    translation_status = "MISSING"
            body = self.resolve_components(body)
            rows.append({
                "position": int(link["position"]), "topic_id": topic["topic_id"], "title": title,
                "topic_version": int(topic["version"]), "language": language, "translation_status": translation_status,
                "body": body, "source_refs": " | ".join(_load_json(topic["source_refs_json"], [])),
                "evidence_status": topic["evidence_status"],
            })
        frame = pd.DataFrame(rows)
        summary = {
            "publication_id": publication_id, "title": pub["title"], "variant": variant or "ALL",
            "target_lang": target_lang or pub["source_lang"], "topics": len(frame),
            "stale_translations": stale_translations, "missing_translations": missing_translations,
        }
        return frame, summary

    def impact_map(self, changed_id: str, *, kind: str = "topic") -> pd.DataFrame:
        impacted_topics: set[str] = set()
        if kind == "topic":
            impacted_topics.add(changed_id)
        elif kind == "component":
            token = f"{{{{component:{changed_id}}}}}"
            for r in self.conn.execute("SELECT topic_id,body FROM topics"):
                if token in (r["body"] or ""):
                    impacted_topics.add(r["topic_id"])
        else:
            raise ValueError("kind deve essere topic o component")
        rows = []
        for topic_id in sorted(impacted_topics):
            topic = self.conn.execute("SELECT * FROM topics WHERE topic_id=?", (topic_id,)).fetchone()
            for link in self.conn.execute("SELECT publication_id,position FROM publication_topics WHERE topic_id=?", (topic_id,)):
                pub = self.conn.execute("SELECT title FROM publications WHERE publication_id=?", (link["publication_id"],)).fetchone()
                rows.append({
                    "changed_kind": kind, "changed_id": changed_id, "topic_id": topic_id, "topic_title": topic["title"],
                    "publication_id": link["publication_id"], "publication_title": pub["title"] if pub else "",
                    "position": link["position"],
                    "include_variants": " | ".join(_load_json(topic["include_variants_json"], [])),
                    "exclude_variants": " | ".join(_load_json(topic["exclude_variants_json"], [])),
                })
        return pd.DataFrame(rows)

    def history(self, object_id: str, *, kind: str = "topic") -> pd.DataFrame:
        if kind == "topic":
            return pd.read_sql_query("SELECT topic_id,version,title,status,evidence_status,created_at FROM topic_versions WHERE topic_id=? ORDER BY version DESC", self.conn, params=(object_id,))
        if kind == "component":
            return pd.read_sql_query("SELECT component_id,version,title,created_at FROM component_versions WHERE component_id=? ORDER BY version DESC", self.conn, params=(object_id,))
        raise ValueError(kind)



def _pdf_running_header(text: str) -> bool:
    t = re.sub(r"\s+", " ", str(text or "")).strip()
    return bool(re.fullmatch(r"(?:Technical data\s+\d+|\d+\s+Technical data)", t, re.I))


def _pdf_semantic_heading(text: str, heading_flag: bool) -> bool:
    """Return True only for title-like PDF blocks.

    PDF font/layout heuristics frequently mark chart labels, bullets, formulas and
    full sentences as headings. Structured migration must be more conservative:
    false-positive topics are worse than keeping a line inside the current topic.
    """
    if not heading_flag:
        return False
    t = re.sub(r"\s+", " ", str(text or "")).strip().lstrip("￭").strip()
    if not t or _pdf_running_header(t):
        return False
    if t.startswith(("•", "- ", "– ", "— ")):
        return False
    if t.lower() in {"note", "note:", "where", "u", "k", "p/pn", "u/un"}:
        return False
    letters = sum(ch.isalpha() for ch in t)
    digits = sum(ch.isdigit() for ch in t)
    words = re.findall(r"[A-Za-zÀ-ÿ][A-Za-zÀ-ÿ0-9+./-]*", t)
    if letters < 3:
        return False
    if len(t) > 120 or len(words) > 14:
        return False
    if t.endswith((".", ":", ";")):
        return False
    # Formula/chart labels must never become topic titles.
    if ("=" in t or "×" in t or "√" in t or "·" in t) and len(words) <= 8:
        return False
    if digits > letters * 1.5 and len(words) < 3:
        return False
    # Long sentence-like blocks are body content even when font heuristics say heading.
    starters = (
        "the ", "this ", "these ", "those ", "when ", "for ", "make ", "derate ",
        "enabling ", "with ", "always ", "use ", "install ", "protect ", "if ", "an ", "a ",
    )
    if t.lower().startswith(starters) and len(words) > 5:
        return False
    return True


def _pdf_block_type(text: str) -> str:
    t = re.sub(r"\s+", " ", str(text or "")).strip()
    low = t.lower()
    if low.startswith(("warning:", "danger:", "caution:")):
        return "warning"
    if low == "note:" or low.startswith("note: "):
        return "note"
    if t.startswith(("•", "- ", "– ", "— ")):
        return "list_item"
    return "paragraph"


def import_pdf_as_topics(repo: ContentRepository, pdf_path: str | Path, *, source_lang: str = "en-US",
                         page_spec: str = "", id_prefix: str = "TOPIC", source_ref_prefix: str | None = None) -> pd.DataFrame:
    """Conservative PDF -> structured-topic migration.

    RC1.9 hardening:
    - accepts only semantic/title-like heading candidates;
    - discards recurring page headers instead of turning them into prose/topics;
    - keeps prose/list/note/warning blocks in reading order;
    - when the industrial CCMS API is available, preserves detected PDF tables as
      first-class structured table blocks instead of flattening them into prose.

    This remains a migration draft. PDF is not a lossless structured-authoring
    source and human review is mandatory before source approval.
    """
    layout = extract_layout(pdf_path, page_spec)
    src_name = source_ref_prefix or Path(pdf_path).name

    # Merge text/table events in physical reading order.
    events: list[dict[str, Any]] = []
    for block in layout.text_blocks:
        txt = re.sub(r"\s+", " ", block.text or "").strip()
        if not txt:
            continue
        y = float(block.bbox[1]) if getattr(block, "bbox", None) else 0.0
        events.append({
            "kind": "text", "page0": int(block.page), "y": y, "text": txt,
            "heading": bool(getattr(block, "heading", False)),
        })
    for table in getattr(layout, "tables", []) or []:
        y = float(table.bbox[1]) if getattr(table, "bbox", None) else 0.0
        events.append({
            "kind": "table", "page0": int(table.page), "y": y,
            "table_index": int(table.index), "rows": table.rows,
        })
    events.sort(key=lambda e: (e["page0"], e["y"], 0 if e["kind"] == "text" else 1))

    groups: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    dropped_headers = 0

    def ensure_current(page0: int) -> dict[str, Any]:
        nonlocal current
        if current is None:
            current = {"title": "Introduction", "items": [], "page": page0 + 1}
        return current

    for ev in events:
        if ev["kind"] == "text":
            txt = ev["text"]
            if _pdf_running_header(txt):
                dropped_headers += 1
                continue
            if _pdf_semantic_heading(txt, ev["heading"]):
                if current and current["items"]:
                    groups.append(current)
                current = {
                    "title": txt.lstrip("￭").strip()[:180],
                    "items": [],
                    "page": ev["page0"] + 1,
                }
            else:
                ensure_current(ev["page0"])["items"].append({
                    "kind": "text", "text": txt, "page": ev["page0"] + 1,
                    "block_type": _pdf_block_type(txt),
                })
        else:
            ensure_current(ev["page0"])["items"].append({
                "kind": "table", "rows": ev["rows"], "page": ev["page0"] + 1,
                "table_index": ev["table_index"],
            })
    if current and current["items"]:
        groups.append(current)

    frame = _write_topic_groups_to_repo(repo, groups, source_lang=source_lang, id_prefix=id_prefix, src_name=src_name)
    frame.attrs["dropped_running_headers"] = dropped_headers
    return frame


def _write_topic_groups_to_repo(repo: ContentRepository, groups: list[dict[str, Any]], *, source_lang: str,
                                 id_prefix: str, src_name: str) -> pd.DataFrame:
    """Format-agnostic CCMS write step shared by PDF, DOCX and IDML migration.

    `groups` is a list of {"title": str, "items": [...], "page": int} where each
    item is one of:
      {"kind":"text","text":str,"page":int,"block_type":str}
      {"kind":"table","rows":[[str,...],...],"page":int,"table_index":int}
      {"kind":"figure","page":int,"caption":str,"alt_text":str,"file_path":str}
    Every item MAY additionally carry an optional "source_refs_extra": [str,...]
    -- appended verbatim to that block/object's source_refs alongside the
    page/story reference. This is how the IDML importer attaches its missing-
    asset marker (never inside caption/alt_text) and a composite cross-revision
    fingerprint for a future change/revision intelligence consumer to read,
    without any schema change and with zero effect on PDF/DOCX, which never
    set this key.
    The caller is responsible for producing this structure from the source
    format; everything below (IDs, status, evidence, table/figure handling) is
    shared, tested domain logic — do not duplicate it in a new importer.
    """
    rows_out: list[dict[str, Any]] = []
    existing_topics = {r[0] for r in repo.conn.execute("SELECT topic_id FROM topics")}
    # Milestone 3D fix: create_figure is a SEPARATE capability, required only
    # for kind="figure" items. A repo that supports block/table but not
    # figures (or a plain non-industrial ContentRepository) must keep writing
    # text/table content exactly as before -- it must never be gated behind a
    # capability only IDML actually needs.
    industrial = all(hasattr(repo, name) for name in ("add_block", "create_table"))
    supports_figures = industrial and hasattr(repo, "create_figure")

    for i, group in enumerate(groups, start=1):
        base_id = f"{id_prefix}-{i:04d}"
        topic_id = base_id
        n = 1
        while topic_id in existing_topics:
            n += 1
            topic_id = f"{base_id}-{n}"
        existing_topics.add(topic_id)

        topic_ref = f"{src_name}#page={group['page']}"
        prose = [it["text"] for it in group["items"] if it["kind"] == "text"]
        body = "\n\n".join(prose)
        repo.create_topic(
            topic_id, group["title"], body, source_lang=source_lang, status="Migrated",
            source_refs=[topic_ref], evidence_status="SOURCE_GROUNDED",
        )

        structured_tables = 0
        figures_created = 0
        if industrial:
            pos = 1
            for item in group["items"]:
                ref = f"{src_name}#page={item['page']}"
                extra_refs = list(item.get("source_refs_extra") or [])
                if item["kind"] == "text":
                    bid = f"{topic_id}-B{pos:03d}"
                    repo.add_block(
                        bid, topic_id, item["block_type"], position=pos, text=item["text"],
                        status="Migrated", source_refs=[ref] + extra_refs, evidence_status="SOURCE_GROUNDED",
                    )
                elif item["kind"] == "table":
                    raw_rows = [[str(cell or "").strip() for cell in row] for row in (item["rows"] or [])]
                    width = max((len(r) for r in raw_rows), default=0)
                    if width:
                        normalized = [r + [""] * (width - len(r)) for r in raw_rows]
                        table_id = f"{topic_id}-T{item['table_index']+1:03d}"
                        # Collision-safe ID for re-import scenarios.
                        base_tid = table_id
                        suffix = 1
                        while repo.conn.execute("SELECT 1 FROM structured_tables WHERE table_id=?", (table_id,)).fetchone():
                            suffix += 1
                            table_id = f"{base_tid}-{suffix}"
                        repo.create_table(
                            table_id, [f"C{c+1}" for c in range(width)], normalized,
                            title=f"Imported table — page {item['page']}", header_rows=0,
                            status="Migrated", source_refs=[ref] + extra_refs,
                        )
                        bid = f"{topic_id}-B{pos:03d}"
                        repo.add_block(
                            bid, topic_id, "table", position=pos, object_ref_id=table_id,
                            status="Migrated", source_refs=[ref], evidence_status="SOURCE_GROUNDED",
                        )
                        structured_tables += 1
                elif item["kind"] == "figure" and not supports_figures:
                    # This repo can store blocks/tables but has no figures
                    # capability (e.g. a plain non-industrial ContentRepository,
                    # or a test double) -- skip only this item, never raise and
                    # never block the rest of the group's text/table content.
                    pass
                elif item["kind"] == "figure":
                    # Milestone 3D (IDML): first producer of the `figures` table,
                    # already fully consumed by publishing_bridges.py. A missing
                    # asset is represented via file_path="" + a technical
                    # source_refs marker (idml_asset_status=missing) — NEVER a
                    # diagnostic sentence inside caption/alt_text, which must
                    # only ever contain content that genuinely came from the
                    # source document.
                    figure_id = f"{topic_id}-F{pos:03d}"
                    base_fid = figure_id
                    suffix = 1
                    while repo.conn.execute("SELECT 1 FROM figures WHERE figure_id=?", (figure_id,)).fetchone():
                        suffix += 1
                        figure_id = f"{base_fid}-{suffix}"
                    repo.create_figure(
                        figure_id, title=item.get("title", "") or "", caption=item.get("caption", "") or "",
                        file_path=item.get("file_path", "") or "", alt_text=item.get("alt_text", "") or "",
                        status="Migrated", source_refs=[ref] + extra_refs,
                    )
                    bid = f"{topic_id}-B{pos:03d}"
                    repo.add_block(
                        bid, topic_id, "figure", position=pos, object_ref_id=figure_id,
                        status="Migrated", source_refs=[ref], evidence_status="SOURCE_GROUNDED",
                    )
                    figures_created += 1
                pos += 1

        rows_out.append({
            "topic_id": topic_id, "title": group["title"], "page": group["page"],
            "source_ref": topic_ref, "characters": len(body),
            "structured_tables": structured_tables, "figures_created": figures_created,
        })

    frame = pd.DataFrame(rows_out)
    frame.attrs["structured_table_count"] = int(frame["structured_tables"].sum()) if not frame.empty else 0
    frame.attrs["figure_count"] = int(frame["figures_created"].sum()) if not frame.empty else 0
    return frame


_LEGACY_DOCX_HEADING_RE = re.compile(r"^\d{1,3}\s*[-–—]\s*[A-ZÀ-Ý][^.]{2,78}$")


def _looks_like_legacy_docx_heading(text: str) -> bool:
    """Conservative fallback for chapter titles with no Word Heading style.

    Real Heading N styles remain the primary signal (checked first, unchanged)
    — this only fires when the style is missing/wrong, and only for a narrow,
    strong pattern: "<number> - <Title>" (e.g. "1 - Introduzione",
    "8 - Gestione segnalazioni - Allarmi"), anchored on the whole line so it
    cannot match inside a normal sentence. Deliberately rejects:
      - numbered lists using ". " or ") " instead of " - " (e.g. "1. Press...");
      - technical ranges/values where digits follow the dash (e.g. "20 - 25°C"),
        because the text after the dash must start with an uppercase letter;
      - long sentences or anything containing a period, which a genuine short
        chapter title normally does not.
    """
    return bool(_LEGACY_DOCX_HEADING_RE.match(str(text or "").strip()))


def _iter_docx_block_items(doc):
    """Yield ("paragraph"|"table", obj) in REAL document order.

    doc.paragraphs and doc.tables are separate flat lists in python-docx and do
    NOT preserve their relative position in the file — a table between two
    paragraphs would otherwise be lost from the sequence and misattributed.
    This walks the underlying XML body directly (the standard python-docx
    recipe for document-order iteration) so a table stays exactly where it
    was, between the paragraphs that precede and follow it.
    """
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    p_tag, tbl_tag = qn("w:p"), qn("w:tbl")
    for child in doc.element.body.iterchildren():
        if child.tag == p_tag:
            yield "paragraph", Paragraph(child, doc)
        elif child.tag == tbl_tag:
            yield "table", Table(child, doc)


def import_docx_as_topics(repo: ContentRepository, docx_path: str | Path, *, source_lang: str = "en-US",
                           id_prefix: str = "TOPIC", source_ref_prefix: str | None = None) -> pd.DataFrame:
    """Conservative Word (.docx) -> structured-topic migration.

    Mirrors import_pdf_as_topics: extracts headings/paragraphs/tables in real
    document order (python-docx has no page concept, so sections are numbered
    sequentially — shown to the guided user as "Sezione N", not a technical
    index), then writes them through the SAME shared repo-writing step used
    by the PDF path (_write_topic_groups_to_repo) so topic/block/table
    handling, status, and evidence tracking are identical for both formats.

    This remains a migration draft. Human review is mandatory before source
    approval, same as the PDF path.
    """
    doc = DocxDocument(str(docx_path))
    src_name = source_ref_prefix or Path(docx_path).name

    groups: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    section_no = 0

    def ensure_current() -> dict[str, Any]:
        nonlocal current, section_no
        if current is None:
            section_no += 1
            current = {"title": "Introduction", "items": [], "page": section_no}
        return current

    for kind, block in _iter_docx_block_items(doc):
        if kind == "paragraph":
            text = re.sub(r"\s+", " ", block.text or "").strip()
            if not text:
                continue
            style = block.style.name if block.style else ""
            # Real Heading N style is checked first and always wins; the legacy
            # pattern is only a fallback for paragraphs that aren't Heading-styled.
            if style.startswith("Heading") or _looks_like_legacy_docx_heading(text):
                if current and current["items"]:
                    groups.append(current)
                section_no += 1
                current = {"title": text[:180], "items": [], "page": section_no}
            else:
                g = ensure_current()
                g["items"].append({"kind": "text", "text": text, "page": g["page"], "block_type": _pdf_block_type(text)})
        else:  # table, encountered exactly where it appears in the document
            rows = [[cell.text for cell in row.cells] for row in block.rows]
            g = ensure_current()
            table_index = sum(1 for it in g["items"] if it["kind"] == "table")
            g["items"].append({"kind": "table", "rows": rows, "page": g["page"], "table_index": table_index})

    if current and current["items"]:
        groups.append(current)

    return _write_topic_groups_to_repo(repo, groups, source_lang=source_lang, id_prefix=id_prefix, src_name=src_name)


def import_idml_as_topics(repo: ContentRepository, idml_path: str | Path, *, source_lang: str = "en-US",
                          id_prefix: str = "IDML", source_ref_prefix: str | None = None,
                          assets_dir: str | Path | None = None) -> pd.DataFrame:
    """Milestone 3D: Adobe InDesign (.idml) -> structured-topic migration.

    idml_import.parse_idml() does the ENTIRE fail-closed parse (security
    validation, designmap/spread/story structural parsing, spread/frame-
    position reading order across stories, table/figure extraction, master-
    spread/orphan-story review flagging) before this function ever touches
    the CCMS. It then feeds the
    SAME shared _write_topic_groups_to_repo() used by PDF/DOCX/DOC -- no
    parallel CCMS-writing path, and this is the first importer to also
    populate the `figures` table (already fully consumed by
    publishing_bridges.py, never written to by any importer before this one).

    Mirrors import_pdf_as_topics/import_docx_as_topics: conservative migration
    draft, human review mandatory before source approval. No layout
    reconstruction and no InDesign round-trip are attempted or implied.

    Fail-closed hardening: the ENTIRE operation is atomic from the caller's
    point of view. Asset bytes are materialized into a hidden STAGING
    directory (never the final assets_dir) while parsed.groups records the
    eventual final path; the CCMS write is snapshotted before starting so any
    exception rolls back every topic/block/table/figure row created during
    THIS attempt. Only once the write has fully succeeded is the staging
    directory atomically renamed into place. On ANY failure -- parse,
    materialization, or CCMS write -- this function guarantees zero new rows
    and no assets_dir is ever created or left half-populated.
    """
    src_name = source_ref_prefix or Path(idml_path).name
    parsed = parse_idml(idml_path)  # fail-closed; nothing touched yet

    assets_dir = Path(assets_dir) if assets_dir is not None else None
    staging_dir = None
    if assets_dir is not None:
        staging_dir = assets_dir.parent / f".{assets_dir.name}.staging-{uuid4().hex[:8]}"
        materialize_figure_assets(idml_path, parsed.groups, staging_dir, record_dir=assets_dir)

    existing_topics_before = {r[0] for r in repo.conn.execute("SELECT topic_id FROM topics")}
    try:
        frame = _write_topic_groups_to_repo(repo, parsed.groups, source_lang=source_lang, id_prefix=id_prefix, src_name=src_name)
    except Exception:
        _rollback_new_ccms_rows(repo, existing_topics_before)
        if staging_dir is not None and staging_dir.exists():
            shutil.rmtree(staging_dir, ignore_errors=True)
        raise

    if staging_dir is not None and staging_dir.exists():
        # materialize_figure_assets only creates staging_dir when there was at
        # least one present_in_package asset to copy -- nothing to finalize
        # otherwise (no figures, or none resolvable).
        if assets_dir.exists():
            shutil.rmtree(assets_dir, ignore_errors=True)  # re-import: fully replace, never merge
        staging_dir.rename(assets_dir)

    # Master-spread content and unthreaded/orphan stories are never silently
    # discarded (Milestone 3D constraint): they are exposed here for the
    # caller (app.import_manual_ui) to persist as reviewable, blocking items
    # in the standard UI, never inserted into the primary reading flow.
    frame.attrs["review_items"] = parsed.review_items
    frame.attrs["missing_asset_count"] = parsed.missing_asset_total
    frame.attrs["ambiguous_asset_count"] = parsed.ambiguous_asset_total
    frame.attrs["figure_total"] = parsed.figure_total
    return frame


def _rollback_new_ccms_rows(repo: ContentRepository, existing_topics_before: set[str]) -> None:
    """Best-effort compensating rollback for a partially-written CCMS import.

    Each repo write method (create_topic/add_block/create_table/create_figure)
    commits immediately on its own -- there is no single outer DB transaction
    to roll back. This instead explicitly deletes every row created by the
    failed attempt: the new topics (FK ON DELETE CASCADE removes their blocks/
    versions) plus the structured_tables/figures rows those blocks pointed at,
    which are not reachable via the topics FK and must be cleaned up
    separately.
    """
    current = {r[0] for r in repo.conn.execute("SELECT topic_id FROM topics")}
    new_ids = current - existing_topics_before
    if not new_ids:
        return
    marks = ",".join("?" for _ in new_ids)
    orphaned = list(repo.conn.execute(
        f"SELECT block_type, object_ref_id FROM content_blocks WHERE topic_id IN ({marks}) "
        f"AND block_type IN ('table','figure') AND object_ref_id<>''", list(new_ids),
    ))
    repo.conn.executemany("DELETE FROM topics WHERE topic_id=?", [(t,) for t in new_ids])
    for block_type, object_ref_id in orphaned:
        table_name = "structured_tables" if block_type == "table" else "figures"
        id_col = "table_id" if block_type == "table" else "figure_id"
        repo.conn.execute(f"DELETE FROM {table_name} WHERE {id_col}=?", (object_ref_id,))
    repo.conn.commit()


def export_publication_html(repo: ContentRepository, publication_id: str, out_path: str | Path, *,
                            variant: str | None = None, target_lang: str | None = None) -> tuple[Path, dict]:
    frame, summary = repo.assemble_publication(publication_id, variant=variant, target_lang=target_lang)
    sections = []
    for _, r in frame.iterrows():
        body_html = "<p>" + html.escape(str(r["body"])).replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>"
        sections.append(
            f"<section data-topic='{html.escape(str(r['topic_id']))}' data-version='{r['topic_version']}'>"
            f"<h2>{html.escape(str(r['title']))}</h2>{body_html}"
            f"<small>Source: {html.escape(str(r['source_refs']))} · {html.escape(str(r['evidence_status']))}</small></section>"
        )
    out_path = Path(out_path)
    out_path.write_text(
        "<!doctype html><html><head><meta charset='utf-8'><title>" + html.escape(summary["title"]) + "</title>"
        "<style>body{font-family:Arial;max-width:920px;margin:40px auto;line-height:1.55;color:#111827}section{margin:0 0 42px}small{color:#64748b}</style></head><body>"
        f"<h1>{html.escape(summary['title'])}</h1><p>Variant: {html.escape(summary['variant'])} · Language: {html.escape(summary['target_lang'])}</p>"
        + "".join(sections) + "</body></html>", encoding="utf-8")
    return out_path, summary


def create_demo_project(out_dir: str | Path) -> dict:
    """Create a deterministic authoring project used by regression and UI demos."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    db = out_dir / "ManualSync_CCMS.sqlite"
    db.unlink(missing_ok=True)
    repo = ContentRepository(db)
    try:
        repo.create_component("SAFETY-POWER", "Power isolation warning", "WARNING: Disconnect power and wait 15 minutes before service.", source_refs=["SafetySpec#4.2"])
        repo.create_topic("INSTALL-001", "Installation", "Install the drive vertically.\n\n{{component:SAFETY-POWER}}", status="Approved", source_refs=["ManualEN#p30"])
        repo.create_topic("R5V2-001", "R5 v2 additional clearance", "Maintain 100 mm top clearance for frame R5 v2.", status="Approved", include_variants=["R5v2"], source_refs=["EngineeringChange#EC-2026-41"])
        repo.create_topic("SERVICE-001", "Service", "Use approved service tools only.\n\n{{component:SAFETY-POWER}}", status="Approved", source_refs=["ServiceSpec#2"])
        repo.create_publication("PUB-ACS", "ACS Installation Manual", ["INSTALL-001","R5V2-001","SERVICE-001"])
        repo.set_translation("INSTALL-001", "it-IT", "Installare il convertitore verticalmente.\n\n{{component:SAFETY-POWER}}", status="Reviewed")
        html_path, summary = export_publication_html(repo, "PUB-ACS", out_dir / "ACS_R5V2.html", variant="R5v2")
        return {"repo": db, "topics": repo.topic_table(), "components": repo.component_table(), "summary": summary, "html": html_path}
    finally:
        repo.close()
