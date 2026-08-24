from __future__ import annotations

import hashlib
import json
import shutil
import zipfile
from pathlib import Path

from docx import Document

from ccms_industrial import IndustrialRepository
from publishing_bridges import _clean_topic_heading, export_publication_docx
from release_manager import _validate_customer_docx, create_release_bundle
from product_version import PRODUCT_BUILD

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "regressione_desktop_r14_out"
if OUT.exists():
    shutil.rmtree(OUT)
OUT.mkdir(parents=True)

results = []
def check(name, cond, detail=""):
    ok = bool(cond)
    results.append({"name": name, "status": "PASS" if ok else "FAIL", "detail": str(detail)})
    if not ok:
        raise AssertionError(f"{name}: {detail}")

# 1) Heading normalization must preserve wording but collapse all legacy whitespace.
raw = "\uFFEDDeclassamento con speciali impostazioni del programma di controllo\ndel convertitore"
clean = _clean_topic_heading(raw)
check(
    "Legacy heading newline becomes one normal space",
    clean == "Declassamento con speciali impostazioni del programma di controllo del convertitore",
    repr(clean),
)
check("Heading contains no line breaks", "\n" not in clean and "\r" not in clean, repr(clean))
check("Heading contains no U+FFED", "\uFFED" not in clean, repr(clean))
check("Tabs/multiple spaces collapse", _clean_topic_heading("A\t  B\nC") == "A B C", repr(_clean_topic_heading("A\t  B\nC")))
check("Normal punctuation preserved", _clean_topic_heading("Fuses (UL) — 400 V") == "Fuses (UL) — 400 V")

# 2) Build a real CCMS fixture with the exact legacy newline pattern seen in the ABB R13 release.
project = OUT / "project"
project.mkdir()
repo = IndustrialRepository(project / "ManualSync_CCMS.sqlite")
repo.create_topic(
    "T1",
    "Derating with special drive control program settings",
    "Legacy body.",
    source_lang="en-US",
    status="Approved",
    source_refs=["SPEC#1"],
)
repo.create_publication("PUB1", "Demo Drive Manual", ["T1"], source_lang="en-US")
repo.add_block("B1", "T1", "paragraph", position=1, text="Install the drive vertically.", status="Approved", source_refs=["SPEC#1"])
repo.set_topic_title_translation(
    "T1",
    "it-IT",
    "\uFFEDDeclassamento con speciali impostazioni del programma di controllo\ndel convertitore",
    status="Reviewed",
)
repo.set_block_translation("B1", "it-IT", "Installare il convertitore verticalmente.", status="Reviewed")
repo.close()

repo = IndustrialRepository(project / "ManualSync_CCMS.sqlite")
out_doc = OUT / "normalized_heading.docx"
export_publication_docx(repo, "PUB1", out_doc, target_lang="it-IT", include_source_refs=False)
repo.close()

doc = Document(out_doc)
headings = [p.text for p in doc.paragraphs if p.style and p.style.name.startswith("Heading")]
expected_heading = "Declassamento con speciali impostazioni del programma di controllo del convertitore"
check("Exported customer DOCX has normalized heading", headings == [expected_heading], headings)
validation = _validate_customer_docx(out_doc)
check("Customer DOCX quality gate passes normalized artifact", validation["headings_checked"] == 1, validation)
check("Customer DOCX exposes no source refs", validation["source_refs_visible"] == 0, validation)

# 3) The quality gate itself must fail closed on an independently malformed DOCX.
bad_doc = Document()
bad_doc.add_heading("Titolo con\ninterruzione", level=1)
bad_path = OUT / "bad_heading.docx"
bad_doc.save(bad_path)
blocked = False
try:
    _validate_customer_docx(bad_path)
except RuntimeError:
    blocked = True
check("Quality gate blocks malformed heading", blocked)

# 4) Controlled release must package only clean current files and keep checksums coherent.
release_dir = OUT / "release"
(release_dir / "ARTIFACTS").mkdir(parents=True)
(release_dir / "AUDIT").mkdir(parents=True)
(release_dir / "ARTIFACTS" / "OLD_R13.docx").write_bytes(b"historical artifact")

bundle = create_release_bundle(
    project,
    "PUB1",
    target_langs=["it-IT"],
    released_by="Regression R14",
    release_id="REL-R14",
    out_dir=release_dir,
)
zip_path = Path(bundle["release_zip"])
check("R14 controlled release ZIP created", zip_path.exists() and zip_path.stat().st_size > 0, zip_path)

with zipfile.ZipFile(zip_path) as z:
    names = set(z.namelist())
    manifest = json.loads(z.read("RELEASE_MANIFEST.json"))
    declared = {"RELEASE_MANIFEST.json"}
    declared.update(a["path"] for a in manifest["artifacts"])
    declared.update(a["path"] for a in manifest["audit_files"])
    check("Historical R13 artifact excluded", "ARTIFACTS/OLD_R13.docx" not in names, sorted(names))
    check("ZIP equals current manifest scope", names == declared, {"names": sorted(names), "declared": sorted(declared)})
    check("Manifest uses R14 product build", manifest["build"] == PRODUCT_BUILD, manifest["build"])
    check("All artifact checksums valid", all(hashlib.sha256(z.read(a["path"])).hexdigest() == a["sha256"] for a in manifest["artifacts"]), manifest["artifacts"])
    check("All audit checksums valid", all(hashlib.sha256(z.read(a["path"])).hexdigest() == a["sha256"] for a in manifest["audit_files"]), manifest["audit_files"])

    it_path = next(a["path"] for a in manifest["artifacts"] if a["language"] == "it-IT" and a["type"] == "DOCX")
    extracted = OUT / "it_release.docx"
    extracted.write_bytes(z.read(it_path))
    release_doc = Document(extracted)
    release_headings = [p.text for p in release_doc.paragraphs if p.style and p.style.name.startswith("Heading")]
    check("Packaged Italian DOCX keeps normalized heading", release_headings == [expected_heading], release_headings)

summary = {
    "build": PRODUCT_BUILD,
    "tests": len(results),
    "pass": sum(r["status"] == "PASS" for r in results),
    "fail": sum(r["status"] == "FAIL" for r in results),
    "results": results,
}
(ROOT / "RISULTATO_REGRESSIONE_DESKTOP_R14.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
(ROOT / "RISULTATO_REGRESSIONE_DESKTOP_R14.txt").write_text(
    "\n".join(f"{r['status']} | {r['name']} | {r['detail']}" for r in results) + f"\n\nTOTAL {summary['pass']}/{summary['tests']} PASS\n",
    encoding="utf-8",
)
print(json.dumps({"tests": summary["tests"], "pass": summary["pass"], "fail": summary["fail"]}, indent=2))
