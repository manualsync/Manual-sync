"""UX Phase 0 -- the four dead ends of the V1 UX audit, and decision 8a.

Every section drives the REAL callbacks the interface wires, on an isolated
data home, and asserts what the user now sees. Where a check guards a fix
whose absence would be silent, a mutant copy of the module with the defect
reintroduced is loaded and the same check must fail on it -- otherwise the
check would prove nothing.

  8a    a CANCELLED job no longer blocks release readiness; a live one does
  P0-4  a job action says what happened, why QA failed segment by segment,
        and "Annulla" asks for confirmation first
  P0-1  a refused translation package names every unapproved item and opens it
  P0-2  the guided import can set the translation languages
  8b    Settings > Project: languages are added freely and removed only when
        nothing was produced for them, after confirmation
  P0-3  Release without a publishing profile explains it and creates one only
        on an explicit, audited confirmation
"""
from __future__ import annotations
import test_isolation  # noqa: E402,F401 -- FIRST import: temporary home + real-data guard, never the user's ManualSyncData (test_isolation.py)

import importlib.util
import json
import os
import shutil
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")

BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE))

# Isolated product data home: nothing here may touch a real ManualSyncData.
_TEST_HOME = Path(tempfile.mkdtemp(prefix="ms_ux_p0_home_"))
os.environ["HOME"] = str(_TEST_HOME)
os.environ["USERPROFILE"] = str(_TEST_HOME)

import app  # noqa: E402
import release_readiness as rr  # noqa: E402
from ccms_industrial import IndustrialRepository  # noqa: E402
from translation_roundtrip import XLIFF_NS  # noqa: E402
from workflow_management import WorkflowManager  # noqa: E402

OUT = Path(tempfile.mkdtemp(prefix="ms_ux_p0_"))
app.DATA_HOME = OUT / "ManualSyncData"
app.DATA_ROOT = app.DATA_HOME / "Projects"
app.ACTIVE_PROJECT_FILE = app.DATA_HOME / "active_project.json"
app.APP_LOG = app.DATA_HOME / "MANUALSYNC_APP.log"
app.DATA_ROOT.mkdir(parents=True, exist_ok=True)

results: list[dict] = []


def check(name, cond, detail=""):
    ok = bool(cond)
    results.append({"test": name, "status": "PASS" if ok else "FAIL", "detail": str(detail)[:400]})
    if not ok:
        raise AssertionError(f"{name}: {detail}")


def make_project(name: str, targets: str = "it-IT, de-DE") -> Path:
    """A project with one approved topic and two approved blocks in PUB-MAIN."""
    p = Path(app.create_project_ui(name, "ACME", "Drive", "MAN-1", "A", "en-US", targets, "Docs")[0])
    repo = IndustrialRepository(p / "ManualSync_CCMS.sqlite")
    try:
        repo.create_topic("T1", "Installation", "legacy", status="Approved", source_refs=["SPEC#1"])
        repo.create_publication("PUB-MAIN", "Machine Manual", ["T1"], source_lang="en-US")
        repo.add_block("B1", "T1", "paragraph", position=1, text="Maintain a clearance of 20 mm around the unit.",
                       status="Approved", source_refs=["SPEC#2"])
        repo.add_block("B2", "T1", "paragraph", position=2, text="Disconnect the supply before opening the cover.",
                       status="Approved", source_refs=["SPEC#3"])
    finally:
        repo.close()
    return p


def lang_zip(project: Path, lang: str) -> Path:
    _msg, _tbl, _master, zips = app.work_package_ui(str(project), "PUB-MAIN", "publication", None)
    return next(Path(z) for z in zips if f"_{lang}_" in Path(z).name)


def returned_xliff(pkg: Path, dest: Path, translate: dict[str, str] | None = None) -> Path:
    """The vendor's file: targets from `translate`, else the source copied over."""
    work = dest.with_suffix("")
    work.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(pkg) as zf:
        zf.extractall(work)
    xlf = next(work.glob("*_WORK_ONLY.xlf"))
    tree = ET.parse(xlf)
    ns = {"x": XLIFF_NS}
    for seg in tree.getroot().findall(".//x:segment", ns):
        src = "".join(seg.find("./x:source", ns).itertext())
        seg.find("./x:target", ns).text = (translate or {}).get(src, src)
        seg.set("state", "translated")
    tree.write(dest, encoding="utf-8", xml_declaration=True)
    return dest


def new_job(project: Path, pkg: Path) -> str:
    app.create_job_ui(str(project), str(pkg), "", "")
    wf = WorkflowManager(project)
    try:
        return str(wf.jobs_table().sort_values("created_at").iloc[-1]["job_id"])
    finally:
        wf.close()


def matrix_row(module, project: Path, lang: str) -> dict:
    r = module.assess_release_readiness(project, "PUB-MAIN", target_langs=[lang])
    return r["matrix"].set_index("language").loc[lang].to_dict()


def load_mutant(module_path: Path, replacements: list[tuple[str, str]], name: str):
    """A uniquely named copy of a module with a defect reintroduced."""
    src = module_path.read_text(encoding="utf-8")
    for old, new in replacements:
        assert src.count(old) == 1, f"mutation anchor not unique in {module_path.name}: {old[:60]!r}"
        src = src.replace(old, new, 1)
    path = OUT / f"{name}.py"
    path.write_text(src, encoding="utf-8")
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


try:
    # ================================================================ 8a
    P8 = make_project("UX 8a readiness")
    pkg_it = lang_zip(P8, "it-IT")
    bad = returned_xliff(pkg_it, OUT / "8a_untranslated.xlf")

    j_live = new_job(P8, pkg_it)
    app.job_action_ui(str(P8), j_live, "Segna inviato", None, None, False, "", False)
    app.job_action_ui(str(P8), j_live, "Ricevi XLIFF + controllo qualità (QA)", str(bad), None, False, "", False)
    row = matrix_row(rr, P8, "it-IT")
    check("8a-1: a live QA-blocked job still blocks its language", row["qa_blocked_jobs"] == 1
          and "job con QA bloccato" in row["issues"], row)

    app.job_action_ui(str(P8), j_live, "Annulla", None, None, False, "", False)
    row = matrix_row(rr, P8, "it-IT")
    check("8a-2: the same job, cancelled, no longer counts as QA-blocked", row["qa_blocked_jobs"] == 0
          and "job con QA bloccato" not in row["issues"], row)

    # Mixed: one cancelled blocked job + one live blocked job -> exactly one.
    j_live2 = new_job(P8, pkg_it)
    app.job_action_ui(str(P8), j_live2, "Segna inviato", None, None, False, "", False)
    app.job_action_ui(str(P8), j_live2, "Ricevi XLIFF + controllo qualità (QA)", str(bad), None, False, "", False)
    row = matrix_row(rr, P8, "it-IT")
    check("8a-3: with a cancelled and a live blocked job, only the live one counts", row["qa_blocked_jobs"] == 1, row)
    app.job_action_ui(str(P8), j_live2, "Annulla", None, None, False, "", False)

    # The whole point of 8a: once every blocker is really cleared, the
    # language becomes releasable again despite the cancelled history.
    good = returned_xliff(lang_zip(P8, "it-IT"), OUT / "8a_translated.xlf", {
        "Installation": "Installazione",
        "Maintain a clearance of 20 mm around the unit.": "Mantenere una distanza di 20 mm attorno all'unità.",
        "Disconnect the supply before opening the cover.": "Scollegare l'alimentazione prima di aprire il coperchio.",
    })
    j_ok = new_job(P8, lang_zip(P8, "it-IT"))
    app.job_action_ui(str(P8), j_ok, "Segna inviato", None, None, False, "", False)
    app.job_action_ui(str(P8), j_ok, "Ricevi XLIFF + controllo qualità (QA)", str(good), None, False, "", False)
    app.job_action_ui(str(P8), j_ok, "Approva traduzione", None, None, False, "Reviewer", True)
    row = matrix_row(rr, P8, "it-IT")
    check("8a-4: after cancelled jobs, a completed translation makes the language READY", row["status"] == "READY"
          and row["qa_blocked_jobs"] == 0, row)

    mutant = load_mutant(BASE / "release_readiness.py", [(
        'qa_blocked = int((live_jobs["qa_status"].astype(str) == "BLOCKED").sum())',
        'qa_blocked = int((lang_jobs["qa_status"].astype(str) == "BLOCKED").sum())')], "ms_ux_p0_mutant_rr")
    mrow = matrix_row(mutant, P8, "it-IT")
    check("8a-5 (mutant): with the defect reintroduced the cancelled jobs block again -- the checks above discriminate",
          mrow["qa_blocked_jobs"] == 2 and mrow["status"] == "BLOCKED", mrow)

    # ================================================================ P0-4
    import authoring_view as av
    import translation_view as tv

    IT_OK = {"Installation": "Installazione",
             "Maintain a clearance of 20 mm around the unit.": "Mantenere una distanza di 20 mm attorno all'unità.",
             "Disconnect the supply before opening the cover.": "Scollegare l'alimentazione prima di aprire il coperchio."}
    P4 = make_project("UX P0-4 quality")
    pkg4 = lang_zip(P4, "it-IT")
    V = app.job_action_view_ui

    def act(job, action, xlf=None, reviewer="", human=False):
        return V(str(P4), job, action, str(xlf) if xlf else None, None, False, reviewer, human)

    def html_text(value):
        import html as _h, re as _re
        return _h.unescape(_re.sub(r"<[^>]+>", " ", str(value)))

    jA = new_job(P4, pkg4)
    o = act(jA, "Segna inviato")
    check("P0-4 a: a normal action reports what happened and the next step", "Completato" in html_text(o[0])
          and "inviato" in html_text(o[0]).lower() and "Ricevi XLIFF" in html_text(o[0]), html_text(o[0]))
    o = act(jA, "Ricevi XLIFF + controllo qualità (QA)", returned_xliff(pkg4, OUT / "p04_untranslated.xlf"))
    txt = html_text(o[0])
    check("P0-4 b: a blocked check says BLOCKED, for which language, and why", "Bloccato" in txt
          and "Controllo qualità non superato" in txt and "traduzione italiana (it-IT)" in txt
          and "uguali al testo originale" in txt, txt)
    check("P0-4 c: the blocked notice says what to do", "Riesegui controllo qualità" in txt, txt)
    readable = o[3]["value"]
    check("P0-4 d: the segment table the user reads is in plain words and lists only failing segments",
          list(readable.columns) == ["N.", "Testo originale", "Traduzione ricevuta", "Problema"]
          and len(readable) == 2 and all("Uguale al testo originale" in p for p in readable["Problema"]),
          readable.to_dict("records"))
    raw = o[4]["value"]
    check("P0-4 e: the engine's raw audit is still delivered, as technical details",
          hasattr(raw, "columns") and {"hard_reason", "segment_id"} <= set(raw.columns), getattr(raw, "columns", raw))
    check("P0-4 f: the technical summary is kept, folded, inside the notice", "Dettagli tecnici" in txt
          and "global_blockers" in txt, txt[-400:])

    jC = new_job(P4, pkg4)
    act(jC, "Segna inviato")
    o = act(jC, "Ricevi XLIFF + controllo qualità (QA)", returned_xliff(pkg4, OUT / "p04_ok.xlf", IT_OK))
    txt = html_text(o[0])
    check("P0-4 g: a passed check says so and asks for the human approval", "Completato" in txt
          and "Controllo qualità superato" in txt and "Approva traduzione" in txt and not o[3]["visible"], txt)
    o = act(jC, "Approva traduzione", reviewer="Anna Bianchi", human=False)
    txt = html_text(o[0])
    job_c = app._job_row(str(P4), jC)
    check("P0-4 h: a refused approval is a notice in place, not a global error, and changes nothing",
          "Attenzione" in txt and "Approvazione umana esplicita obbligatoria" in txt and job_c["status"] == "QA_READY", txt)

    jB = new_job(P4, pkg4)
    act(jB, "Segna inviato")
    repo, svc = av.open_service(P4)
    try:
        svc.save_block("B1", text="Maintain a clearance of 25 mm around the unit.", expected_version=1, actor="Mario Rossi")
    finally:
        repo.close()
    o = act(jB, "Ricevi XLIFF + controllo qualità (QA)", returned_xliff(pkg4, OUT / "p04_ok_b.xlf", IT_OK))
    txt = html_text(o[0])
    check("P0-4 i: when the source changed after sending, the notice says exactly that -- the case that used to show 0 errors",
          "Bloccato" in txt and "Il testo originale è cambiato dopo l'invio" in txt
          and "Annulla questo lavoro" in txt and "nuovo pacchetto" in txt, txt)

    # Decisions: an action the state does not allow, and confirmations.
    jD = new_job(P4, pkg4)
    o = act(jD, "Approva traduzione")
    txt = html_text(o[0])
    check("P0-4 j: an action the state does not allow explains the state in words, never its code",
          "Azione non disponibile" in txt and "Da inviare" in txt and "READY_TO_SEND" not in txt
          and "Segna inviato" in txt, txt)
    o = act(jA, "Annulla")
    check("P0-4 k: 'Annulla' first asks -- nothing is cancelled yet",
          "Vuoi annullare il lavoro per la traduzione italiana (it-IT)?" in html_text(o[0])
          and o[7]["visible"] and o[8]["visible"] and o[9] == jA and app._job_row(str(P4), jA)["status"] == "QA_BLOCKED",
          (html_text(o[0]), o[9]))
    back = app.job_cancel_back_ui()
    check("P0-4 l: 'No, torna indietro' withdraws the question and changes nothing",
          back[0] == "" and not back[1]["visible"] and back[3] == "" and app._job_row(str(P4), jA)["status"] == "QA_BLOCKED",
          back)
    # The field now names another job; the confirmation still acts on jA.
    o = app.job_cancel_confirmed_ui(str(P4), jA, None, None, False, "", False)
    check("P0-4 m: the confirmation cancels exactly the job it asked about", app._job_row(str(P4), jA)["status"] == "CANCELLED"
          and app._job_row(str(P4), jB)["status"] == "QA_BLOCKED" and "Lavoro annullato" in html_text(o[0]), html_text(o[0]))
    check("P0-4 n: no job given / unknown job -> a plain explanation",
          "Nessun lavoro indicato" in html_text(act("", "Segna inviato")[0])
          and "Lavoro non trovato" in html_text(act("JOB-NOPE", "Segna inviato")[0]), "")

    # Every reason the engine can write becomes a sentence; unknown ones are kept verbatim.
    cases = {
        "segmento mancante nel file restituito": "Segmento assente",
        "sorgente XLIFF modificata": "testo originale è stato modificato",
        "target vuoto/incompleto": "Traduzione mancante",
        "numeri sorgente/traduzione non coerenti: ['20'] ≠ ['25']": "originale: 20 — traduzione: 25",
        "codici sorgente/traduzione non coerenti: ['X-100'] ≠ ['X-200']": "originale: X-100 — traduzione: X-200",
        "target identico alla sorgente nonostante lingue diverse (en-US→it-IT)": "Uguale al testo originale",
        "target quasi identico alla sorgente (97.0%) con lingue diverse (en-US→it-IT)": "Quasi uguale",
        "target linguisticamente più compatibile con en-US che con it-IT (stopword gate)": "lingua di destinazione",
        "tag inline sconosciuto X": "Altro controllo non superato: tag inline sconosciuto X",
    }
    bad = {k: tv._segment_problems(k) for k, v in cases.items() if v not in " ".join(s for _, s in tv._segment_problems(k))}
    check("P0-4 o: every engine reason is translated, and an unknown one is kept verbatim", not bad, bad)

    job_b = app._job_row(str(P4), jB)
    mutant_tv = load_mutant(BASE / "translation_view.py", [(
        'if summary and (not summary.get("source_version_ok", True) or not summary.get("source_snapshot_ok", True)):',
        'if False:')], "ms_ux_p0_mutant_tv")
    m_html, _ = mutant_tv.qa_notice(job_b, None)
    check("P0-4 p (mutant): without the source-change reason the notice would fall back to a vague message",
          "Il testo originale è cambiato" not in html_text(m_html), html_text(m_html))

    # ================================================================ P0-1
    import project_view as pv
    import translation_impact as timp
    W = app.work_package_view_ui

    P1n = make_project("UX P0-1 no languages", targets="")
    try:
        app.work_package_ui(str(P1n), "PUB-MAIN", "publication", None)
        engine_msg = ""
    except ValueError as exc:
        engine_msg = str(exc)
    check("P0-1 a: the engine still refuses a project without languages (fail-closed kept)",
          engine_msg.startswith(app._NO_TARGET_LANGS_PREFIX), engine_msg)
    o = W(str(P1n), "", "publication", None, "PUB-MAIN")
    txt = html_text(o[0])
    check("P0-1 b: the user is told languages are missing and gets the button to add them",
          "Bloccato" in txt and "Mancano le lingue di destinazione" in txt and o[6]["visible"] and not o[5]["visible"], txt)

    P1 = make_project("UX P0-1 in review")
    repo, svc = av.open_service(P1)
    try:
        svc.save_block("B2", text="Disconnect the supply before opening the front cover.", expected_version=1,
                       actor="Mario Rossi")
    finally:
        repo.close()
    app.studio_submit_review_ui(str(P1), "PUB-MAIN", "block", "B2", False, "Mario Rossi")
    o = W(str(P1), "", "publication", None, "PUB-MAIN")
    txt = html_text(o[0])
    check("P0-1 c: one content item in review -> exactly that item, its state in words, and why",
          "1 contenuto non è ancora approvato" in txt and "In revisione" in txt and "front cover" in txt
          and "Maintain a clearance" not in txt and "contenuto approvato" in txt, txt)
    check("P0-1 d: that item can be picked and opened", o[4]["visible"] and o[4]["value"] == "block:B2" and o[5]["visible"],
          o[4])
    bundle = app.studio_search_open_ui(str(P1), "PUB-MAIN", o[4]["value"], False)
    st_id_at = next(i for i, c in enumerate(app.ST_OUT_FULL) if c is app.st_id)
    check("P0-1 e: 'Apri nell'editor' opens exactly that content", bundle[st_id_at] == "B2", bundle[st_id_at])
    check("P0-1 f: the manual in context is used -- no identifier to type", o[4]["value"] and "non trovato" not in txt, "")
    o = W(str(P1), "PUB-NOPE", "publication", None, "PUB-MAIN")
    check("P0-1 g: an unknown object is a plain explanation", "Contenuto da tradurre non trovato" in html_text(o[0])
          and "PUB-NOPE" in html_text(o[0]), html_text(o[0]))

    # Parity with the engine gate over the states it can meet.
    PP = make_project("UX P0-1 parity")
    repo = IndustrialRepository(PP / "ManualSync_CCMS.sqlite")
    try:
        # A publication cannot be CREATED empty, but it becomes empty when its
        # chapters are unlinked -- the engine's "EMPTY" state.
        repo.create_publication("PUB-EMPTY", "Empty", ["T1"], source_lang="en-US")
        repo.conn.execute("DELETE FROM publication_topics WHERE publication_id='PUB-EMPTY'")
        repo.conn.commit()
        states = [
            ("all approved", []),
            ("block Draft", [("UPDATE content_blocks SET status='Draft' WHERE block_id='B1'", ())]),
            ("block In Review", [("UPDATE content_blocks SET status='In Review' WHERE block_id='B1'", ())]),
            ("topic Migrated", [("UPDATE topics SET status='Migrated' WHERE topic_id='T1'", ())]),
            ("archived Draft block", [("UPDATE content_blocks SET status='Draft', archived_at='2026-01-01' WHERE block_id='B2'", ())]),
        ]
        mismatches = []
        for label, sql in states:
            repo.conn.execute("UPDATE topics SET status='Approved' WHERE topic_id='T1'")
            repo.conn.execute("UPDATE content_blocks SET status='Approved', archived_at=NULL")
            for q, a in sql:
                repo.conn.execute(q, a)
            repo.conn.commit()
            items = pv.unapproved_for_packages(repo, "publication", "PUB-MAIN")
            engine = timp._publication_status(repo, "PUB-MAIN")
            if (not items) != (engine == "Approved") or any(i.status in pv.APPROVED_STATUSES for i in items):
                mismatches.append((label, engine, [i.status for i in items]))
        items = pv.unapproved_for_packages(repo, "publication", "PUB-EMPTY")
        if not (items and items[0].status == "EMPTY" and timp._publication_status(repo, "PUB-EMPTY") == "EMPTY"):
            mismatches.append(("empty publication", timp._publication_status(repo, "PUB-EMPTY"), items))
        repo.conn.execute("UPDATE content_blocks SET status='Draft', archived_at='2026-01-01' WHERE block_id='B2'")
        repo.conn.commit()
        arch = pv.unapproved_for_packages(repo, "publication", "PUB-MAIN")
        check("P0-1 h: an archived Draft block is listed and flagged as archived", arch and arch[-1].archived, arch)
        check("P0-1 i: the list is empty if and only if the engine accepts the manual (6 states)", not mismatches, mismatches)
        mutant_pv = load_mutant(BASE / "project_view.py", [(
            'APPROVED_STATUSES = frozenset({"Approved", "Published"})',
            'APPROVED_STATUSES = frozenset({"Approved", "Published", "In Review"})')], "ms_ux_p0_mutant_pv")
        repo.conn.execute("UPDATE content_blocks SET status='Approved', archived_at=NULL")
        repo.conn.execute("UPDATE content_blocks SET status='In Review' WHERE block_id='B1'")
        repo.conn.commit()
        check("P0-1 j (mutant): a weakened rule is caught -- the parity check discriminates",
              mutant_pv.unapproved_for_packages(repo, "publication", "PUB-MAIN") == []
              and timp._publication_status(repo, "PUB-MAIN") != "Approved", "")
        repo.conn.execute("UPDATE content_blocks SET status='Approved'")
        repo.conn.commit()
    finally:
        repo.close()

    o = W(str(PP), "", "publication", None, "PUB-MAIN")
    txt = html_text(o[0])
    direct = app.work_package_ui(str(PP), "PUB-MAIN", "publication", None)
    check("P0-1 k: with everything approved the packages are generated exactly as before",
          "Completato" in txt and "pacchetti lingua creati" in txt and "traduzione italiana (it-IT)" in txt
          and sorted(Path(f).name for f in o[3]) == sorted(Path(f).name for f in direct[3]), txt)

    orig = app.work_package_ui
    app.work_package_ui = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom-xyz"))
    try:
        o = W(str(PP), "", "publication", None, "PUB-MAIN")
    finally:
        app.work_package_ui = orig
    txt = html_text(o[0])
    check("P0-1 l: an unexpected failure is explained, and its real reason is kept in the details",
          "Errore" in txt and "Non è stato possibile generare i pacchetti" in txt and "boom-xyz" in txt, txt)

    # ================================================================ P0-2 (guided import languages)
    import docx
    from project_lifecycle import read_project
    import gradio as gr

    def manual_docx(path: Path) -> Path:
        d = docx.Document()
        d.add_heading("Installazione", level=1)
        d.add_paragraph("Mantenere una distanza di 20 mm intorno alla macchina durante l'installazione.")
        d.add_paragraph("Scollegare l'alimentazione prima di aprire il coperchio anteriore.")
        d.save(path)
        return path

    def projects_now() -> set[str]:
        return {p.name for p in app.DATA_ROOT.iterdir()}

    doc = manual_docx(OUT / "guided_manual.docx")
    out = app.guided_confirm_language_ui("UX P0-2 guidato", str(doc), "Italiano (it-IT)", ["de-DE", "fr-FR", "it-IT"])
    meta = read_project(out[0])
    check("P0-2 a: the languages picked in the guided import are saved; the document's own language is left out",
          meta["target_langs"] == ["de-DE", "fr-FR"] and meta["source_lang"] == "it-IT", meta)
    msg = str(out[3])
    check("P0-2 b: the result says which languages were set and why the source one was not",
          "Lingue di traduzione impostate: Tedesco (de-DE), Francese (fr-FR)." in msg
          and "Italiano (it-IT) è la lingua del documento" in msg
          and msg.rstrip().endswith("Prossimo passo: rivedi il contenuto importato prima di procedere."), msg)
    check("P0-2 c: the import itself ran exactly as before (content in the review table)",
          len(out[11]) > 0 and out[12] == str(doc) and out[13] == "it-IT", (len(out[11]), out[12], out[13]))

    out = app.guided_confirm_language_ui("UX P0-2 senza lingue", str(doc), "Italiano (it-IT)")
    msg = str(out[3])
    check("P0-2 d: without languages the flow and its message are unchanged (no language line)",
          read_project(out[0])["target_langs"] == [] and "Lingue di traduzione" not in msg
          and "\n\nProssimo passo: rivedi" in msg and "lingua del documento" not in msg, msg)

    before = projects_now()
    try:
        app.guided_confirm_language_ui("UX P0-2 errore", str(doc), "Italiano (it-IT)", ["de-DE", "tedesco"])
        err = ""
    except gr.Error as exc:
        err = exc.message
    check("P0-2 e: a mistyped language stops the import BEFORE anything is created, naming the value",
          "tedesco" in err and projects_now() == before, (err, projects_now() - before))

    orig_detect = app.detect_docx_language
    app.detect_docx_language = lambda path: ("it", 0.99, "test")
    try:
        out = app.guided_new_project_from_pdf_ui("UX P0-2 rilevata", str(doc), ["pt_br", "PT-BR"])
    finally:
        app.detect_docx_language = orig_detect
    check("P0-2 f: the detected-language path carries the languages too; a typed code is normalised once",
          read_project(out[0])["target_langs"] == ["pt-BR"], read_project(out[0]))

    out = app.guided_confirm_language_ui("UX P0-2 guidato", str(doc), "Italiano (it-IT)", ["es-ES"])
    check("P0-2 g: an existing project is not silently changed, and the user is told where to add languages",
          read_project(out[0])["target_langs"] == ["de-DE", "fr-FR"] and "esisteva già" in str(out[3]), out[3])
    check("P0-2 h: the guided picker offers named languages, the source picker's ones included",
          ("Tedesco (de-DE)", "de-DE") in pv.language_choices() and ("Italiano (it-IT)", "it-IT") in pv.language_choices(),
          pv.language_choices()[:4])

    # ================================================================ 8b (Settings › Project)
    S = app.settings_project_ui

    def langs(project: Path) -> list[str]:
        return read_project(project)["target_langs"]

    o = S(None)
    check("8b a: no open project -> a plain explanation and no controls",
          "Nessun progetto aperto" in html_text(o[0]) and not o[1]["visible"] and not o[3]["visible"], html_text(o[0]))

    PB = make_project("UX 8b lingue")  # source en-US, targets it-IT + de-DE
    o = S(str(PB))
    facts = html_text(o[0])
    check("8b b: the project's details are shown read-only, languages by name",
          "UX 8b lingue" in facts and "Inglese (en-US)" in facts and "Tedesco (de-DE), Italiano (it-IT)" in facts
          and "ACME" in facts and "MAN-1" in facts, facts)
    add_codes = [c for _, c in o[1]["choices"]]
    check("8b c: the add list leaves out the source and the languages already present",
          "en-US" not in add_codes and "it-IT" not in add_codes and "de-DE" not in add_codes and "fr-FR" in add_codes,
          add_codes)
    check("8b d: the remove list holds exactly the translation languages -- never the source",
          [c for _, c in o[3]["choices"]] == ["de-DE", "it-IT"], o[3]["choices"])

    o = app.settings_add_languages_ui(str(PB), ["fr-FR", "en-US", "de-DE", "pt_br", "fr-fr"])
    txt = html_text(o[5])
    check("8b e: adding saves the new languages once, normalised",
          langs(PB) == ["de-DE", "fr-FR", "it-IT", "pt-BR"], langs(PB))
    check("8b f: the notice says what was added and why the rest was ignored",
          "Completato" in txt and "Lingue aggiunte" in txt and "Francese (fr-FR)" in txt and "Portoghese (pt-BR)" in txt
          and "Inglese (en-US): è la lingua del documento originale" in txt
          and "Tedesco (de-DE): è già tra le lingue" in txt, txt)
    o = app.settings_add_languages_ui(str(PB), ["nl-NL", "olandese!"])
    check("8b g: a value that is not a language code saves nothing and is named",
          langs(PB) == ["de-DE", "fr-FR", "it-IT", "pt-BR"] and "olandese!" in html_text(o[5])
          and "Codice lingua non riconosciuto" in html_text(o[5]), html_text(o[5]))
    o = app.settings_add_languages_ui(str(PB), [])
    check("8b h: adding nothing asks for a choice", "Nessuna lingua scelta" in html_text(o[5]), html_text(o[5]))

    # One language per kind of data, each with ONLY that kind.
    PB_ZIP = lang_zip(PB, "de-DE")
    new_job(PB, PB_ZIP)                                       # de-DE: a translation job
    repo = IndustrialRepository(PB / "ManualSync_CCMS.sqlite")
    try:
        repo.set_block_translation("B1", "it-IT", "Mantenere 20 mm.")  # it-IT: a translated text
        repo.conn.execute(
            "INSERT INTO publication_releases(release_id,publication_id,profile_id,target_lang,snapshot_json,"
            "snapshot_sha256,released_by,created_at) VALUES('REL-X','PUB-MAIN','','fr-FR','{}','x','QA','2026-01-01')")
        repo.conn.commit()                                    # fr-FR: a release
    finally:
        repo.close()
    from cat_tm import TranslationMemory
    (PB / "TranslationMemory").mkdir(exist_ok=True)
    tm = TranslationMemory(PB / "TranslationMemory" / "Vendor_TM.sqlite")
    try:                                                      # pt-BR: memory, stored in the legacy form
        tm.add("en-US", "pt_BR", "Disconnect the supply.", "Desligue a alimentação.")
        tm.conn.execute("UPDATE tm_entries SET target_lang='pt_br'")
        tm.conn.commit()
    finally:
        tm.close()
    expected = {"de-DE": ("jobs", "1 lavoro di traduzione"), "it-IT": ("translations", "1 testo tradotto"),
                "fr-FR": ("releases", "1 rilascio pubblicato"), "pt-BR": ("memory", "1 frase nella memoria")}
    wrong = {}
    for code, (field, words) in expected.items():
        usage = pv.language_usage(PB, code)
        o = app.settings_remove_language_ui(str(PB), code)
        txt = html_text(o[5])
        others = [f for f in ("translations", "jobs", "releases", "memory") if f != field and getattr(usage, f)]
        if not (getattr(usage, field) == 1 and not others and "Bloccato" in txt and words in txt
                and not o[6]["visible"] and o[8] == "" and code in langs(PB)):
            wrong[code] = (usage, txt[:200])
    check("8b i: a language with translations, a job, a release or memory cannot be removed -- each named", not wrong, wrong)

    app.settings_add_languages_ui(str(PB), ["nl-NL"])
    o = app.settings_remove_language_ui(str(PB), "nl-NL")
    check("8b j: a language with no data asks for confirmation first -- nothing removed yet",
          "Rimuovere Olandese (nl-NL)?" in html_text(o[5]) and o[6]["visible"] and o[7]["visible"] and o[8] == "nl-NL"
          and "nl-NL" in langs(PB) and o[3]["interactive"] is False, html_text(o[5]))
    o = app.settings_remove_back_ui(str(PB))
    check("8b k: 'No, mantieni' keeps the language and closes the question",
          "nl-NL" in langs(PB) and not o[6]["visible"] and o[8] == "" and o[5] == "", o[8])
    app.settings_remove_language_ui(str(PB), "nl-NL")
    o = app.settings_remove_confirmed_ui(str(PB), "nl-NL")
    check("8b l: confirming removes exactly that language and says so",
          langs(PB) == ["de-DE", "fr-FR", "it-IT", "pt-BR"] and "Lingua rimossa" in html_text(o[5])
          and not o[6]["visible"], (langs(PB), html_text(o[5])))

    app.settings_add_languages_ui(str(PB), ["sv-SE"])
    o = app.settings_remove_language_ui(str(PB), "sv-SE")
    tm = TranslationMemory(PB / "ManualSync_CONFIRMED_TM.sqlite")
    try:
        tm.add("en-US", "sv-SE", "Disconnect the supply.", "Koppla från matningen.")
    finally:
        tm.close()
    o = app.settings_remove_confirmed_ui(str(PB), o[8])
    check("8b m: data recorded between the question and the answer still blocks the removal",
          "sv-SE" in langs(PB) and "Bloccato" in html_text(o[5]), html_text(o[5]))

    o = app.settings_remove_language_ui(str(PB), "en-US")
    check("8b n: the source language cannot be removed (it is not a translation language)",
          "en-US" == read_project(PB)["source_lang"] and "non è tra le lingue di traduzione" in html_text(o[5]),
          html_text(o[5]))
    o = app.settings_remove_language_ui(str(PB), None)
    check("8b o: removing with nothing chosen asks for a choice", "Scegli la lingua da rimuovere" in html_text(o[5]), "")

    app.settings_add_languages_ui(str(PB), ["da-DK"])
    orig_usage = pv.language_usage
    pv.language_usage = lambda *a, **k: (_ for _ in ()).throw(OSError("disk-xyz"))
    try:
        o = app.settings_remove_language_ui(str(PB), "da-DK")
    finally:
        pv.language_usage = orig_usage
    check("8b p: if the check itself fails the language stays (fail-closed) and the reason is kept",
          "da-DK" in langs(PB) and "Verifica non riuscita" in html_text(o[5]) and "disk-xyz" in html_text(o[5])
          and not o[6]["visible"], html_text(o[5]))

    mutant_jobs = load_mutant(BASE / "project_view.py", [(
        "return any((self.translations, self.jobs, self.releases, self.memory))",
        "return any((self.translations, self.releases, self.memory))")], "ms_ux_p0_mutant_pv_jobs")
    mutant_tm = load_mutant(BASE / "project_view.py", [(
        'for tm in sorted(base.rglob("*.sqlite")):', 'for tm in []:')], "ms_ux_p0_mutant_pv_tm")
    check("8b q (mutants): ignoring jobs or translation memories would let a language with data go -- caught",
          pv.language_usage(PB, "de-DE").in_use and not mutant_jobs.language_usage(PB, "de-DE").in_use
          and pv.language_usage(PB, "pt-BR").in_use and not mutant_tm.language_usage(PB, "pt-BR").in_use, "")

    import i18n
    import re as _re
    nested = {}
    for lang in ("it", "en"):
        cat = json.loads((BASE / "locales" / f"{lang}.json").read_text(encoding="utf-8"))
        for key, text in cat.items():
            if key.startswith("ux.") and "{language}" in text:
                sample = text.replace("{language}", i18n.translate("ux.lang.translation", lang, name="x", code="de-DE"))
                if _re.search(r"\([^()]*\(", sample):
                    nested[f"{lang}:{key}"] = sample
    check("UX text: no notice nests a language in double parentheses", not nested, nested)

    # ================================================================ P0-3 (Release: publishing profile)
    import publishing_output as po
    N = app._PROFILE_DIM_SLOTS

    def profiles(project: Path, inactive: bool = True) -> list[dict]:
        repo, svc = av.open_service(project)
        try:
            return svc.list_variant_profiles(include_inactive=inactive)
        finally:
            repo.close()

    def audit_count(project: Path) -> int:
        repo = IndustrialRepository(project / "ManualSync_CCMS.sqlite")
        try:
            return int(repo.conn.execute("SELECT COUNT(*) FROM audit_events").fetchone()[0])
        finally:
            repo.close()

    def form_visible(o) -> bool:  # name, save and cancel are shown together
        return all(o[i].get("visible") for i in (2, 3 + N, 4 + N))

    P3 = make_project("UX P0-3 profilo")
    audit0 = audit_count(P3)
    o = app.release_profile_notice_ui(str(P3))
    txt = html_text(o[0])
    check("P0-3 a: no profile -> the page says a profile is needed and why, and offers to create one",
          "Attenzione" in txt and "Serve un profilo di pubblicazione" in txt and "anteprima ed esportazione" in txt
          and o[1]["visible"] and not form_visible(o), txt)
    check("P0-3 b: showing the notice creates nothing (no profile, no audit event)",
          profiles(P3) == [] and audit_count(P3) == audit0, (profiles(P3), audit_count(P3), audit0))

    o = app.release_profile_open_ui(str(P3))
    check("P0-3 c: 'Crea profilo' opens the form with the name 'Standard' and, without variants, no picker",
          form_visible(o) and o[2]["value"] == "Standard" and not any(o[3 + i]["visible"] for i in range(N))
          and not o[1]["visible"] and "non ha varianti" in html_text(o[0]) and o[-1] == [], html_text(o[0]))
    o = app.release_profile_notice_ui(str(P3))  # "Annulla" is wired to this
    check("P0-3 d: 'Annulla' closes the form and still creates nothing",
          not form_visible(o) and o[1]["visible"] and profiles(P3) == [], "")

    C = app.release_profile_create_ui
    o = C(str(P3), "Standard", "", [])
    check("P0-3 e: without the operator's name nothing is created and the user is told where to write it",
          profiles(P3) == [] and "Manca il tuo nome" in html_text(o[0]) and "Operatore" in html_text(o[0])
          and "visible" not in o[2], html_text(o[0]))
    o = C(str(P3), "  ", "Mario Rossi", [])
    check("P0-3 f: an empty name is refused, nothing created", profiles(P3) == []
          and "Manca il nome del profilo" in html_text(o[0]), html_text(o[0]))

    o = C(str(P3), "Standard", "Mario Rossi", [])
    made = profiles(P3)
    check("P0-3 g: confirming creates ONE active profile with a readable id and no variant restriction",
          [(p["profile_id"], p["name"], p["dimensions"], p["active"]) for p in made]
          == [("PROF-STANDARD", "Standard", {}, True)], made)
    repo = IndustrialRepository(P3 / "ManualSync_CCMS.sqlite")
    try:
        ev = repo.conn.execute("SELECT actor FROM audit_events WHERE action='VARIANT_PROFILE_CREATED' "
                               "AND object_id='PROF-STANDARD'").fetchall()
    finally:
        repo.close()
    check("P0-3 h: the creation is audited under the operator's name, like Settings › Varianti",
          [r[0] for r in ev] == ["Mario Rossi"], [tuple(r) for r in ev])
    check("P0-3 i: the new profile is selected at once in the Release picker and the shared state",
          o[-2]["value"] == "PROF-STANDARD" and ("Standard", "PROF-STANDARD") in [tuple(c) for c in o[-2]["choices"]]
          and o[-1] == "PROF-STANDARD" and "Completato" in html_text(o[0]) and not form_visible(o)
          and not o[1]["visible"], (o[-2], o[-1]))
    preview = html_text(app.studio_preview_refresh_ui(str(P3), "PUB-MAIN", "PROF-STANDARD", False))
    repo, svc = av.open_service(P3)
    try:
        res = po.export_pdf(repo, svc, "PUB-MAIN", "PROF-STANDARD")
    finally:
        repo.close()
    check("P0-3 j: with it, the preview shows the manual and the PDF export produces a real file",
          "Maintain a clearance" in preview and Path(res["pdf"]).stat().st_size > 1000 and res["pages"] >= 1,
          preview[:200])
    o = app.release_profile_notice_ui(str(P3))
    check("P0-3 k: once a profile exists the notice and the offer disappear", o[0] == "" and not o[1]["visible"], "")

    PI = make_project("UX P0-3 inattivo")
    repo, svc = av.open_service(PI)
    try:
        svc.create_variant_profile("PROF-OLD", "Vecchio", {}, actor="Writer")
        svc.set_variant_profile_active("PROF-OLD", False, actor="Writer")
    finally:
        repo.close()
    check("P0-3 l: only INACTIVE profiles still count as 'no profile' (the picker would be empty)",
          "Serve un profilo" in html_text(app.release_profile_notice_ui(str(PI))[0]), "")
    o = C(str(PI), "Vecchio", "Mario Rossi", [])
    check("P0-3 m: a name already used gets its own unused id -- nothing is overwritten",
          sorted(p["profile_id"] for p in profiles(PI)) == ["PROF-OLD", "PROF-VECCHIO"], profiles(PI))

    PD = make_project("UX P0-3 varianti")
    repo, svc = av.open_service(PD)
    try:
        svc.create_condition_dimension("market", "Mercato", actor="Writer")
        svc.create_condition_value("eu", "market", "Europa", actor="Writer")
        svc.create_condition_value("us", "market", "Stati Uniti", actor="Writer")
    finally:
        repo.close()
    o = app.release_profile_open_ui(str(PD))
    check("P0-3 n: with variants the form asks for each one by name, 'Nessuna scelta' first",
          o[3]["visible"] and o[3]["label"] == "Mercato"
          and [tuple(c) for c in o[3]["choices"]] == [("Nessuna scelta", app._PROFILE_DIM_NONE), ("Europa", "eu"),
                                                        ("Stati Uniti", "us")]
          and o[-1] == ["market"] and "ha delle varianti" in html_text(o[0]), o[3])
    blank = [app._PROFILE_DIM_NONE] * N
    o = C(str(PD), "Europa", "Mario Rossi", ["market"], "eu", *blank[1:])
    check("P0-3 o: the chosen variant is saved in the profile", [p["dimensions"] for p in profiles(PD)]
          == [{"market": "eu"}], profiles(PD))
    o = C(str(PD), "Rotto", "Mario Rossi", ["market"], "zz", *blank[1:])
    check("P0-3 p: a refusal by the service is explained, its reason kept, nothing created",
          "Profilo non creato" in html_text(o[0]) and "zz" in html_text(o[0]) and len(profiles(PD)) == 1,
          html_text(o[0]))

    o = app.release_export_profile_check_ui(str(make_project("UX P0-3 export")), None, [])
    check("P0-3 q: pressing Export with no profile at all shows the way forward again",
          "Serve un profilo" in html_text(o[0]) and o[1]["visible"], "")
    o = app.release_export_profile_check_ui(str(P3), "PROF-STANDARD", ["x"])
    check("P0-3 r: with a profile selected the export click leaves this area untouched",
          all(isinstance(u, dict) and set(u) <= {"__type__"} for u in o[:-1]) and o[-1] == ["x"], o[:3])

    deps = {getattr(app.demo.fns.get(d.get("id")).fn, "__name__", ""): d
            for d in app.demo.get_config_file()["dependencies"] if app.demo.fns.get(d.get("id"))}
    create_dep = deps.get("release_profile_create_ui", {})
    check("P0-3 s: 'Crea e usa questo profilo' runs once per click (a double click cannot make two profiles)",
          create_dep.get("trigger_mode") == "once", create_dep.get("trigger_mode"))
    check("P0-3 t: the preview/export callbacks themselves are still wired unchanged",
          [d["outputs"] for n, d in deps.items() if n in ("studio_export_pdf_ui", "studio_export_html_ui")]
          == [[app.pv_export_status._id], [app.pv_export_status._id]], "")
    mutant_id = load_mutant(BASE / "project_view.py", [("while pid in taken:", "while False:")],
                            "ms_ux_p0_mutant_pv_id")
    check("P0-3 u (mutant): without the unique suffix the id would collide with an existing profile",
          mutant_id.profile_id_for("Vecchio", ["PROF-VECCHIO"]) == "PROF-VECCHIO"
          and pv.profile_id_for("Vecchio", ["PROF-VECCHIO"]) == "PROF-VECCHIO-2", "")
except BaseException as _crash:
    # A crash must never read as "0 failed" in the summary.
    if not isinstance(_crash, AssertionError):
        results.append({"test": "suite interrupted by an unexpected error", "status": "FAIL",
                        "detail": f"{type(_crash).__name__}: {_crash}"[:400]})
    raise
finally:
    for _m in [m for m in sys.modules if m.startswith("ms_ux_p0_mutant_")]:
        sys.modules.pop(_m, None)
    summary = {"suite": "UX Phase 0 -- dead ends, 8a, 8b", "tests": len(results),
               "pass": sum(r["status"] == "PASS" for r in results),
               "fail": sum(r["status"] == "FAIL" for r in results), "results": results}
    (BASE / "RISULTATO_UX_P0.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: summary[k] for k in ("suite", "tests", "pass", "fail")}, ensure_ascii=False))
    shutil.rmtree(OUT, ignore_errors=True)
    shutil.rmtree(_TEST_HOME, ignore_errors=True)
