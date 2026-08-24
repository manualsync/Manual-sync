from __future__ import annotations

import html
import json
import os
import re
import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from ccms_industrial import IndustrialRepository

PUBLISHING_BUILD = "0.18.3-heading-whitespace-normalization-20260818"
IDML_PLACEHOLDER = "[[MANUALSYNC_CONTENT]]"


def _clean_topic_heading(value: Any) -> str:
    """Normalize semantic headings for customer-facing publication.

    Legacy imports can contain migration-only U+FFED glyphs and hard line breaks
    inside a heading. A Word heading must remain a single semantic line: collapsing
    whitespace to one normal space prevents strings such as ``controllo\ndel`` from
    rendering or being extracted later as ``controllodel`` while preserving the
    approved words and punctuation.
    """
    text = str(value or "").replace("\uFFED", " ")
    return re.sub(r"\s+", " ", text, flags=re.UNICODE).strip()


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _add_source_ref(doc: Document, source_refs: str, evidence: str) -> None:
    if not source_refs:
        return
    p = doc.add_paragraph()
    p.style = doc.styles["Caption"]
    run = p.add_run(f"Source: {source_refs} · {evidence}")
    run.font.size = Pt(8)
    run.font.italic = True


def export_publication_docx(repo: IndustrialRepository, publication_id: str, out_path: str | Path, *,
                            profile_id: str | None = None, dimensions: dict[str, Any] | None = None,
                            target_lang: str | None = None, include_source_refs: bool = True) -> tuple[Path, dict[str, Any]]:
    frame, summary = repo.assemble_structured_publication(
        publication_id, profile_id=profile_id, dimensions=dimensions, target_lang=target_lang
    )
    pub = repo.conn.execute("SELECT title FROM publications WHERE publication_id=?", (publication_id,)).fetchone()
    title = pub["title"] if pub else publication_id

    doc = Document()
    now = datetime.now(timezone.utc)
    props = doc.core_properties
    props.title = str(title)
    props.subject = f"ManualSync controlled technical publication · {target_lang or 'source'}"
    props.author = "ManualSync"
    props.last_modified_by = "ManualSync"
    props.comments = ""
    props.created = now
    props.modified = now
    try:
        props.language = str(target_lang or "en-US")
    except Exception:
        pass
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(24)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(16)

    t = doc.add_paragraph(style="Title")
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.add_run(title)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if str(target_lang or "").lower().startswith("it"):
        meta_text = f"Profilo: {profile_id or 'ALL'} · Lingua: {target_lang or 'sorgente'}"
    else:
        meta_text = f"Profile: {profile_id or 'ALL'} · Language: {target_lang or 'source'}"
    meta.add_run(meta_text).italic = True

    last_topic = None
    for _, row in frame.iterrows():
        if row["topic_id"] != last_topic:
            doc.add_heading(_clean_topic_heading(row["topic_title"]), level=1)
            last_topic = row["topic_id"]

        bt = str(row["block_type"])
        content = row["content"]
        if bt in {"paragraph", "component_ref", "code"}:
            p = doc.add_paragraph()
            if bt == "code":
                p.style = doc.styles["No Spacing"]
            p.add_run(str(content))
        elif bt == "warning":
            table = doc.add_table(rows=1, cols=1)
            table.autofit = True
            cell = table.cell(0,0)
            _set_cell_shading(cell, "FFF2CC")
            run = cell.paragraphs[0].add_run("WARNING\n" + str(content))
            run.bold = True
        elif bt == "note":
            p = doc.add_paragraph()
            r = p.add_run("NOTE: " + str(content))
            r.italic = True
        elif bt == "list_item":
            doc.add_paragraph(str(content), style="List Bullet")
        elif bt == "table" and isinstance(content, dict):
            cols = [str(x) for x in content.get("columns", [])]
            rows = content.get("rows", []) or []
            if content.get("caption"):
                cap = doc.add_paragraph(style="Caption")
                cap.add_run(str(content["caption"]))
            if cols:
                tbl = doc.add_table(rows=1, cols=len(cols))
                tbl.style = "Table Grid"
                hdr = tbl.rows[0]
                _set_repeat_table_header(hdr)
                for i, col in enumerate(cols):
                    hdr.cells[i].text = col
                    _set_cell_shading(hdr.cells[i], "D9EAF7")
                    for run in hdr.cells[i].paragraphs[0].runs:
                        run.bold = True
                for data_row in rows:
                    cells = tbl.add_row().cells
                    for i, value in enumerate(list(data_row)[:len(cols)]):
                        cells[i].text = str(value)
        elif bt == "figure" and isinstance(content, dict):
            fp = Path(str(content.get("file_path", "")))
            if fp.exists() and fp.is_file():
                try:
                    doc.add_picture(str(fp), width=Inches(4.8))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception:
                    doc.add_paragraph(f"[Figure file not renderable: {fp.name}]")
            else:
                doc.add_paragraph(f"[FIGURE: {content.get('figure_id','')}] {content.get('alt_text','')}")
            if content.get("caption"):
                cap = doc.add_paragraph(style="Caption")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.add_run(str(content["caption"]))

        if include_source_refs:
            _add_source_ref(doc, str(row.get("source_refs", "")), str(row.get("evidence_status", "")))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Generated by ManualSync")

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path, summary


def import_docx_outline(docx_path: str | Path) -> list[dict[str, Any]]:
    """Conservative Word bridge: extracts heading-defined topics, paragraphs and tables.

    It does not pretend to preserve arbitrary Word layout. The returned data keeps
    structural evidence (paragraph style, table boundaries) for later migration.
    """
    doc = Document(str(docx_path))
    items: list[dict[str, Any]] = []
    current_topic = "INTRO"
    for p in doc.paragraphs:
        text = (p.text or "").strip()
        if not text:
            continue
        style = p.style.name if p.style else ""
        if style.startswith("Heading"):
            current_topic = text
            items.append({"kind":"heading", "topic":current_topic, "style":style, "text":text})
        else:
            items.append({"kind":"paragraph", "topic":current_topic, "style":style, "text":text})
    for ti, table in enumerate(doc.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        items.append({"kind":"table", "topic":current_topic, "table_index":ti, "rows":rows})
    return items


def _block_plain_text(row) -> str:
    bt = str(row["block_type"])
    c = row["content"]
    if bt == "table" and isinstance(c, dict):
        lines = []
        cols = c.get("columns", []) or []
        if c.get("caption"):
            lines.append(str(c["caption"]))
        if cols:
            lines.append("\t".join(str(x) for x in cols))
        for rr in c.get("rows", []) or []:
            lines.append("\t".join(str(x) for x in rr))
        return "\r".join(lines)
    if bt == "figure" and isinstance(c, dict):
        return f"[FIGURE {c.get('figure_id','')}] {c.get('caption','')}"
    prefix = {"warning":"WARNING: ", "note":"NOTE: ", "list_item":"• "}.get(bt, "")
    return prefix + str(c)


def create_idml_template_fixture(out_path: str | Path) -> Path:
    """Create a deterministic structural fixture for bridge regression only.

    The fixture mimics the Story XML location used by IDML packages but is not
    presented as an Adobe-authored production template. Production export is
    template-driven: users provide a valid IDML template exported by InDesign.
    """
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    story = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Story Self="u1"><StoryPreference OpticalMarginAlignment="false"/><InCopyExportOption/><ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/$ID/[Basic Paragraph]"><CharacterStyleRange AppliedCharacterStyle="CharacterStyle/$ID/[No character style]"><Content>{IDML_PLACEHOLDER}</Content></CharacterStyleRange></ParagraphStyleRange></Story>'''
    designmap = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Document DOMVersion="19.0"><idPkg:Story src="Stories/Story_u1.xml" xmlns:idPkg="http://ns.adobe.com/AdobeInDesign/idml/1.0/packaging"/></Document>'''
    with zipfile.ZipFile(out_path, "w") as zf:
        zf.writestr("mimetype", "application/vnd.adobe.indesign-idml-package", compress_type=zipfile.ZIP_STORED)
        zf.writestr("designmap.xml", designmap)
        zf.writestr("Stories/Story_u1.xml", story)
    return out_path


def patch_idml_template(repo: IndustrialRepository, publication_id: str, template_idml: str | Path, out_idml: str | Path, *,
                        profile_id: str | None = None, dimensions: dict[str, Any] | None = None,
                        target_lang: str | None = None, placeholder: str = IDML_PLACEHOLDER) -> tuple[Path, dict[str, Any]]:
    """Inject ManualSync content into a valid template-driven IDML package.

    The bridge intentionally depends on an IDML template exported by InDesign.
    It preserves all untouched package parts and replaces a dedicated story
    placeholder with structured paragraphs. Complex InDesign-native tables,
    anchored graphics and master-page styling are not reconstructed from scratch.
    """
    template_idml = Path(template_idml)
    out_idml = Path(out_idml)
    if not zipfile.is_zipfile(template_idml):
        raise ValueError("Il template IDML non è un pacchetto ZIP valido")
    frame, summary = repo.assemble_structured_publication(publication_id, profile_id=profile_id, dimensions=dimensions, target_lang=target_lang)

    tmp = out_idml.parent / (out_idml.stem + "_tmp")
    if tmp.exists():
        shutil.rmtree(tmp)
    tmp.mkdir(parents=True)
    with zipfile.ZipFile(template_idml, "r") as zin:
        zin.extractall(tmp)

    story_files = sorted((tmp / "Stories").glob("*.xml")) if (tmp / "Stories").exists() else []
    target_story = None
    target_content = None
    for sf in story_files:
        root = ET.parse(sf).getroot()
        for content in root.iter("Content"):
            if (content.text or "") == placeholder:
                target_story = sf
                target_content = content
                break
        if target_story:
            break
    if not target_story or target_content is None:
        shutil.rmtree(tmp, ignore_errors=True)
        raise ValueError(f"Placeholder {placeholder} non trovato nelle Stories del template IDML")

    tree = ET.parse(target_story)
    root = tree.getroot()
    content_node = None
    for c in root.iter("Content"):
        if (c.text or "") == placeholder:
            content_node = c; break
    if content_node is None:
        shutil.rmtree(tmp, ignore_errors=True)
        raise ValueError("Placeholder perso durante il parsing")

    parent_char = None
    parent_para = None
    for p in root.iter("ParagraphStyleRange"):
        for ch in p.findall("CharacterStyleRange"):
            for c in ch.findall("Content"):
                if c is content_node:
                    parent_char = ch; parent_para = p; break
            if parent_para is not None: break
        if parent_para is not None: break
    if parent_para is None:
        shutil.rmtree(tmp, ignore_errors=True)
        raise ValueError("Struttura Story IDML non riconosciuta")

    children = list(root)
    insert_at = children.index(parent_para)
    root.remove(parent_para)

    paragraphs = []
    last_topic = None
    for _, row in frame.iterrows():
        if row["topic_id"] != last_topic:
            paragraphs.append((str(row["topic_title"]), "ParagraphStyle/$ID/Heading 1"))
            last_topic = row["topic_id"]
        paragraphs.append((_block_plain_text(row), "ParagraphStyle/$ID/[Basic Paragraph]"))

    for offset, (text, pstyle) in enumerate(paragraphs):
        p = ET.Element("ParagraphStyleRange", {"AppliedParagraphStyle": pstyle})
        ch = ET.SubElement(p, "CharacterStyleRange", {"AppliedCharacterStyle":"CharacterStyle/$ID/[No character style]"})
        c = ET.SubElement(ch, "Content")
        c.text = text
        root.insert(insert_at + offset, p)

    tree.write(target_story, encoding="UTF-8", xml_declaration=True)
    bridge_manifest = {
        "manualsync_build": PUBLISHING_BUILD,
        "publication_id": publication_id,
        "profile_id": profile_id or "",
        "target_lang": target_lang or "",
        "story": str(target_story.relative_to(tmp)).replace(os.sep, "/"),
        "inserted_paragraphs": len(paragraphs),
        "limitations": [
            "Template-driven bridge: a valid IDML template exported by InDesign is required for production use.",
            "Structured tables are injected as tab-delimited story text in v0.18.",
            "Figures are referenced textually; placed graphics are not reconstructed in v0.18.",
        ],
    }
    (tmp / "ManualSyncBridge.json").write_text(json.dumps(bridge_manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    out_idml.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(out_idml, "w") as zout:
        mime = tmp / "mimetype"
        if mime.exists():
            zout.write(mime, "mimetype", compress_type=zipfile.ZIP_STORED)
        for p in sorted(tmp.rglob("*")):
            if not p.is_file() or p == mime:
                continue
            zout.write(p, str(p.relative_to(tmp)).replace(os.sep, "/"), compress_type=zipfile.ZIP_DEFLATED)
    shutil.rmtree(tmp, ignore_errors=True)
    return out_idml, bridge_manifest


def import_idml_story_outline(idml_path: str | Path) -> list[dict[str, str]]:
    if not zipfile.is_zipfile(idml_path):
        raise ValueError("IDML non leggibile come package")
    rows = []
    with zipfile.ZipFile(idml_path, "r") as zf:
        for name in sorted(n for n in zf.namelist() if n.startswith("Stories/") and n.endswith(".xml")):
            root = ET.fromstring(zf.read(name))
            for p in root.iter("ParagraphStyleRange"):
                style = p.attrib.get("AppliedParagraphStyle", "")
                text = "".join((c.text or "") for c in p.iter("Content"))
                if text.strip():
                    rows.append({"story":name, "paragraph_style":style, "text":text})
    return rows
