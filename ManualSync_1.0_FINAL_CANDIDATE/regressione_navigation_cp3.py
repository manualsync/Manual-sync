"""Checkpoint 3 -- navigation and application structure.

Drives the real callbacks and walks the real Gradio wiring, on an isolated
data home. Where a check guards a fix whose absence would be silent, a
mutant copy with the defect reintroduced must fail the same check.

  H  the header, sidebar and Home follow the SESSION's project, after every
     event that changes it
  P  the project selector opens only real projects, and remembers the choice
  S  no section keeps showing the previous project's results
  G  the guided import is confirmed where the user can see it
  L  interface language: header, immediate, remembered; not the project's
  N  sidebar highlight, quick actions, section names, the manual in context
"""
from __future__ import annotations
import test_isolation  # noqa: E402,F401 -- FIRST import: temporary home + real-data guard, never the user's ManualSyncData (test_isolation.py)

import importlib.util
import json
import os
import re
import shutil
import sys
import tempfile
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")

BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE))

_TEST_HOME = Path(tempfile.mkdtemp(prefix="ms_nav_cp3_home_"))
os.environ["HOME"] = str(_TEST_HOME)
os.environ["USERPROFILE"] = str(_TEST_HOME)

import app  # noqa: E402
import gradio as gr  # noqa: E402
import i18n  # noqa: E402
import shell_view as sv  # noqa: E402
from ccms_industrial import IndustrialRepository  # noqa: E402

OUT = Path(tempfile.mkdtemp(prefix="ms_nav_cp3_"))


def isolate(module):
    module.DATA_HOME = OUT / "ManualSyncData"
    module.DATA_ROOT = module.DATA_HOME / "Projects"
    module.ACTIVE_PROJECT_FILE = module.DATA_HOME / "active_project.json"
    module.APP_LOG = module.DATA_HOME / "MANUALSYNC_APP.log"
    module.DATA_ROOT.mkdir(parents=True, exist_ok=True)


isolate(app)
results: list[dict] = []


def check(name, cond, detail=""):
    ok = bool(cond)
    results.append({"test": name, "status": "PASS" if ok else "FAIL", "detail": str(detail)[:400]})
    if not ok:
        raise AssertionError(f"{name}: {detail}")


def text(value) -> str:
    import html as _h
    return _h.unescape(re.sub(r"<[^>]+>", " ", str(value)))


def skipped(u) -> bool:
    return isinstance(u, dict) and u.get("__type__") == "update" and len(u) == 1


def make_project(module, name, client, targets, marker):
    p = Path(module.create_project_ui(name, client, "Macchina", "MAN-1", "A", "en-US", targets, "Docs")[0])
    repo = IndustrialRepository(p / "ManualSync_CCMS.sqlite")
    try:
        repo.create_topic("T1", f"{marker} chapter", "legacy", status="Approved", source_refs=["SPEC#1"])
        repo.create_publication("PUB-MAIN", f"{marker} manual", ["T1"], source_lang="en-US")
        repo.add_block("B1", "T1", "paragraph", position=1, text=f"{marker}: keep a clearance of 20 mm.",
                       status="Approved", source_refs=["SPEC#2"])
    finally:
        repo.close()
    return p


def load_mutant(replacements, name):
    src = (BASE / "app.py").read_text(encoding="utf-8")
    for old, new in replacements:
        assert src.count(old) == 1, f"mutation anchor not unique: {old[:70]!r}"
        src = src.replace(old, new, 1)
    path = OUT / f"{name}.py"
    path.write_text(src, encoding="utf-8")
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    isolate(mod)
    return mod


def graph(module):
    """Dependencies of a built app, with each one's callable name."""
    deps = module.demo.get_config_file()["dependencies"]
    fn = {}
    for d in deps:
        bf = module.demo.fns.get(d["id"])
        fn[d["id"]] = getattr(getattr(bf, "fn", None), "__name__", "")
    return deps, fn


def chain_after(deps, first):
    out, cur = [], first
    while True:
        nxt = [x for x in deps if x.get("trigger_after") == cur["id"]]
        if not nxt:
            return out
        cur = nxt[0]
        out.append(cur)


def descendants(deps, first):
    """Every step that runs after `first`, following all branches."""
    out, todo = [], [first]
    while todo:
        cur = todo.pop(0)
        nxt = [x for x in deps if x.get("trigger_after") == cur["id"]]
        out.extend(nxt)
        todo.extend(nxt)
    return out


def triggered_by(deps, component, event):
    return [d for d in deps if any(t[0] == component._id and t[1] == event for t in d.get("targets", []))]


def project_change_gaps(module):
    """Events that write project_state but are not followed by the
    project-change routine (project_changed_ui)."""
    deps, fn = graph(module)
    gaps = []
    for d in deps:
        if module.project_state._id in d.get("outputs", []) and d.get("trigger_after") is None:
            names = {fn[x["id"]] for x in descendants(deps, d)}
            if "project_changed_ui" not in names:
                gaps.append(fn[d["id"]])
    return gaps


def walk(block):
    yield block
    for c in getattr(block, "children", []) or []:
        yield from walk(c)


try:
    A = make_project(app, "Progetto Alfa", "Cliente Alfa", "it-IT, de-DE", "ALFA")
    B = make_project(app, "Progetto Beta", "", "fr-FR", "BETA")
    app.create_job_ui(str(A), str(next(z for z in app.work_package_ui(str(A), "PUB-MAIN", "publication", None)[3]
                                       if "_de-DE_" in z)), "", "")
    (app.DATA_ROOT / "Not a project").mkdir()
    (app.DATA_ROOT / "Half project").mkdir()
    (app.DATA_ROOT / "Half project" / "ManualSync_Project.json").write_text("{}", encoding="utf-8")

    # ============================================================ H: the shell
    o = app.shell_ui(None, "PUB-MAIN", "it")
    check("H1: with no project the header selects nothing, the card says so, quick actions are hidden",
          o[0]["value"] is None and "Nessun progetto aperto" in text(o[4]) and not o[6]["visible"]
          and not o[7]["visible"], (o[0]["value"], text(o[4])[:80]))
    o = app.shell_ui(str(A), "PUB-MAIN", "it")
    labels = [c[0] for c in o[0]["choices"]]
    check("H2: the header lists exactly the real projects, by name, and selects the open one",
          labels == ["Progetto Alfa — Cliente Alfa", "Progetto Beta"] and o[0]["value"] == A.name
          and "Progetto Alfa" in text(o[3]) and "Progetto Alfa" in text(o[4]) and o[6]["visible"], labels)
    check("H3: no filesystem path reaches the header (labels or values)",
          not [c for c in o[0]["choices"] if re.search(r"[\\/:]", c[0] + c[1])], o[0]["choices"])
    app._set_active_project(B)  # another session moved the on-disk pointer
    o = app.shell_ui(str(A), "PUB-MAIN", "it")
    check("H4: the shell draws THIS session's project, never the on-disk pointer",
          o[0]["value"] == A.name and "Progetto Alfa" in text(o[4]) and "Progetto Alfa" in text(o[3]),
          (o[0]["value"], text(o[4])[:60]))
    check("H5: every event that changes the project is followed by the shell routine",
          not project_change_gaps(app), project_change_gaps(app))
    deps, fn = graph(app)
    header_writers = sorted({fn[d["id"]] for d in deps if app.ws_project._id in d.get("outputs", [])})
    check("H6: the header project is a selector written only by the shell (and the language labels)",
          isinstance(app.ws_project, gr.Dropdown)
          and header_writers == ["project_changed_ui", "shell_ui", "switch_language_ui"],
          header_writers)

    mutant = load_mutant([
        ("choices, key = sv.project_choices(DATA_ROOT, project_dir)",
         "choices, key = sv.project_choices(DATA_ROOT, _load_active_project())"),
        ("                _project_card(project_dir or None),",
         "                _project_card(_load_active_project() or None),"),
        ("for _project_evt in (_open_evt, _create_evt, _restore_evt, _import_evt, _confirm_evt):",
         "for _project_evt in (_open_evt, _create_evt, _import_evt, _confirm_evt):"),
    ], "ms_nav_cp3_mutant_app")
    mo = mutant.shell_ui(str(A), "PUB-MAIN", "it")
    check("H7 (mutant): a shell reading the on-disk pointer would name the wrong project -- caught",
          mo[0]["value"] == B.name and "Progetto Beta" in text(mo[4]), mo[0]["value"])
    check("H8 (mutant): an event that changes the project without the routine is caught",
          project_change_gaps(mutant) == ["restore_project_view_ui"], project_change_gaps(mutant))

    # ============================================================ P: the selector
    app._set_active_project(A)
    o = app.open_project_ui(B.name, str(A), "it")
    check("P1: choosing another project opens it, remembers it for the next start and says so",
          Path(o[0]) == B.resolve() and json.loads(app.ACTIVE_PROJECT_FILE.read_text(encoding="utf-8"))["project_dir"]
          == str(B.resolve()) and "Progetto «Progetto Beta» aperto" in text(o[1]), text(o[1]))
    o = app.open_project_ui(B.name, str(B), "it")
    check("P2: choosing the project already open changes nothing", skipped(o[0]) and skipped(o[1]), o)
    refused = {}
    for bad in ("..", "Not a project", "Half project", "C:\\Windows", "../Projects/" + A.name, "Nessuno"):
        before = app.ACTIVE_PROJECT_FILE.read_text(encoding="utf-8")
        o = app.open_project_ui(bad, str(B), "it")
        if not (skipped(o[0]) and "Progetto non disponibile" in text(o[1])
                and app.ACTIVE_PROJECT_FILE.read_text(encoding="utf-8") == before):
            refused[bad] = o
    check("P3: anything but a real project folder of the data root is refused, nothing changes", not refused, refused)
    check("P4: folders that are not complete projects are never offered",
          [e.key for e in sv.list_projects(app.DATA_ROOT)] == [A.name, B.name],
          [e.key for e in sv.list_projects(app.DATA_ROOT)])
    open_dep = triggered_by(deps, app.ws_project, "select")
    check("P5: the selector reacts to the USER's choice only (.select), never to the shell setting it",
          len(open_dep) == 1 and fn[open_dep[0]["id"]] == "open_project_ui"
          and not triggered_by(deps, app.ws_project, "change"), [fn[d["id"]] for d in open_dep])

    # ============================================================ S: no stale screens
    o = app.project_scope_ui(str(B), str(A), "it")
    check("S1: after a change every registered result is put back to a fresh page, and the change is recorded",
          o[0] == str(B) and len(o) == len(app.PROJECT_SCOPED) + 1 and not any(skipped(x) for x in o[1:]), len(o))
    o = app.project_scope_ui(str(B), str(B), "it")
    check("S2: without a change nothing is touched (e.g. the import asked for the document language)",
          all(skipped(x) for x in o), "")

    scoped_ids = {c._id for c, _r in app.PROJECT_SCOPED}
    in_tabs = {c._id for c in walk(app.workspace_tabs)}
    open_chain = descendants(deps, open_dep[0])
    ALLOWED = {app.home_notice._id: "the report of the change itself",
               app.new_project_panel._id: "Home intake form, not project data",
               app.legacy_onboarding_panel._id: "Home intake form, not project data",
               app.legacy_multi_panel._id: "Home intake form, not project data"}

    def uncovered(scoped):
        project_data = {}
        for d in deps:
            if app.project_state._id in d.get("inputs", []):
                for o_ in d.get("outputs", []):
                    if o_ in in_tabs:
                        project_data.setdefault(o_, fn[d["id"]])
        # The registry is covered only through `scoped` itself: project_scope_ui's
        # outputs ARE the registry, so counting them too would hide a gap.
        registry = {c._id for c, _r in app.PROJECT_SCOPED}
        reloaded = {o_ for d in open_chain for o_ in d.get("outputs", [])} - registry
        covered = set(scoped) | reloaded | set(ALLOWED)
        # A reset component re-runs its own .change listeners, which refill what they write.
        for d in deps:
            if any(t[0] in scoped and t[1] == "change" for t in d.get("targets", [])):
                covered |= set(d.get("outputs", []))
        return {o_: f for o_, f in project_data.items() if o_ not in covered}

    names = {v._id: k for k, v in vars(app).items() if isinstance(v, gr.blocks.Block)}
    gaps = uncovered(scoped_ids)
    check("S3: every component that shows project data is cleared or reloaded on a project change",
          not gaps, {names.get(k, k): v for k, v in gaps.items()})
    check("S4 (mutant): forgetting one result panel is caught by S3",
          app.t_table._id in uncovered(scoped_ids - {app.t_table._id}), "")
    change = [d for d in open_chain if fn[d["id"]] == "project_changed_ui"]
    outs = change[0]["outputs"] if change else []
    check("S4b: everything visible changes in ONE step, before the heavy reloads, with no output twice",
          len(change) == 1 and change[0]["trigger_after"] == open_dep[0]["id"] and len(outs) == len(set(outs))
          and app.w_jobs._id in outs and app.ws_project._id in outs and app.t_table._id in outs
          and all(d["trigger_after"] in {change[0]["id"]} | {x["id"] for x in open_chain}
                  for d in open_chain if d is not change[0]), len(outs))
    o = app.project_changed_ui(str(B), "PUB-MAIN", str(A), None, "en")
    check("S4c: and that step returns a value for every one of its outputs, in the session language",
          len(o) == len(outs) and "Project name" in text(o[len(outs) - 1 - 3 - len(app._SETTINGS_PROJECT_OUT)]),
          len(o))
    o = app.workflow_tables_ui(str(A))
    ob = app.workflow_tables_ui(str(B))
    check("S5: Lavorazioni is reloaded for the new project: Alfa's job does not appear under Beta",
          o[0]["visible"] and len(o[0]["value"]) == 1 and not ob[0]["visible"], (o[0]["visible"], ob[0]["visible"]))
    check("S6: without a project the editor reload does nothing -- no error toast during an import",
          all(skipped(x) for x in app.studio_reload_ui("", "PUB-MAIN", False)), "")
    bundle = app.studio_reload_ui(str(B), "PUB-MAIN", False)
    outline = next(v for c, v in zip(app.ST_OUT_FULL, bundle) if c is app.st_outline_radio)
    check("S7: the editor reloads the NEW project's content", "BETA chapter" in json.dumps(outline, ensure_ascii=False)
          and "ALFA" not in json.dumps(outline, ensure_ascii=False), str(outline)[:200])

    # ============================================================ G: guided import
    import docx
    doc = docx.Document()
    doc.add_heading("Installazione", level=1)
    doc.add_paragraph("Mantenere una distanza di 20 mm intorno alla macchina durante l'installazione.")
    manual = OUT / "guidato.docx"
    doc.save(manual)
    o = app.guided_confirm_view_ui("Progetto Gamma", str(manual), "Italiano (it-IT)", ["de-DE"], "it")
    notice = text(o[-1])
    check("G1: after the guided import the Home says what was imported, where, and with which languages",
          len(o) == 17 and "Manuale importato nel progetto «Progetto Gamma»" in notice
          and "Lingua del documento: Italiano (it-IT)" in notice and "Lingue di traduzione: Tedesco (de-DE)" in notice
          and "Rivedi il contenuto importato" in notice, notice)
    check("G2: its next step is on screen: the review button is shown, outside the panel that hides",
          o[6]["visible"] and not any(c is app.lg_review for c in walk(app.legacy_onboarding_panel))
          and not any(c is app.home_notice for c in walk(app.new_project_panel)), "")
    orig = app.detect_docx_language
    app.detect_docx_language = lambda path: ("it", 0.2, "test")
    try:
        o = app.guided_import_view_ui("Progetto Delta", str(manual), None, "it")
    finally:
        app.detect_docx_language = orig
    check("G3: when the language is uncertain the question is shown ABOVE the choices, nothing is reported yet",
          "Non riconosco con sicurezza" in str(o[7].get("value")) and o[7]["visible"] and skipped(o[3])
          and skipped(o[-1]), (o[7], o[3]))
    imp = next(d for d in deps if fn[d["id"]] == "<lambda>"
               and any(t[0] == app.home_import_manual._id for t in d.get("targets", [])))
    upd = app.demo.fns[imp["id"]].fn()
    check("G4: «Importa un manuale» opens the guided import even while a project is open",
          imp["outputs"][0] == app.legacy_onboarding_panel._id and upd[0]["visible"], upd)

    # ============================================================ L: interface language
    header = next(b for b in walk(app.demo) if "ms-ctxbar" in (getattr(b, "elem_classes", None) or []))
    sidebar = next(b for b in walk(app.demo) if isinstance(b, gr.Sidebar))
    check("L1: the interface language sits in the header, not in the sidebar (decision 8d)",
          any(c is app.ws_lang for c in walk(header)) and not any(c is app.ws_lang for c in walk(sidebar)), "")
    lang_dep = triggered_by(deps, app.ws_lang, "change")[0]
    after = [fn[d["id"]] for d in chain_after(deps, lang_dep)]
    check("L2: switching language re-renders the shell and the UX Phase 0 areas at once",
          {"shell_ui", "settings_project_ui", "release_profile_notice_ui"} <= set(after) and after.count("translation_panel_ui") == 2, after)
    CP2 = {"job_action_view_ui", "job_cancel_confirmed_ui", "work_package_view_ui", "settings_add_languages_ui",
           "settings_remove_language_ui", "settings_remove_confirmed_ui", "settings_remove_back_ui",
           "settings_project_ui", "release_profile_notice_ui", "release_profile_open_ui", "release_profile_create_ui",
           "release_export_profile_check_ui"}
    wrong = sorted({fn[d["id"]] for d in deps if fn[d["id"]] in CP2 and (d.get("inputs") or [None])[0] != app.ws_lang._id})
    check("L3: every UX Phase 0 callback receives the session language", not wrong, wrong)
    en = text(app._localized(app.settings_project_ui)("en", str(A))[0])
    check("L4: and renders in it (Settings › Project in English)",
          "Project name" in en and "Original document language" in en and "German (de-DE)" in en, en[:200])
    fake = lambda f: (lambda lang, *a: f(*a))  # a wrapper that forgets the language
    check("L5 (mutant): a wrapper that forgot the language would stay Italian -- caught",
          "Project name" not in text(fake(app.settings_project_ui)("en", str(A))[0]), "")
    js = app._SHELL_PERSIST_JS
    check("L6: the language is remembered on this computer, the operator only for this browser session (8c)",
          "#ms-ui-lang" in js and "#ms-operator" in js and "put(localStorage, LANG" in js
          and "put(sessionStorage, OP" in js and "put(localStorage, OP" not in js
          and "get(sessionStorage, OP)" in js and "get(localStorage, OP)" not in js, "")
    check("L6b: a restored operator never replaces a name the user has started typing",
          "if (e.isTrusted) typed = true" in js and "typed ||" in js and "if (!el.value)" in js, "")
    op_writers = [fn[d["id"]] for d in deps if app.ws_operator._id in d.get("outputs", [])]
    check("L7: restoring the operator is browser-side only -- no server event can overwrite it",
          op_writers == ["switch_language_ui"], op_writers)
    with i18n.use_language("it"):
        ui_label = app.t("context.ui_language")
    check("L8: the interface language is named as such, distinct from the project's languages",
          ui_label == "Lingua dell'interfaccia" and ui_label not in ("Lingue di traduzione",
                                                                    "Lingua del documento originale"), ui_label)

    # ============================================================ N: navigation
    def selected(u):
        return u.get("selected") if isinstance(u, dict) else getattr(u, "selected", None)

    nav_fail = {}
    for btn, tab in ((app.nav_impact, "impact"), (app.nav_home, "home"), (app.nav_release, "release")):
        d = [x for x in triggered_by(deps, btn, "click") if x["outputs"] and x["outputs"][0] == app.workspace_tabs._id][0]
        upd = app.demo.fns[d["id"]].fn()
        lit = [names.get(oid) for u, oid in zip(upd[1:], d["outputs"][1:]) if u.get("variant") == "primary"]
        if selected(upd[0]) != tab or lit != [names.get(btn._id)]:
            nav_fail[tab] = (selected(upd[0]), lit)
    check("N1: every section lights up its own sidebar entry -- Analisi impatto included", not nav_fail, nav_fail)
    # Checkpoint 4: nine pages, seven entries. Each two-page section is one
    # entry plus a sub-navigation; whichever of its pages is open, THAT
    # section's entry is the one lit, and every page stays one click away.
    sub_fail = {}
    for btn, tab, entry in ((app.subnav_impact_to_rev, "revision", app.nav_impact),
                            (app.subnav_rev_to_impact, "impact", app.nav_impact),
                            (app.subnav_tr_to_jobs, "jobs", app.nav_translation),
                            (app.subnav_jobs_to_tr, "translations", app.nav_translation)):
        d = [x for x in triggered_by(deps, btn, "click") if x["outputs"] and x["outputs"][0] == app.workspace_tabs._id][0]
        upd = app.demo.fns[d["id"]].fn()
        lit = [names.get(oid) for u, oid in zip(upd[1:], d["outputs"][1:]) if u.get("variant") == "primary"]
        if selected(upd[0]) != tab or lit != [names.get(entry._id)]:
            sub_fail[tab] = (selected(upd[0]), lit)
    reach = set()
    for b_ in list(app._NAV_BUTTONS) + [app.subnav_impact_to_rev, app.subnav_rev_to_impact,
                                         app.subnav_tr_to_jobs, app.subnav_jobs_to_tr]:
        for d in triggered_by(deps, b_, "click"):
            if d["outputs"] and d["outputs"][0] == app.workspace_tabs._id:
                reach.add(selected(app.demo.fns[d["id"]].fn()[0]))
    pages = {"home", "authoring", "impact", "revision", "translations", "jobs", "compliance", "release", "settings"}
    check("N1b: the sub-navigation opens its page and keeps the section's entry lit; every page is reachable",
          not sub_fail and reach == pages, (sub_fail, sorted(pages - reach)))
    back_tr = [fn[x["id"]] for x in chain_after(deps, [x for x in triggered_by(deps, app.subnav_jobs_to_tr, "click")
                                                        if x["outputs"][0] == app.workspace_tabs._id][0])]
    check("N1c: coming back to Traduzioni from its sub-page re-reads the reliability panel, like the sidebar entry",
          "translation_panel_ui" in back_tr, back_tr)
    quick = {}
    for btn, tab, entry in ((app.home_revision, "revision", app.nav_impact),
                            (app.home_translation, "translations", app.nav_translation),
                            (app.home_compliance, "compliance", app.nav_compliance),
                            (app.home_release, "release", app.nav_release)):
        d = [x for x in triggered_by(deps, btn, "click") if x["outputs"] and x["outputs"][0] == app.workspace_tabs._id][0]
        upd = app.demo.fns[d["id"]].fn()
        lit = [names.get(oid) for u, oid in zip(upd[1:], d["outputs"][1:]) if u.get("variant") == "primary"]
        if selected(upd[0]) != tab or lit != [names.get(entry._id)]:
            quick[tab] = (selected(upd[0]), lit)
    check("N2: each quick action opens exactly its page and lights up the matching sidebar entry (its section)", not quick, quick)
    tr_chain = [fn[x["id"]] for x in chain_after(deps, [x for x in triggered_by(deps, app.nav_translation, "click")
                                                        if x["outputs"][0] == app.workspace_tabs._id][0])]
    panel = text(app.translation_panel_ui(str(A), "topic", "T1", "it")[0])
    check("N3: arriving on Traduzioni re-reads its reliability panel for the open project",
          "translation_panel_ui" in tr_chain and "tedesca" in panel.lower(), (tr_chain, panel[:160]))
    stale = []
    for lg in ("it", "en"):
        for k, v in json.loads((BASE / "locales" / f"{lg}.json").read_text(encoding="utf-8")).items():
            # Section names the sidebar no longer uses, in either language.
            # Checkpoint 4 adds the names retired with the seven-entry sidebar.
            if re.search(r"Redazione (&amp;|&) contenuti|Authoring (&amp;|&) content|Job (&amp;|&) fornitori"
                         r"|<b>Release &rarr;|Lavorazioni e fornitori|<b>Rilascio</b>|in Rilascio\b|Vai a Rilascio"
                         if lg == "it" else r"Authoring (&amp;|&) content|Go to Release\b|<b>Release</b>", v):
                stale.append(f"{lg}:{k}")
    for n, line in enumerate((BASE / "app.py").read_text(encoding="utf-8").splitlines(), 1):
        code = line.split("#", 1)[0]
        if re.search(r"Redazione (&amp;|&) contenuti", code):
            stale.append(f"app.py:{n}")
    check("N4: no text sends the user to a section name that no longer exists", not stale, stale)
    pub_sync = {fn[d["id"]]: [names.get(o_, o_) for o_ in d["outputs"]]
                for d in triggered_by(deps, app.cp_pub, "change")}
    back = [names.get(d["outputs"][0]) for d in triggered_by(deps, app.c_pub, "blur") + triggered_by(deps, app.r_pub, "blur")]
    sh = app.shell_ui(str(A), "PUB-NOPE", "it")
    check("N5: one manual in context -- the header drives Contenuti and Pubblicazione, and they report back",
          list(pub_sync.values()) == [["c_pub", "r_pub"]] and back == ["cp_pub", "cp_pub"]
          and sh[1]["value"] == sh[2] == "PUB-MAIN" and sh[1]["choices"] == [("ALFA manual", "PUB-MAIN")],
          (pub_sync, back, sh[1]))
except BaseException as _crash:
    if not isinstance(_crash, AssertionError):
        results.append({"test": "suite interrupted by an unexpected error", "status": "FAIL",
                        "detail": f"{type(_crash).__name__}: {_crash}"[:400]})
    raise
finally:
    for _m in [m for m in sys.modules if m.startswith("ms_nav_cp3_mutant_")]:
        sys.modules.pop(_m, None)
    summary = {"suite": "Checkpoint 3 -- navigation and application structure", "tests": len(results),
               "pass": sum(r["status"] == "PASS" for r in results),
               "fail": sum(r["status"] == "FAIL" for r in results), "results": results}
    (BASE / "RISULTATO_NAVIGATION_CP3.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: summary[k] for k in ("suite", "tests", "pass", "fail")}, ensure_ascii=False))
    shutil.rmtree(OUT, ignore_errors=True)
    shutil.rmtree(_TEST_HOME, ignore_errors=True)
