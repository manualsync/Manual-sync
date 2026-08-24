from __future__ import annotations

import hashlib
import json
import shutil
import zipfile
from pathlib import Path

from docx import Document

from cat_tm import technical_signature, source_target_safe
from ccms_industrial import IndustrialRepository
from publishing_bridges import export_publication_docx
from release_manager import create_release_bundle

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "regressione_desktop_r13_out"
if OUT.exists():
    shutil.rmtree(OUT)
OUT.mkdir(parents=True)

results = []
def check(name, cond, detail=""):
    ok = bool(cond)
    results.append({"name": name, "status": "PASS" if ok else "FAIL", "detail": str(detail)})
    if not ok:
        raise AssertionError(f"{name}: {detail}")

# 1) CAT: ordinary lower-case hyphenated prose must NOT become immutable codes.
short_sig = technical_signature("Calculating the short-circuit current of the installation")
marine_sig = technical_signature("For ratings of marine-type approved drives")
check("Natural short-circuit is not a model code", "SHORT-CIRCUIT" not in short_sig.codes, short_sig)
check("Natural marine-type is not a model code", "MARINE-TYPE" not in marine_sig.codes, marine_sig)

# Real technical identifiers are still protected.
tech_sig = technical_signature("Use ACS880-01 with option C132 and module PHD-12.")
check("Real model ACS880-01 remains protected", "ACS880-01" in tech_sig.codes, tech_sig)
check("Real option C132 remains protected", "C132" in tech_sig.codes, tech_sig)
check("Real code PHD-12 remains protected", "PHD-12" in tech_sig.codes, tech_sig)

ok, reason = source_target_safe(
    "Calculating the short-circuit current of the installation",
    "Calcolo della corrente di cortocircuito dell’installazione",
)
check("Translated hyphenated prose passes bilingual technical gate", ok, reason)

# 2) Build a small complete release project.
project = OUT / "project"
project.mkdir()
repo = IndustrialRepository(project / "ManualSync_CCMS.sqlite")
repo.create_topic("T1", "Electrical ratings", "Legacy body.", source_lang="en-US", status="Approved", source_refs=["SPEC#1"])
repo.create_publication("PUB1", "Demo Drive Manual", ["T1"], source_lang="en-US")
repo.add_block("B1", "T1", "paragraph", position=1, text="Install the drive vertically.", status="Approved", source_refs=["SPEC#1"])
repo.set_topic_title_translation("T1", "it-IT", "\uFFEDValori nominali elettrici", status="Reviewed")
repo.set_block_translation("B1", "it-IT", "Installare il convertitore verticalmente.", status="Reviewed")
repo.close()

# 3) Customer publishing strips the legacy migration-only U+FFED heading marker.
repo = IndustrialRepository(project / "ManualSync_CCMS.sqlite")
out_doc = OUT / "heading_cleanup.docx"
export_publication_docx(repo, "PUB1", out_doc, target_lang="it-IT", include_source_refs=False)
repo.close()
doc = Document(out_doc)
headings = [p.text for p in doc.paragraphs if p.style and p.style.name.startswith("Heading")]
check("Legacy U+FFED heading marker removed", headings == ["Valori nominali elettrici"], headings)
meta_text = "\n".join(p.text for p in doc.paragraphs[:4])
check("Italian publication metadata labels localized", "Profilo:" in meta_text and "Lingua: it-IT" in meta_text, meta_text)

# 4) Seed historical files in persistent release folders, then create a new release.
release_dir = OUT / "release"
(release_dir / "ARTIFACTS").mkdir(parents=True)
(release_dir / "AUDIT").mkdir(parents=True)
(release_dir / "ARTIFACTS" / "OLD_R11.docx").write_bytes(b"historical artifact")
(release_dir / "AUDIT" / "OLD_R11.txt").write_text("historical audit", encoding="utf-8")

bundle = create_release_bundle(
    project,
    "PUB1",
    target_langs=["it-IT"],
    released_by="Regression R13",
    release_id="REL-R13",
    out_dir=release_dir,
)
zip_path = Path(bundle["release_zip"])
check("Release ZIP created", zip_path.exists() and zip_path.stat().st_size > 0, zip_path)

with zipfile.ZipFile(zip_path) as z:
    names = set(z.namelist())
    manifest = json.loads(z.read("RELEASE_MANIFEST.json"))
    declared = {"RELEASE_MANIFEST.json"}
    declared.update(a["path"] for a in manifest["artifacts"])
    declared.update(a["path"] for a in manifest["audit_files"])
    check("Historical artifact excluded from new ZIP", "ARTIFACTS/OLD_R11.docx" not in names, sorted(names))
    check("Historical audit excluded from new ZIP", "AUDIT/OLD_R11.txt" not in names, sorted(names))
    check("Release ZIP contains exactly manifest-declared current files", names == declared, {"names": sorted(names), "declared": sorted(declared)})
    check("All artifact checksums valid", all(hashlib.sha256(z.read(a["path"])).hexdigest() == a["sha256"] for a in manifest["artifacts"]), manifest["artifacts"])
    check("All audit checksums valid", all(hashlib.sha256(z.read(a["path"])).hexdigest() == a["sha256"] for a in manifest["audit_files"]), manifest["audit_files"])

summary = {
    "build": "1.0.0-desktop-r13-release-isolation-cat-safety-ui-cleanup-20260818",
    "tests": len(results),
    "pass": sum(r["status"] == "PASS" for r in results),
    "fail": sum(r["status"] == "FAIL" for r in results),
    "results": results,
}
(ROOT / "RISULTATO_REGRESSIONE_DESKTOP_R13.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
(ROOT / "RISULTATO_REGRESSIONE_DESKTOP_R13.txt").write_text(
    "\n".join(f"{r['status']} | {r['name']} | {r['detail']}" for r in results) + f"\n\nTOTAL {summary['pass']}/{summary['tests']} PASS\n",
    encoding="utf-8",
)
print(json.dumps({"tests": summary["tests"], "pass": summary["pass"], "fail": summary["fail"]}, indent=2))
