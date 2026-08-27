"""Milestone 3B.1 regression suite: legacy Word .doc onboarding.

Covers the isolated conversion module (legacy_doc_conversion.py) with a fake
soffice subprocess (no real LibreOffice required for CI/dev-machine runs),
and the app.py routing/provenance/no-regression behavior with a mocked
converter. This suite does NOT replace the mandatory manual acceptance test
with a real industrial .doc file via a real LibreOffice conversion.
"""
from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile
from pathlib import Path
from uuid import uuid4

BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE))

import gradio as gr
from docx import Document as DocxDocument

import legacy_doc_conversion as ldc
from ccms_industrial import IndustrialRepository
import app

TMP = Path(tempfile.mkdtemp(prefix="manualsync_3b1_"))
results = []


def check(name, cond, detail=""):
    ok = bool(cond)
    results.append({"test": name, "status": "PASS" if ok else "FAIL", "detail": str(detail)})
    if not ok:
        raise AssertionError(f"{name}: {detail}")


def make_project() -> Path:
    project_dir = TMP / f"proj_{uuid4().hex[:8]}"
    project_dir.mkdir(parents=True)
    IndustrialRepository(str(project_dir / "ManualSync_CCMS.sqlite")).close()
    return project_dir


def build_fixture_docx(path: Path) -> Path:
    d = DocxDocument()
    d.add_paragraph("1 - Introduzione")  # legacy heading pattern, no Heading style
    d.add_paragraph("Testo introduttivo del capitolo uno.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Parametro"; t.cell(0, 1).text = "Valore"
    t.cell(1, 0).text = "Tensione"; t.cell(1, 1).text = "24V"
    d.add_paragraph("Testo dopo la tabella, stesso capitolo.")
    d.save(str(path))
    return path


SIM_SCRIPT_SOURCE = '''
import os, sys
from pathlib import Path

def main():
    args = sys.argv[1:]
    mode = os.environ.get("MANUALSYNC_TEST_DOC_SIM", "success")
    outdir = None
    for i, a in enumerate(args):
        if a == "--outdir" and i + 1 < len(args):
            outdir = args[i + 1]
    doc_arg = args[-1]
    out_path = Path(outdir) / (Path(doc_arg).stem + ".docx")

    if mode == "timeout":
        import time
        time.sleep(600)
        return 0
    if mode == "crash":
        return 1
    if mode == "exit0_no_output":
        return 0
    if mode == "zero_byte":
        out_path.write_bytes(b"")
        return 0
    if mode == "corrupt":
        out_path.write_bytes(b"not a real docx, just garbage bytes")
        return 0
    if mode == "success":
        from docx import Document
        d = Document()
        d.add_paragraph("1 - Introduzione")
        d.add_paragraph("Testo introduttivo del capitolo uno.")
        tb = d.add_table(rows=2, cols=2)
        tb.cell(0, 0).text = "Parametro"; tb.cell(0, 1).text = "Valore"
        tb.cell(1, 0).text = "Tensione"; tb.cell(1, 1).text = "24V"
        d.add_paragraph("Testo dopo la tabella, stesso capitolo.")
        d.save(str(out_path))
        return 0
    return 1

if __name__ == "__main__":
    sys.exit(main())
'''

from contextlib import contextmanager


@contextmanager
def _fake_success_converter(doc_path):
    yield FIXTURE_DOCX


@contextmanager
def _fake_failing_converter(doc_path):
    raise ldc.LegacyDocConverterUnavailable("simulated: no converter available")
    yield None  # pragma: no cover - unreachable, keeps this a generator function


try:
    FIXTURE_DOCX = build_fixture_docx(TMP / "fixture_converted.docx")
    sim_path = TMP / "sim_soffice.py"
    sim_path.write_text(SIM_SCRIPT_SOURCE, encoding="utf-8")
    fake_doc = TMP / "sample_legacy.doc"
    fake_doc.write_bytes(b"placeholder legacy .doc content; sim converter ignores it")

    def run_sim(mode):
        os.environ["MANUALSYNC_TEST_DOC_SIM"] = mode
        return [sys.executable, str(sim_path)]

    # ------------------------------------------------------------------
    # Group A: legacy_doc_conversion.py in isolation (real subprocess,
    # fake soffice binary) -- no LibreOffice required.
    # ------------------------------------------------------------------
    created_dirs = []
    _orig_mkdtemp = ldc.tempfile.mkdtemp

    def _tracking_mkdtemp(*a, **kw):
        d = _orig_mkdtemp(*a, **kw)
        created_dirs.append(d)
        return d

    ldc.tempfile.mkdtemp = _tracking_mkdtemp
    try:
        # NOTE: this used to assert find_soffice_binary() is None as a "real
        # fact" about the dev machine. That stopped being true once LibreOffice
        # was installed (with explicit authorization) for the 3B.1 manual
        # acceptance test -- a legitimate, expected environment change, not a
        # product regression. The "converter absent" behavior below is now
        # verified by forcing that condition via monkeypatch instead of
        # depending on incidental machine state, so this suite passes
        # identically whether or not LibreOffice happens to be installed.
        detected = ldc.find_soffice_binary()
        check("find_soffice_binary() resolves without raising (informational; real state noted, not asserted)",
              True, f"currently resolves to: {detected}")

        _orig_find_soffice = ldc.find_soffice_binary
        ldc.find_soffice_binary = lambda: None
        try:
            with ldc.convert_legacy_doc_to_docx(fake_doc):
                pass
            check("converter absent raises LegacyDocConverterUnavailable", False, "no exception")
        except ldc.LegacyDocConverterUnavailable:
            check("converter absent raises LegacyDocConverterUnavailable", True)
        finally:
            ldc.find_soffice_binary = _orig_find_soffice

        try:
            with ldc.convert_legacy_doc_to_docx(TMP / "not_a_doc.txt", soffice_command=run_sim("success")):
                pass
            check(".doc-only guard rejects non-.doc input", False, "no exception")
        except ValueError:
            check(".doc-only guard rejects non-.doc input", True)

        try:
            with ldc.convert_legacy_doc_to_docx(fake_doc, timeout=1.0, soffice_command=run_sim("timeout")):
                pass
            check("timeout fails closed", False, "no exception")
        except ldc.LegacyDocConversionError as exc:
            check("timeout fails closed", "imeout" in str(exc), str(exc))

        try:
            with ldc.convert_legacy_doc_to_docx(fake_doc, soffice_command=run_sim("crash")):
                pass
            check("nonzero exit / crash fails closed", False, "no exception")
        except ldc.LegacyDocConversionError:
            check("nonzero exit / crash fails closed", True)

        try:
            with ldc.convert_legacy_doc_to_docx(fake_doc, soffice_command=run_sim("exit0_no_output")):
                pass
            check("exit 0 with missing output fails closed (exit code not trusted)", False, "no exception")
        except ldc.LegacyDocConversionError:
            check("exit 0 with missing output fails closed (exit code not trusted)", True)

        try:
            with ldc.convert_legacy_doc_to_docx(fake_doc, soffice_command=run_sim("zero_byte")):
                pass
            check("zero-byte output fails closed", False, "no exception")
        except ldc.LegacyDocConversionError:
            check("zero-byte output fails closed", True)

        try:
            with ldc.convert_legacy_doc_to_docx(fake_doc, soffice_command=run_sim("corrupt")):
                pass
            check("corrupt/unopenable output fails closed", False, "no exception")
        except ldc.LegacyDocConversionError:
            check("corrupt/unopenable output fails closed", True)

        produced_holder = {}
        with ldc.convert_legacy_doc_to_docx(fake_doc, soffice_command=run_sim("success")) as produced:
            produced_holder["p"] = produced
            check("successful conversion yields an existing non-empty file",
                  produced.exists() and produced.stat().st_size > 0, str(produced))
            opened = DocxDocument(str(produced))
            paras = [p.text for p in opened.paragraphs]
            check("converted content preserves legacy heading text", "1 - Introduzione" in paras, paras)
            check("converted content preserves the table", len(opened.tables) == 1, len(opened.tables))
        check("temp output cleaned up once the with-block exits",
              not produced_holder["p"].exists(), str(produced_holder["p"]))

        check("no leaked temporary conversion directories across all Group A cases",
              all(not Path(d).exists() for d in created_dirs), created_dirs)
    finally:
        ldc.tempfile.mkdtemp = _orig_mkdtemp
        os.environ.pop("MANUALSYNC_TEST_DOC_SIM", None)

    # ------------------------------------------------------------------
    # Group B: app.py routing + provenance, with the converter mocked
    # (no LibreOffice needed to test routing/provenance logic).
    # ------------------------------------------------------------------
    _orig_convert = app.convert_legacy_doc_to_docx

    doc_project_dir = None
    try:
        project_dir = make_project()
        doc_path = TMP / "LineControl_Manuale.doc"
        doc_path.write_bytes(b"placeholder; converter is mocked for this test")
        app.convert_legacy_doc_to_docx = _fake_success_converter
        msg, topics, blocks = app.import_manual_ui(str(project_dir), str(doc_path), "it-IT", "PUB-MAIN", "Test Doc")
        check(".doc import (routed through mocked converter) produces topics", not topics.empty, str(topics))
        refs = " | ".join(blocks["source_refs"].astype(str))
        check(".doc provenance references the ORIGINAL .doc filename", "LineControl_Manuale.doc" in refs, refs)
        check(".doc provenance does not leak the temporary .docx name", ".docx" not in refs, refs)
        titles = list(topics.sort_values("topic_id")["title"])
        check(".doc import reuses 3B legacy heading inference unchanged", "1 - Introduzione" in titles, titles)
        simple = app._simple_review_table_for_project(str(project_dir))
        table_rows = simple[simple["Tipo di contenuto"] == "Tabella"]
        check(".doc import reuses 3B table preview unchanged",
              not table_rows.empty and all("Tabella" in str(p) for p in table_rows["Anteprima del testo"]),
              list(table_rows["Anteprima del testo"]))
        check(".doc import reuses 3B 'Sezione N' labeling unchanged",
              all(str(p).startswith("Sezione ") for p in simple["Pagina di origine"] if p is not None),
              list(simple["Pagina di origine"]))
        doc_project_dir = project_dir
    finally:
        app.convert_legacy_doc_to_docx = _orig_convert

    try:
        project_dir = make_project()
        doc_path = TMP / "Broken.doc"
        doc_path.write_bytes(b"x")
        app.convert_legacy_doc_to_docx = _fake_failing_converter
        try:
            app.import_manual_ui(str(project_dir), str(doc_path), "it-IT", "PUB-MAIN", "Test")
            check("missing converter raises a guided gr.Error", False, "no exception")
        except gr.Error as exc:
            check("missing converter error text matches the required guided message",
                  exc.message == app._LEGACY_DOC_FRIENDLY_ERROR, exc.message)
    finally:
        app.convert_legacy_doc_to_docx = _orig_convert

    # guided onboarding entry point: routing + converter-failure wrapping
    try:
        calls = {"docx": 0, "pdf": 0}

        def fake_docx_lang(path):
            calls["docx"] += 1
            return ("it", 0.95, "contenuto")

        def fake_pdf_lang(path):
            calls["pdf"] += 1
            raise AssertionError("detect_document_language must not be used for .doc input")

        captured = {}

        def fake_success_outputs(name, pdf_file, source_lang):
            captured.update(name=name, pdf_file=pdf_file, source_lang=source_lang)
            return ("SENTINEL_SUCCESS",)

        orig_docx_lang, orig_pdf_lang, orig_success = app.detect_docx_language, app.detect_document_language, app._guided_success_outputs
        app.detect_docx_language, app.detect_document_language, app._guided_success_outputs = fake_docx_lang, fake_pdf_lang, fake_success_outputs
        app.convert_legacy_doc_to_docx = _fake_success_converter
        try:
            doc_path = TMP / "Guided.doc"
            doc_path.write_bytes(b"placeholder")
            result = app.guided_new_project_from_pdf_ui("Nome progetto", str(doc_path))
            check(".doc guided routing uses detect_docx_language, never the PDF detector",
                  calls["docx"] == 1 and calls["pdf"] == 0, calls)
            check(".doc guided routing keeps the ORIGINAL .doc path flowing onward (provenance-safe)",
                  captured.get("pdf_file") == str(doc_path), captured)
            check(".doc guided routing reaches success outputs", result == ("SENTINEL_SUCCESS",), result)
        finally:
            app.detect_docx_language, app.detect_document_language, app._guided_success_outputs = orig_docx_lang, orig_pdf_lang, orig_success
            app.convert_legacy_doc_to_docx = _orig_convert

        app.convert_legacy_doc_to_docx = _fake_failing_converter
        try:
            doc_path = TMP / "GuidedBroken.doc"
            doc_path.write_bytes(b"placeholder")
            try:
                app.guided_new_project_from_pdf_ui("Nome progetto", str(doc_path))
                check("guided .doc converter failure raises a guided gr.Error", False, "no exception")
            except gr.Error as exc:
                check("guided .doc converter failure message matches the required guided text",
                      exc.message == app._LEGACY_DOC_FRIENDLY_ERROR, exc.message)
        finally:
            app.convert_legacy_doc_to_docx = _orig_convert
    finally:
        app.convert_legacy_doc_to_docx = _orig_convert

    # ------------------------------------------------------------------
    # Group C: no-regression -- PDF (3A), native .docx (3B), human approval.
    # ------------------------------------------------------------------
    project_dir = make_project()
    pdf_src = BASE / "benchmark" / "siemens" / "old_EN.pdf"
    pdf_path = TMP / "old_EN.pdf"
    shutil.copy2(pdf_src, pdf_path)
    msg, topics, blocks = app.import_manual_ui(str(project_dir), str(pdf_path), "en-US", "PUB-MAIN", "PDF Test")
    check("PDF import (3A) unaffected by 3B.1 changes", not topics.empty, str(topics))

    project_dir = make_project()
    docx_path = TMP / "Native.docx"
    shutil.copy2(FIXTURE_DOCX, docx_path)
    calls_native = {"n": 0}

    def spy_converter(*a, **kw):
        calls_native["n"] += 1
        raise AssertionError("legacy .doc converter must never be invoked for a native .docx import")

    app.convert_legacy_doc_to_docx = spy_converter
    try:
        msg, topics, blocks = app.import_manual_ui(str(project_dir), str(docx_path), "it-IT", "PUB-MAIN", "Native")
        check("native .docx import (3B) unaffected by 3B.1 changes", not topics.empty, str(topics))
        check("native .docx import never invokes the .doc converter", calls_native["n"] == 0, calls_native)
        titles = list(topics.sort_values("topic_id")["title"])
        check("native .docx legacy heading inference still intact", "1 - Introduzione" in titles, titles)
    finally:
        app.convert_legacy_doc_to_docx = _orig_convert

    check("a .doc-originated project exists for the human-approval check", doc_project_dir is not None, "setup failure")
    try:
        app.approve_source_ui(str(doc_project_dir), "PUB-MAIN", "QA Reviewer", False)
        check("approval without explicit confirmation still fails closed", False, "no exception raised")
    except gr.Error as exc:
        check("approval without explicit confirmation still fails closed", "conferma esplicita" in exc.message, exc.message)
    approval_msg = app.approve_source_ui(str(doc_project_dir), "PUB-MAIN", "QA Reviewer", True)
    check("human approval workflow unchanged for .doc-derived content", "approvata" in approval_msg.casefold(), approval_msg)

    # ------------------------------------------------------------------
    # Group D: architectural guardrails (no second importer, no Word/Java deps).
    # ------------------------------------------------------------------
    src_ldc = (BASE / "legacy_doc_conversion.py").read_text(encoding="utf-8")
    check("no shell=True in the conversion subprocess call", "shell=True" not in src_ldc, "found shell=True")
    check("no Microsoft Word COM automation dependency introduced",
          not any(tok in src_ldc for tok in ("win32com", "comtypes", "Dispatch(")), "forbidden token found")
    check("no Java/Apache POI bridge dependency introduced",
          not any(tok in src_ldc for tok in ("jpype", "py4j", "jaydebeapi")), "forbidden token found")
    src_app = (BASE / "app.py").read_text(encoding="utf-8")
    check("app.py introduces no Word COM / Java bridge tokens",
          not any(tok in src_app for tok in ("win32com", "comtypes", "jpype", "py4j")), "forbidden token found")

    payload = {
        "build": "Milestone 3B.1 - legacy .doc onboarding",
        "tests": len(results),
        "pass": sum(r["status"] == "PASS" for r in results),
        "fail": sum(r["status"] == "FAIL" for r in results),
        "results": results,
    }
    (BASE / "RISULTATO_LEGACY_DOC_3B1.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (BASE / "RISULTATO_LEGACY_DOC_3B1.txt").write_text(
        "\n".join(f"{r['status']} | {r['test']} | {r['detail']}" for r in results) + f"\n\nPASS: {payload['pass']}/{payload['tests']}\n",
        encoding="utf-8",
    )
    try:
        print(json.dumps(payload, ensure_ascii=False, indent=2))
    except UnicodeEncodeError:
        pass  # Windows console codepage can't render some emoji in result details; RISULTATO_* files already hold the full UTF-8 output.
    if payload["fail"]:
        raise SystemExit(1)
finally:
    shutil.rmtree(TMP, ignore_errors=True)
