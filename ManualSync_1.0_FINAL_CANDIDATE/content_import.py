"""CP3.3: the PDF and DOCX readers of the topic import, moved verbatim out of authoring_ccms.

They read a document -- a PDF through the pdf_backend reading contract, a DOCX
through python-docx -- and turn it into topic groups, which they hand to the
repository-side write step that stayed in authoring_ccms
(_write_topic_groups_to_repo, the only code that writes, and holds every SQL
statement of the import). IDML already had its own reader (idml_import);
its repository transaction, import_idml_as_topics, also stayed in authoring_ccms.

Before CP3.3 these readers lived inside authoring_ccms, so the content
repository -- and every service built on it -- loaded PyMuPDF, numpy, PIL and
python-docx just to open a project. The code below is unchanged except for one
line: import_pdf_as_topics asks pdf_backend for the layout instead of calling
layout_engine.extract_layout directly. No SQL here.

authoring_ccms still serves these names (lazily, on first access), so
`from authoring_ccms import import_pdf_as_topics` keeps working.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document as DocxDocument

import pdf_backend
from authoring_ccms import ContentRepository, _write_topic_groups_to_repo


def _pdf_running_header(text: str) -> bool:
    t = re.sub(r"\s+", " ", str(text or "")).strip()
    return bool(re.fullmatch(r"(?:Technical data\s+\d+|\d+\s+Technical data)", t, re.I))


def _pdf_semantic_heading(text: str, heading_flag: bool) -> bool:
    """Return True only for title-like PDF blocks.

    PDF font/layout heuristics frequently mark chart labels, bullets, formulas and
    full sentences as headings. Structured migration must be more conservative:
    false-positive topics are worse than keeping a line inside the current topic.
    """
    if not heading_flag:
        return False
    t = re.sub(r"\s+", " ", str(text or "")).strip().lstrip("￭").strip()
    if not t or _pdf_running_header(t):
        return False
    if t.startswith(("•", "- ", "– ", "— ")):
        return False
    if t.lower() in {"note", "note:", "where", "u", "k", "p/pn", "u/un"}:
        return False
    letters = sum(ch.isalpha() for ch in t)
    digits = sum(ch.isdigit() for ch in t)
    words = re.findall(r"[A-Za-zÀ-ÿ][A-Za-zÀ-ÿ0-9+./-]*", t)
    if letters < 3:
        return False
    if len(t) > 120 or len(words) > 14:
        return False
    if t.endswith((".", ":", ";")):
        return False
    # Formula/chart labels must never become topic titles.
    if ("=" in t or "×" in t or "√" in t or "·" in t) and len(words) <= 8:
        return False
    if digits > letters * 1.5 and len(words) < 3:
        return False
    # Long sentence-like blocks are body content even when font heuristics say heading.
    starters = (
        "the ", "this ", "these ", "those ", "when ", "for ", "make ", "derate ",
        "enabling ", "with ", "always ", "use ", "install ", "protect ", "if ", "an ", "a ",
    )
    if t.lower().startswith(starters) and len(words) > 5:
        return False
    return True


def _pdf_block_type(text: str) -> str:
    t = re.sub(r"\s+", " ", str(text or "")).strip()
    low = t.lower()
    if low.startswith(("warning:", "danger:", "caution:")):
        return "warning"
    if low == "note:" or low.startswith("note: "):
        return "note"
    if t.startswith(("•", "- ", "– ", "— ")):
        return "list_item"
    return "paragraph"


def import_pdf_as_topics(repo: ContentRepository, pdf_path: str | Path, *, source_lang: str = "en-US",
                         page_spec: str = "", id_prefix: str = "TOPIC", source_ref_prefix: str | None = None) -> pd.DataFrame:
    """Conservative PDF -> structured-topic migration.

    RC1.9 hardening:
    - accepts only semantic/title-like heading candidates;
    - discards recurring page headers instead of turning them into prose/topics;
    - keeps prose/list/note/warning blocks in reading order;
    - when the industrial CCMS API is available, preserves detected PDF tables as
      first-class structured table blocks instead of flattening them into prose.

    This remains a migration draft. PDF is not a lossless structured-authoring
    source and human review is mandatory before source approval.
    """
    layout = pdf_backend.default_reader().extract_layout(pdf_path, page_spec)
    src_name = source_ref_prefix or Path(pdf_path).name

    # Merge text/table events in physical reading order.
    events: list[dict[str, Any]] = []
    for block in layout.text_blocks:
        txt = re.sub(r"\s+", " ", block.text or "").strip()
        if not txt:
            continue
        y = float(block.bbox[1]) if getattr(block, "bbox", None) else 0.0
        events.append({
            "kind": "text", "page0": int(block.page), "y": y, "text": txt,
            "heading": bool(getattr(block, "heading", False)),
        })
    for table in getattr(layout, "tables", []) or []:
        y = float(table.bbox[1]) if getattr(table, "bbox", None) else 0.0
        events.append({
            "kind": "table", "page0": int(table.page), "y": y,
            "table_index": int(table.index), "rows": table.rows,
        })
    events.sort(key=lambda e: (e["page0"], e["y"], 0 if e["kind"] == "text" else 1))

    groups: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    dropped_headers = 0

    def ensure_current(page0: int) -> dict[str, Any]:
        nonlocal current
        if current is None:
            current = {"title": "Introduction", "items": [], "page": page0 + 1}
        return current

    for ev in events:
        if ev["kind"] == "text":
            txt = ev["text"]
            if _pdf_running_header(txt):
                dropped_headers += 1
                continue
            if _pdf_semantic_heading(txt, ev["heading"]):
                if current and current["items"]:
                    groups.append(current)
                current = {
                    "title": txt.lstrip("￭").strip()[:180],
                    "items": [],
                    "page": ev["page0"] + 1,
                }
            else:
                ensure_current(ev["page0"])["items"].append({
                    "kind": "text", "text": txt, "page": ev["page0"] + 1,
                    "block_type": _pdf_block_type(txt),
                })
        else:
            ensure_current(ev["page0"])["items"].append({
                "kind": "table", "rows": ev["rows"], "page": ev["page0"] + 1,
                "table_index": ev["table_index"],
            })
    if current and current["items"]:
        groups.append(current)

    frame = _write_topic_groups_to_repo(repo, groups, source_lang=source_lang, id_prefix=id_prefix, src_name=src_name)
    frame.attrs["dropped_running_headers"] = dropped_headers
    return frame


_LEGACY_DOCX_HEADING_RE = re.compile(r"^\d{1,3}\s*[-–—]\s*[A-ZÀ-Ý][^.]{2,78}$")


def _looks_like_legacy_docx_heading(text: str) -> bool:
    """Conservative fallback for chapter titles with no Word Heading style.

    Real Heading N styles remain the primary signal (checked first, unchanged)
    — this only fires when the style is missing/wrong, and only for a narrow,
    strong pattern: "<number> - <Title>" (e.g. "1 - Introduzione",
    "8 - Gestione segnalazioni - Allarmi"), anchored on the whole line so it
    cannot match inside a normal sentence. Deliberately rejects:
      - numbered lists using ". " or ") " instead of " - " (e.g. "1. Press...");
      - technical ranges/values where digits follow the dash (e.g. "20 - 25°C"),
        because the text after the dash must start with an uppercase letter;
      - long sentences or anything containing a period, which a genuine short
        chapter title normally does not.
    """
    return bool(_LEGACY_DOCX_HEADING_RE.match(str(text or "").strip()))


def _iter_docx_block_items(doc):
    """Yield ("paragraph"|"table", obj) in REAL document order.

    doc.paragraphs and doc.tables are separate flat lists in python-docx and do
    NOT preserve their relative position in the file — a table between two
    paragraphs would otherwise be lost from the sequence and misattributed.
    This walks the underlying XML body directly (the standard python-docx
    recipe for document-order iteration) so a table stays exactly where it
    was, between the paragraphs that precede and follow it.
    """
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    p_tag, tbl_tag = qn("w:p"), qn("w:tbl")
    for child in doc.element.body.iterchildren():
        if child.tag == p_tag:
            yield "paragraph", Paragraph(child, doc)
        elif child.tag == tbl_tag:
            yield "table", Table(child, doc)


def import_docx_as_topics(repo: ContentRepository, docx_path: str | Path, *, source_lang: str = "en-US",
                           id_prefix: str = "TOPIC", source_ref_prefix: str | None = None) -> pd.DataFrame:
    """Conservative Word (.docx) -> structured-topic migration.

    Mirrors import_pdf_as_topics: extracts headings/paragraphs/tables in real
    document order (python-docx has no page concept, so sections are numbered
    sequentially — shown to the guided user as "Sezione N", not a technical
    index), then writes them through the SAME shared repo-writing step used
    by the PDF path (_write_topic_groups_to_repo) so topic/block/table
    handling, status, and evidence tracking are identical for both formats.

    This remains a migration draft. Human review is mandatory before source
    approval, same as the PDF path.
    """
    doc = DocxDocument(str(docx_path))
    src_name = source_ref_prefix or Path(docx_path).name

    groups: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    section_no = 0

    def ensure_current() -> dict[str, Any]:
        nonlocal current, section_no
        if current is None:
            section_no += 1
            current = {"title": "Introduction", "items": [], "page": section_no}
        return current

    for kind, block in _iter_docx_block_items(doc):
        if kind == "paragraph":
            text = re.sub(r"\s+", " ", block.text or "").strip()
            if not text:
                continue
            style = block.style.name if block.style else ""
            # Real Heading N style is checked first and always wins; the legacy
            # pattern is only a fallback for paragraphs that aren't Heading-styled.
            if style.startswith("Heading") or _looks_like_legacy_docx_heading(text):
                if current and current["items"]:
                    groups.append(current)
                section_no += 1
                current = {"title": text[:180], "items": [], "page": section_no}
            else:
                g = ensure_current()
                g["items"].append({"kind": "text", "text": text, "page": g["page"], "block_type": _pdf_block_type(text)})
        else:  # table, encountered exactly where it appears in the document
            rows = [[cell.text for cell in row.cells] for row in block.rows]
            g = ensure_current()
            table_index = sum(1 for it in g["items"] if it["kind"] == "table")
            g["items"].append({"kind": "table", "rows": rows, "page": g["page"], "table_index": table_index})

    if current and current["items"]:
        groups.append(current)

    return _write_topic_groups_to_repo(repo, groups, source_lang=source_lang, id_prefix=id_prefix, src_name=src_name)
